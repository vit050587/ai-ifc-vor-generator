"""Разбор и рендер типового шаблона ВОР.

Шаблон — ``data/Типовой ВОР по КР.xlsx``, лист ``ВОР КР_ЖС``. Структура листа:

* колонки: ``A`` — № п/п, ``B`` — Наименование работ, ``C`` — Ед. изм.,
  ``D`` — Объем, ``E`` — Расчет объема, ``F`` — Обоснование объема работ;
* строки-заголовки: часть здания (``Подземная часть здания …``), раздел
  (``Раздел N. …``), подраздел (``Подраздел N. …``), геометрический/блочный
  заголовок (``Стены толщиной … мм``, ``Плита рампы`` и т. п.);
* строки-работы — целочисленный № п/п (``1``, ``2``, …);
* строки-ресурсы — дробный № п/п (``1.1``, ``2.3``, …);
* строки-варианты — без номера, но с наименованием (``-объемом до 5 м3``).

Модуль предоставляет:

* :func:`parse_template` — разбор листа в список :class:`TemplateRow`
  с восстановленным контекстом (часть/раздел/подраздел/тип конструкции);
* :func:`render_template` — запись заполненного ВОР в копию шаблона
  с сохранением структуры и форматирования.
"""

import copy
import os
import re
from dataclasses import dataclass, field
from typing import Dict, List, Optional, Tuple

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from src.core.logger import setup_logger
from . import mapping

logger = setup_logger(__name__)

# Имя типового шаблона (в каталоге data/).
TEMPLATE_FILENAME = "Типовой ВОР по КР.xlsx"
DEFAULT_SHEET = "ВОР КР_ЖС"

# Каталог data/ относительно корня проекта.
_DATA_DIR = os.path.join(
    os.path.dirname(os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))),
    "data",
)


def resolve_template_path(data_dir: str = None, filename: str = TEMPLATE_FILENAME) -> str:
    """Находит файл типового шаблона в каталоге ``data/``.

    Имя файла на диске может быть записано в разной Unicode-форме (NFC/NFD),
    поэтому сравнение выполняется после нормализации.
    """
    import unicodedata

    data_dir = data_dir or _DATA_DIR
    target = unicodedata.normalize("NFC", filename)
    try:
        for name in os.listdir(data_dir):
            if unicodedata.normalize("NFC", name) == target:
                return os.path.join(data_dir, name)
    except FileNotFoundError:
        pass
    return os.path.join(data_dir, filename)


# Путь к типовому шаблону по умолчанию.
DEFAULT_TEMPLATE_PATH = resolve_template_path()

# Количество колонок шаблона (A..F).
COLUMNS = 6
# Первая строка с данными (строки 1–3 — шапка таблицы).
FIRST_DATA_ROW = 4

_INT_RE = re.compile(r"^\d+$")
_DEC_RE = re.compile(r"^\d+[.,]\d+$")


@dataclass
class TemplateRow:
    """Строка типового шаблона с восстановленным контекстом."""

    index: int                       # 1-based номер строки в листе
    kind: str                        # part|section|subsection|block|work|resource|variant|empty
    number: str = ""                 # значение колонки A
    name: str = ""                   # значение колонки B
    unit: str = ""                   # значение колонки C
    part_text: Optional[str] = None  # текст текущей части здания
    part: Optional[str] = None       # underground|above|parking
    section: Optional[str] = None    # текст текущего раздела
    subsection: Optional[str] = None  # текст текущего подраздела
    element_kind: Optional[str] = None  # тип конструкции (mapping.detect_kind)
    geometry: Optional[str] = None   # толщина/сечение/площадь текущего блока
    block_index: Optional[int] = None  # индекс строки-заголовка блока
    parking: bool = False            # строка относится к части «Автостоянка»


def _cell(ws, row: int, col: int):
    return ws.cell(row=row, column=col).value


def _text(value) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()


def _classify(a: str, b: str, c: str) -> str:
    """Классифицирует строку шаблона по значениям колонок A/B/C."""
    if not a and not b and not c:
        return "empty"
    if a.startswith("Раздел "):
        return "section"
    if a.startswith("Подраздел "):
        return "subsection"
    if _INT_RE.match(a):
        return "work"
    if _DEC_RE.match(a):
        return "resource"
    if a and not b and not c:
        # Заголовок без номера и наименования — часть здания или блок.
        return "part" if mapping.detect_part(a) else "block"
    if not a and b:
        return "variant"
    return "other"


def parse_template(path: str = None, sheet_name: str = DEFAULT_SHEET) -> List[TemplateRow]:
    """Разбирает лист типового шаблона в список :class:`TemplateRow`."""
    path = path or DEFAULT_TEMPLATE_PATH
    wb = load_workbook(path, data_only=False)
    if sheet_name not in wb.sheetnames:
        raise ValueError(f"В шаблоне нет листа «{sheet_name}»")
    ws = wb[sheet_name]

    rows: List[TemplateRow] = []
    part_text = part = None
    section = subsection = None
    element_kind = None
    geometry = None
    block_index = None
    # Признак части «Автостоянка»: в шаблоне её подразделы озаглавлены как
    # «Подземная/Надземная часть», поэтому признак запоминается до конца листа
    # (автостоянка — последний раздел шаблона) и переносится во все строки.
    parking = False

    for r in range(FIRST_DATA_ROW, ws.max_row + 1):
        a = _text(_cell(ws, r, 1))
        b = _text(_cell(ws, r, 2))
        c = _text(_cell(ws, r, 3))
        kind = _classify(a, b, c)

        if kind == "empty":
            rows.append(TemplateRow(index=r, kind="empty", parking=parking))
            continue

        if kind == "part":
            part_text = a
            part = mapping.detect_part(a)
            section = subsection = None
            element_kind = mapping.detect_kind(a)
            geometry = None
            block_index = None
            if part == "parking":
                parking = True
        elif kind == "section":
            section = a
            subsection = None
            element_kind = mapping.detect_kind(a) or element_kind
            geometry = None
            block_index = None
        elif kind == "subsection":
            subsection = a
            element_kind = mapping.detect_kind(a) or element_kind
            geometry = None
            block_index = None
        elif kind == "block":
            block_index = r
            geometry = mapping.detect_geometry(a)
            block_kind = mapping.detect_kind(a)
            if block_kind:
                element_kind = block_kind

        rows.append(TemplateRow(
            index=r, kind=kind, number=a, name=b, unit=c,
            part_text=part_text, part=part, section=section,
            subsection=subsection, element_kind=element_kind,
            geometry=geometry if kind in ("block", "work", "resource", "variant") else None,
            block_index=block_index, parking=parking,
        ))

    logger.info("Шаблон ВОР разобран: строк %d (лист «%s»)", len(rows), sheet_name)
    return rows


# =====================================================================
#  Рендер заполненного ВОР
# =====================================================================

@dataclass
class OutputRow:
    """Строка итогового файла ВОР.

    ``src_index`` — номер строки-образца в шаблоне (для копирования стилей и
    текстов). ``volume``/``justification`` переопределяют значения колонок
    ``D``/``F`` (``None`` — оставить значение шаблона/пусто).
    """

    src_index: int
    label: Optional[str] = None          # переопределение колонки A (заголовок блока)
    volume: Optional[float] = None
    justification: Optional[str] = None
    name_override: Optional[str] = None  # переопределение колонки B
    extra: Dict[int, object] = field(default_factory=dict)


def render_template(output_rows: List[OutputRow], path: str = None,
                    sheet_name: str = DEFAULT_SHEET, output_path: str = None) -> str:
    """Записывает заполненный ВОР в копию шаблона.

    Структура листа (заголовки, порядок разделов, форматирование) сохраняется:
    строки-заголовки копируются из шаблона, а строки работ/ресурсов получают
    заполненные столбцы «Объем» и «Обоснование объема работ». Дополнительные
    строки (при группировке по геометрии) копируют стиль строки-образца.
    """
    path = path or DEFAULT_TEMPLATE_PATH
    if not output_path:
        raise ValueError("Не задан путь для сохранения ВОР")

    # Шаблон-источник стилей/текстов (только чтение).
    src = load_workbook(path, data_only=False)
    src_ws = src[sheet_name]

    out = load_workbook(path, data_only=False)
    ws = out[sheet_name]

    max_row = ws.max_row
    if max_row >= FIRST_DATA_ROW:
        ws.delete_rows(FIRST_DATA_ROW, max_row - FIRST_DATA_ROW + 1)

    for offset, row in enumerate(output_rows):
        target = FIRST_DATA_ROW + offset
        sample = src_ws[row.src_index]

        # Переносим стили строки-образца (и высоту строки).
        for col in range(1, COLUMNS + 1):
            src_cell = sample[col - 1]
            dst_cell = ws.cell(row=target, column=col)
            dst_cell._style = copy.copy(src_cell._style)
        ws.row_dimensions[target].height = src_ws.row_dimensions[row.src_index].height

        # Значения колонок A/B/C берём из строки-образца.
        ws.cell(row=target, column=1).value = _cell(src_ws, row.src_index, 1)
        ws.cell(row=target, column=2).value = _cell(src_ws, row.src_index, 2)
        ws.cell(row=target, column=3).value = _cell(src_ws, row.src_index, 3)
        ws.cell(row=target, column=5).value = _cell(src_ws, row.src_index, 5)

        if row.label is not None:
            ws.cell(row=target, column=1).value = row.label
        if row.name_override is not None:
            ws.cell(row=target, column=2).value = row.name_override
        if row.volume is not None:
            cell = ws.cell(row=target, column=4)
            cell.value = row.volume
            cell.number_format = "0.00"
        if row.justification is not None:
            ws.cell(row=target, column=6).value = row.justification
        for col, value in row.extra.items():
            ws.cell(row=target, column=col).value = value

    os.makedirs(os.path.dirname(os.path.abspath(output_path)) or ".", exist_ok=True)
    out.save(output_path)
    logger.info("ВОР сохранён: %s (строк %d)", output_path, len(output_rows))
    return output_path
