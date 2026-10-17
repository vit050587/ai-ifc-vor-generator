"""Local text-PDF/XLSX VOR parser. Column headings, not fixed column numbers."""
import re
import json
from decimal import Decimal, InvalidOperation
from pathlib import Path

from .works_comparison.units import unit_and_factor


class VorParseError(ValueError):
    pass


class QuantityColumnRequired(VorParseError):
    def __init__(self, choices):
        self.choices = list(dict.fromkeys(choices))
        self.key = json.dumps(self.choices, ensure_ascii=False)
        super().__init__('Найдено несколько столбцов объёма. Выберите нужный: ' + ', '.join(self.choices))


def clean(value):
    return re.sub(r'\s+', ' ', str(value if value is not None else '').replace('\xa0', ' ')).strip()


def heading(value):
    return clean(value).lower().replace('ё','е').replace('.', '').replace('-', '').replace(' ','')


def quantity_kind(value):
    text = heading(value)
    if any(word in text for word in ('обоснован','расчет','расчёт','разниц','отклон','было','стоим','цен','примеч','код','шифр')):
        return None
    if text in ('стало','послекорректировки','попроекту','проектныйобъем','проектныйобьем','актуальныйобъем'):
        return 'current'
    if any(word in text for word in ('объем','обьем','количество','колво','колич')):
        return 'quantity'
    return None


def choose_quantity(candidates, selection=None):
    labels = [clean(c[1]) for c in candidates]
    candidates = [(i, f'{label} [колонка {i+1}]' if labels.count(clean(label))>1 else label, kind) for i,label,kind in candidates]
    current = [c for c in candidates if c[2] == 'current']
    eligible = current or candidates
    signature = json.dumps(list(dict.fromkeys(c[1] for c in eligible)), ensure_ascii=False)
    selected = selection.get(signature) if isinstance(selection, dict) else selection
    if selected:
        found = [c for c in eligible if clean(c[1]) == clean(selected)]
        if len(found) == 1:
            return found[0]
    if len(eligible) == 1:
        return eligible[0]
    raise QuantityColumnRequired([c[1] for c in eligible])


def number(value):
    text = re.sub(r'\s+', '', str(value if value is not None else '')).replace(',', '.')
    if not re.fullmatch(r'[+-]?\d+(?:\.\d+)?', text):
        return None
    try:
        value = Decimal(text)
        return float(value) if value.is_finite() else None
    except InvalidOperation:
        return None


def unit(value):
    text = clean(value)
    # Line-wrapped punctuation is common in text PDFs: "куб.м\n.".
    canonical, factor = unit_and_factor(text)
    return text if factor is not None else None


def row_record(name, units, quantity, location, number_value='', section='', code=None, raw=None):
    return {'workAndCostName':clean(name), 'unit':clean(units), 'quantity':quantity,
            'rateCodeAndResourceCodes':{'code':clean(code) or None}, 'rowNumber':location.get('row'),
            'sourceLocation':location,'sourceQuantityColumn':location.get('quantityColumn'),
            'sourceRawQuantity':raw, 'positionNumber':clean(number_value), 'sectionName':section}


def payload(records, filename, diagnostics):
    if not records:
        raise VorParseError('Не удалось извлечь строки ВОР с наименованием и единицей. Для сканированного PDF требуется OCR; для текстового проверьте заголовки таблицы.')
    sections = []
    for record in records:
        section = record['sectionName'] or 'ВОР заказчика'
        if not sections or sections[-1]['sectionName'] != section:
            sections.append({'sectionName':section,'positions':[]})
        sections[-1]['positions'].append({'number':record['positionNumber'] or len(sections[-1]['positions'])+1,'rows':[record]})
    return {'status':'completed','filename':filename,'parser':'local_vor_v3',
            'parseDiagnostics':diagnostics,'resultData':[{'estimateNumber':filename,'workDescription':'ВОР заказчика','form':'ВОР','sections':sections}]}


def parse_xlsx(path, selection=None, progress=None):
    from openpyxl import load_workbook
    workbook = load_workbook(path,read_only=True,data_only=True)
    records, diagnostics = [], []
    try:
        for sheet in workbook:
            schema = None
            section = sheet.title
            for row_index, values in enumerate(sheet.iter_rows(values_only=True),1):
                cells = [clean(v) for v in values]
                name_columns = [i for i,c in enumerate(cells) if any(k in heading(c) for k in ('наименование','наименовани'))]
                units_columns = [i for i,c in enumerate(cells) if heading(c) in ('едизм','единицаизмерения','единицыизмерения')]
                if name_columns and units_columns:
                    candidates=[(i,c,quantity_kind(c)) for i,c in enumerate(cells) if quantity_kind(c)]
                    # For multi-row headers collect a lower header on the next iteration.
                    schema={'name':name_columns[0],'unit':units_columns[0],'candidates':candidates,'row':row_index,
                            'code':next((i for i,c in enumerate(cells) if heading(c) in ('шифртсн','шифр','код','кодрасценки')),None)}
                    continue
                if not schema:
                    continue
                lower=[(i,c,quantity_kind(c)) for i,c in enumerate(cells) if quantity_kind(c)]
                if lower and row_index <= schema['row']+2 and not unit(cells[schema['unit']]):
                    schema['candidates']=lower
                    continue
                units=unit(cells[schema['unit']])
                name=cells[schema['name']]
                if not units:
                    texts=[c for c in cells[:schema['unit']] if c]
                    if texts and not any(number(c) is not None for c in cells[schema['unit']:]):
                        section=' / '.join(texts)
                    continue
                if not name or not schema['candidates']:
                    continue
                selected=choose_quantity(schema['candidates'],selection)
                raw=values[selected[0]]
                value=number(raw)
                records.append(row_record(name,units,value,{'sheet':sheet.title,'row':row_index,'quantityColumn':selected[1]},
                                          cells[0],section,cells[schema['code']] if schema['code'] is not None else None,raw))
                if value is None:
                    diagnostics.append(f'{sheet.title}, строка {row_index}: объём не распознан: {raw}')
            if progress:
                progress(f'Лист «{sheet.title}»: извлечено {len(records)} строк')
    finally:
        workbook.close()
    return payload(records,Path(path).name,diagnostics)


def _pdf_headers(words):
    headers=[]
    for name in words:
        if not heading(name[4]).startswith('наименование'):
            continue
        nearby=[w for w in words if name[1]-3 <= w[1] <= name[1]+65]
        units=[w for w in nearby if heading(w[4]) in ('ед','едизм','единица') and w[0]>name[0]]
        if not units:
            continue
        u=units[0];ux=(u[0]+u[2])/2
        current=[w for w in nearby if quantity_kind(w[4])=='current' and w[0]>u[2]]
        candidates=current or [w for w in nearby if quantity_kind(w[4]) and w[0]>u[2] and heading(w[4]) not in ('объема','объемаработ')]
        if not candidates:
            continue
        # Revision subheaders have a parent Кол-во; use the child header row.
        if current:
            candidates=[w for w in nearby if abs(w[1]-current[0][1])<5 and (quantity_kind(w[4]) or heading(w[4]) in ('было','разница'))]
        quantity_candidates=[(i,clean(w[4]),quantity_kind(w[4])) for i,w in enumerate(candidates) if quantity_kind(w[4])]
        all_centers=sorted({(w[0]+w[2])/2 for w in candidates})
        header_bottom=max(w[3] for w in candidates)
        headers.append({'top':name[1]-2,'bottom':header_bottom+2,'unit_x':ux,'name_left':name[0]-15,
                        'unit_left':u[0]-10,'columns':candidates,'centers':all_centers,'candidates':quantity_candidates})
    return sorted(headers,key=lambda h:h['top'])


def _section_lines(words, left, right):
    lines = {}
    for w in words:
        if left <= w[0] and w[2] < right:
            lines.setdefault(round(w[1]/3), []).append(w)
    result=[]
    for line in lines.values():
        text=_join_words(line)
        if re.match(r'^(?:\d+\.\s+[А-Я]|Раздел|Подраздел|Деформационный|Секции|Секция\s)',text):
            y=min(w[1] for w in line)
            # Section titles span the whole table, unlike work-name cells.
            full_line=[w for w in words if w[0]>=left and abs(w[1]-y)<3]
            if any(right<=w[0]<=right+55 and unit(w[4]) for w in full_line):
                continue
            result.append((y,_join_words(full_line)))
    return sorted(result)


def _join_words(words):
    return clean(' '.join(w[4] for w in sorted(words,key=lambda w:(round(w[1]/3),w[0]))))


def _pdf_grid(page):
    vertical, horizontal = [], []
    for drawing in page.get_drawings():
        for item in drawing.get('items', []):
            if item[0]=='re':
                rect=item[1]
                if rect.width<=1 and rect.height>5:
                    vertical.append(((rect.x0+rect.x1)/2,rect.y0,rect.y1))
                elif rect.height<=1 and rect.width>30:
                    horizontal.append(((rect.y0+rect.y1)/2,rect.x0,rect.x1))
            elif item[0]=='l':
                a,b=item[1:3]
                if abs(a.x-b.x)<1 and abs(a.y-b.y)>5:
                    vertical.append(((a.x+b.x)/2,min(a.y,b.y),max(a.y,b.y)))
                elif abs(a.y-b.y)<1 and abs(a.x-b.x)>30:
                    horizontal.append(((a.y+b.y)/2,min(a.x,b.x),max(a.x,b.x)))
    return vertical,horizontal


def _grid_schema(schema, grid, y):
    h=dict(schema)
    vertical,_=grid
    xs=sorted({round(x,1) for x,top,bottom in vertical if top-1<=y<=bottom+1})
    left=[x for x in xs if x<h['unit_x']-4]
    right=[x for x in xs if x>h['unit_x']+4]
    if len(left)>=2 and right:
        # Actual cell borders, not the centred header text.
        h['name_left']=left[-2]+1
        h['unit_left']=left[-1]
        h['unit_right']=right[0]
    return h


def parse_pdf(path, selection=None, progress=None):
    import fitz
    document=fitz.open(path)
    records, diagnostics=[],[]
    schema=None
    section=section_title='ВОР заказчика'
    try:
        for page_number,page in enumerate(document,1):
            words=page.get_text('words')
            headers=_pdf_headers(words)
            grid=_pdf_grid(page)
            segments=[]
            if schema:
                segments.append((0,headers[0]['top'] if headers else page.rect.height-45,schema))
            for i,header in enumerate(headers):
                segments.append((header['bottom'],headers[i+1]['top'] if i+1<len(headers) else page.rect.height-45,header))
            for top,bottom,h in segments:
                if not h['candidates']:
                    continue
                probe=next((w for w in words if top<=w[1]<bottom and abs((w[0]+w[2])/2-h['unit_x'])<19 and unit(w[4])),None)
                if probe:
                    h=_grid_schema(h,grid,(probe[1]+probe[3])/2)
                selected=choose_quantity(h['candidates'],selection)
                qword=h['columns'][selected[0]];qx=(qword[0]+qword[2])/2
                centers=h['centers'];left=max([x for x in centers if x<qx],default=h['unit_x'])
                right=min([x for x in centers if x>qx],default=qx+(qx-left))
                qleft,qright=(left+qx)/2,(right+qx)/2
                section_lines=_section_lines(words,h['name_left']-40,h['unit_left'])
                unit_words=[w for w in words if top<=w[1]<bottom and abs((w[0]+w[2])/2-h['unit_x'])<19 and unit(w[4])]
                if unit_words and records and top==0:
                    prefix=_join_words([w for w in words if 35<=w[1]<unit_words[0][1]-5 and h['name_left']<=w[0] and w[2]<h['unit_left']+5])
                    previous=records[-1]
                    if (previous['sourceLocation'].get('page')==page_number-1 and 'ГОСТ' in previous['workAndCostName']
                            and re.fullmatch(r'(?:Р\s+)?(?:ГОСТ\s+)?\d{4,5}(?:-\d{2,4})?[.,]?',prefix)):
                        separator='' if previous['workAndCostName'].endswith('-') else ' '
                        previous['workAndCostName']+=separator+prefix
                        previous['sourceLocation'].setdefault('continuedPages',[]).append(page_number)
                for j,uw in enumerate(unit_words):
                    preceding_numbers=[w[1] for w in words if w[0]<h['name_left'] and re.fullmatch(r'\d+(?:\.\d+)*',w[4]) and uw[1]-5<=w[1]<=uw[1]+2]
                    start=(max(preceding_numbers) if preceding_numbers else uw[1])-2
                    end=unit_words[j+1][1]-2 if j+1<len(unit_words) else bottom
                    # A numbered row without a unit must not be swallowed by this row.
                    following_numbers=[w[1] for w in words if w[0]<h['name_left']-2 and w[0]>h['name_left']-45 and re.fullmatch(r'\d+(?:\.\d+)*',w[4]) and uw[1]+5<w[1]<end]
                    if following_numbers:
                        end=min(following_numbers)-2
                    row_y=(uw[1]+uw[3])/2
                    borders=[y for y,x0,x1 in grid[1] if x0<=h['name_left']+15<=x1 and top-2<=y<bottom+2]
                    lower=[y for y in borders if y<row_y]
                    upper=[y for y in borders if y>row_y]
                    if lower and upper:
                        start=max(start,max(lower))
                        end=min(end,min(upper))
                    preceding=[label for y,label in section_lines if y < start]
                    if preceding:
                        for label in preceding:
                            if not re.match(r'^(?:Деформационный|Секции|Секция\s)', label):
                                section_title=label
                        last=preceding[-1]
                        section=section_title+' > '+last if re.match(r'^(?:Деформационный|Секции|Секция\s)', last) else last
                    next_section=[y for y,label in section_lines if start+10 < y < end]
                    if next_section:
                        end=min(next_section)
                    row_words=[w for w in words if start<=(w[1]+w[3])/2<end]
                    # Stop the last row at the next section/header rather than absorbing it.
                    headings=[w[1] for w in row_words if w[0]<h['name_left'] and re.fullmatch(r'\d+\.',w[4]) and w[1]>uw[1]+15]
                    if headings:
                        end=min(headings);row_words=[w for w in row_words if w[1]<end]
                    names=[w for w in row_words if h['name_left']<=w[0] and w[2]<h['unit_left']+5]
                    name=_join_words(names)
                    if not name:
                        continue
                    qwords=[w for w in row_words if qleft<=(w[0]+w[2])/2<qright and re.fullmatch(r'[+−–-]?\d*(?:[.,]\d*)?', w[4])]
                    raw=''.join(w[4] for w in sorted(qwords,key=lambda w:(w[1],w[0])))
                    value=number(raw)
                    row_numbers=[w[4] for w in row_words if w[0]<h['name_left'] and abs(w[1]-uw[1])<5 and re.fullmatch(r'\d+(?:\.\d+)*',w[4])]
                    if not row_numbers and not (lower and upper) and not headers:
                        # A carried schema is not evidence that a prose/formula page is a table.
                        continue
                    unit_cell=_join_words([w for w in row_words if abs((w[0]+w[2])/2-h['unit_x'])<22])
                    records.append(row_record(name,unit_cell if unit(unit_cell) else uw[4],value,{'page':page_number,'row':len(records)+1,'quantityColumn':selected[1]},
                                              row_numbers[0] if row_numbers else '',section,raw=raw))
                    if value is None:
                        diagnostics.append(f'Страница {page_number}, {name[:60]}: объём не распознан: {raw[:60]}')
            if headers:
                schema=headers[-1]
            if progress:
                progress(f'PDF: страница {page_number}/{len(document)}, извлечено строк: {len(records)}')
    finally:
        document.close()
    return payload(records,Path(path).name,diagnostics)


def parse_vor(path, quantity_column=None, progress=None):
    ext=Path(path).suffix.lower()
    if ext=='.pdf':
        return parse_pdf(path,quantity_column,progress)
    if ext=='.xlsx':
        return parse_xlsx(path,quantity_column,progress)
    raise VorParseError('Локальный парсер ВОР поддерживает PDF с текстом и XLSX')
