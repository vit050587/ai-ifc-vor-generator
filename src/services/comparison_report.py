"""VOR comparison with explicit matching links and traceable source positions."""
import os
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from .comparison_vor import prepare_report


def build_comparison_excel(comparison_dir, result):
    report = prepare_report(result)
    workbook = Workbook()
    ws = workbook.active
    ws.title = 'Сравнение ВОР'
    headers = ['Наименование работы из сформированного ВОР','Объём из сформированного ВОР','Наименование работы из ВОР заказчика','Объём из ВОР заказчика','Различие','Обоснование модели','Степень уверенности','Комментарии']
    ws.append(['Сравнение сформированного ВОР с ВОР заказчика'])
    ws.merge_cells(start_row=1,start_column=1,end_row=1,end_column=len(headers))
    ws.append(['Расхождение = ВОР заказчика − сформированный ВОР. Процент — от объёма сформированного ВОР. Плюс: заказчик завысил, минус: занизил. Единицы без множителей. Уверенность относится к сопоставлению работ, не к точности объёма. Пустая ячейка означает отсутствие установленной пары; нулевой объём отображается числом 0.'])
    similarity=report.get('compositionSimilarity') or {}
    score=similarity.get('score')
    if score is not None:
        ws.cell(2,1).value='Сходство состава работ: '+str(score)+'%. Сопоставлено групп работ сформированного ВОР: '+str(similarity['generatedMatched'])+' из '+str(similarity['generatedTotal'])+'; заказчика: '+str(similarity['customerMatched'])+' из '+str(similarity['customerTotal'])+'. '+ws.cell(2,1).value+' Монтаж/демонтаж опалубки считаются одной группой работ; пустые группы в оценку не входят.'
    ws.merge_cells(start_row=2,start_column=1,end_row=2,end_column=len(headers))
    ws.append(headers)
    previous = None
    for number,row in enumerate(report['vorRows'],1):
        hierarchy = tuple(row[k] for k in ('part','section','subsection','elementGroup'))
        if hierarchy != previous:
            common = 0
            if previous:
                for old,new in zip(previous,hierarchy):
                    if old!=new:break
                    common+=1
            for depth,label in enumerate(hierarchy):
                if depth < common or not label:continue
                ws.append([label])
                ws.merge_cells(start_row=ws.max_row,start_column=1,end_row=ws.max_row,end_column=len(headers))
                ws.cell(ws.max_row,1).fill = PatternFill('solid',fgColor=['244A73','D8E5F2','E8F0F7','F1F5F9'][depth])
                ws.cell(ws.max_row,1).font = Font(bold=True,color='FFFFFF' if depth==0 else '244A73',size=11)
                ws.cell(ws.max_row,1).alignment = Alignment(indent=depth,wrap_text=True,vertical='center')
                ws.row_dimensions[ws.max_row].height = 34
            previous=hierarchy
        ws.append([row['generatedWork'],row['generatedQuantity'],row['customerWork'],row['customerQuantity'],row['difference'],row['modelReason'],row['comparisonConfidence'],row['comments']])
        unit_text=str(row['unit'] or '').replace('"','')
        for column in (2,4,5):
            ws.cell(ws.max_row,column).number_format='#,##0.###" '+unit_text+'"'
        ws.row_dimensions[ws.max_row].height=85
        fill='E7F4EA' if row['status'] in ('matched','operation_linked') else 'FFF3D6' if row['status'] in ('quantity_mismatch','only_ifc','only_project') else 'FDE9E7'
        ws.cell(ws.max_row,8).fill=PatternFill('solid',fgColor=fill)
    matching=workbook.create_sheet('Сопоставление работ')
    matching.append(['Сопоставление работ сформированного ВОР и ВОР заказчика'])
    matching.merge_cells('A1:N1')
    matching.append(['Номер объединяет строки одной связи. Повторённые строки обоих ВОР в пределах одного номера учитываются один раз. Объёмы операций комплекса не суммируются. Предложенные связи требуют проверки; уверенность сопоставления не является точностью объёма.'])
    matching.merge_cells('A2:N2')
    matching.append(['№ связи','Работа сформированного ВОР','Объект / группа','Источник модели','Ед. изм. сформированного ВОР','Объём операции','Работа ВОР заказчика','Объект / раздел заказчика','Источник заказчика','Ед. изм. заказчика','Объём заказчика','Тип связи','Уверенность сопоставления','Статус / объяснение'])
    relation_labels={'equivalent':'Одна операция','aggregate':'Одна операция на нескольких участках',
                     'complex':'Комплекс операций — не суммировать','unresolved':'Не установлена'}
    def source_label(position):
        label=position.get('sourceFile') or '; '.join(position.get('sourceFiles') or [])
        loc=position.get('sourceLocation') or {}
        places=[f"{title}: {loc[key]}" for key,title in [('page','стр.'),('sheet','лист'),('row','строка')] if loc.get(key) is not None]
        return '; '.join(filter(None,[label,', '.join(places)]))
    def object_label(position):
        element=position.get('element') or {};props=position.get('properties') or {}
        return ', '.join(str(v) for v in [element.get('category'),element.get('typeName'),element.get('level'),props.get('location'),props.get('material'),position.get('groupId')] if v)
    for number,row in enumerate(report['vorRows'],1):
        if row.get('isEmptyGroup'):continue
        group=row['detail'];semantic=group.get('semanticMatch') or {}
        customer=(group.get('projectGroup') or {}).get('positions',[])
        generated=(group.get('ifcGroup') or {}).get('positions',[])
        relation=relation_labels.get(semantic.get('relation'),'По шифру / наименованию' if customer and generated else 'Не установлена')
        explanation=('Подтверждённая связь' if semantic.get('confirmed') and row['status']!='matching_unresolved' else 'Предложенная связь — требуется проверка') if semantic else row['result']
        explanation='; '.join(filter(None,[explanation,semantic.get('reason')]))
        # A repeated customer volume describes the same link; it is not a new quantity.
        for cp in customer or [{}]:
            for gp in generated or [{}]:
                matching.append([number,gp.get('name') or gp.get('workAndCostName'),object_label(gp),source_label(gp),gp.get('unit'),gp.get('quantity'),
                                 cp.get('workAndCostName') or cp.get('name') or (row['customerWork'] if row.get('operationLink') else None),'; '.join(filter(None,[(cp.get('_construction_group') or {}).get('label'),cp.get('sectionName'),(cp.get('_construction') or {}).get('name')])),source_label(cp),cp.get('unit'),cp.get('quantity'),
                                 relation,('По правилам' if semantic.get('method')=='rules' else semantic.get('confidence')),explanation])
    detail=workbook.create_sheet('Подробности')
    detail.append(['Код','Сторона','Работа','Группа','Источник','Ед. без множителя','Объём без множителя','Исходная единица','Исходный объём','Коэффициент','Вердикт ИИ','Уверенность ИИ','Обоснование','Недостающие данные','Предупреждения'])
    detail_rows=report['vorRows']
    if report.get('csFinalModels'):
        from .works_comparison.ifc_positions_extractor import IfcPositionsExtractor
        from .works_comparison.units import normalize_positions
        source_positions=normalize_positions(IfcPositionsExtractor(report['csFinalModels']).run())
        detail_rows=[{'code':None,'detail':{'ifcGroup':{'positions':source_positions}}}]+[{'code':None,'detail':{'projectGroup':entry.get('projectGroup')}} for entry in report.get('comparison',[]) if entry.get('projectGroup')]
    for row in detail_rows:
        for key,label in [('ifcGroup','Сформированный ВОР'),('projectGroup','ВОР заказчика')]:
            for p in (row['detail'].get(key) or {}).get('positions',[]):
                v=p.get('validation_result') or {}
                detail.append([p.get('code') or row['code'],label,p.get('name') or p.get('workAndCostName'),p.get('groupId'),p.get('sourceFile') or '; '.join(p.get('sourceFiles') or []),p.get('unit'),p.get('quantity'),p.get('originalUnit'),p.get('originalQuantity'),p.get('unitFactor'),v.get('verdict'),v.get('confidence'),v.get('reason'),'; '.join(v.get('missingData') or []),'; '.join(v.get('warnings') or [])])
    for sheet,header_row in [(ws,3),(matching,3),(detail,1)]:
        sheet.sheet_view.showGridLines=False
        sheet.freeze_panes=f'D{header_row+1}'
        if sheet!=ws:sheet.auto_filter.ref=f'A{header_row}:{get_column_letter(sheet.max_column)}{sheet.max_row}'
        for cell in sheet[header_row]:
            cell.fill=PatternFill('solid',fgColor='244A73');cell.font=Font(bold=True,color='FFFFFF');cell.alignment=Alignment(wrap_text=True,vertical='center')
        for cells in sheet:
            for cell in cells:
                if isinstance(cell.value,str):cell.data_type='s'
                if cell.row>header_row and cell.column>1:
                    cell.alignment=Alignment(wrap_text=True,vertical='top')
                if isinstance(cell.value,(float,int)) and not (sheet==ws and cell.column in (2,4,5)):cell.number_format='#,##0.###'
        sheet.row_dimensions[header_row].height=38
        for i in range(1,sheet.max_column+1):sheet.column_dimensions[get_column_letter(i)].width=24
    widths=[54,22,54,22,22,65,30,65]
    for i,width in enumerate(widths,1):ws.column_dimensions[get_column_letter(i)].width=width
    ws.row_dimensions[1].height=28;ws.row_dimensions[2].height=64
    ws.print_title_rows='1:3';ws.sheet_properties.pageSetUpPr.fitToPage=True
    ws.page_setup.orientation='landscape';ws.page_setup.paperSize=ws.PAPERSIZE_A3;ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0
    for i in (3,13,14,15):detail.column_dimensions[get_column_letter(i)].width=55
    for i,width in enumerate([10,50,35,35,14,18,50,42,35,14,18,35,20,65],1):
        matching.column_dimensions[get_column_letter(i)].width=width
    for cells in matching.iter_rows(min_row=4):
        cells[12].number_format='0%'
        matching.row_dimensions[cells[0].row].height=85
    matching.row_dimensions[1].height=28;matching.row_dimensions[2].height=42
    matching.print_title_rows='1:3'
    target=Path(comparison_dir)/'result.xlsx';temporary=target.with_name('result.tmp.xlsx')
    workbook.save(temporary);os.replace(temporary,target)
    return str(target)
