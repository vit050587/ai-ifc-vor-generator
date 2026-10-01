"""
Локальный справочник расценок ТСН — замена цифрового сборника (larix).

Источник — data/price_cost.xlsx (выгрузка цифрового сборника):
  - лист «_Получение_расценок_из_дерева_s»: шифр, наименование, ед. изм., статус, id;
  - лист «_Получние_параметры_позиции_sel»: стоимость (ЗП, ЭМ, МР, прямые затраты).

При первом обращении Excel конвертируется в data/local_works.json
(быстрый кеш), дальше читается только JSON.

Отладка:
    python -m src.services.local_works           # пересобрать кеш, показать статистику
    python -m src.services.local_works 3.6-71    # показать работы таблицы
"""

import json
import os
import re
import sys
from datetime import datetime, timezone
from typing import Any, Dict, List, Optional

from src.core.logger import setup_logger

logger = setup_logger(__name__)

PRICE_COST_XLSX = os.getenv("PRICE_COST_PATH", "/app/data/price_cost.xlsx")
LOCAL_WORKS_JSON = os.getenv("LOCAL_WORKS_PATH", "/app/data/local_works.json")

SHEET_RATES = "_Получение_расценок_из_дерева_s"
SHEET_PARAMS = "_Получние_параметры_позиции_sel"

# Период-заглушка: в локальном режиме период ТСН — это выгрузка price_cost.xlsx
LOCAL_PERIOD: Dict[str, Any] = {
    "title": "Локальный справочник (price_cost.xlsx)",
    "id": 1,
    "dateStart": None,
    "baseTypeCode": "TSN",
}

# Начало имени колонки (нижний регистр) -> поле в формате ответа larix.
# В именах колонок выгрузки есть опечатки («стомость»), поэтому даны варианты.
COST_COLUMNS = [
    (("базисная прямые затраты",), "directCosts"),
    (("текущие прямые затраты",), "curDirectCosts"),
    (("базисная зарплата рабочих",), "salary"),
    (("текущая зарплата рабочих",), "curSalary"),
    (("базисная стоимость эксплуатации",), "operationOfMachines"),
    (("текущая стоимость эксплуатации",), "curOperationOfMachines"),
    (("базисная стоимость материалов",), "costOfMaterialResources"),
    (("текущая стомость материалов", "текущая стоимость материалов"), "curCostOfMaterialResources"),
    (("базисная стоимость всего",), "totalCost"),
    (("текущее стоимость всего", "текущая стоимость всего"), "curTotalCost"),
]

_cache: Optional[Dict[str, Any]] = None


def _to_float(value: Any) -> Optional[float]:
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def _to_int(value: Any) -> Optional[int]:
    try:
        return int(float(value))
    except (TypeError, ValueError):
        return None


def table_code_of(pressmark: str) -> str:
    """Шифр расценки -> шифр таблицы: «3.6-71-1» -> «3.6-71»."""
    pm = str(pressmark).strip()
    return pm.rsplit("-", 1)[0] if pm.count("-") >= 2 else pm


def _pm_sort_key(pm: str):
    return [(0, int(p), "") if p.isdigit() else (1, 0, p) for p in re.split(r"[-.]", pm)]


def build_local_works(xlsx_path: str = PRICE_COST_XLSX,
                      out_path: str = LOCAL_WORKS_JSON) -> Dict[str, Any]:
    """Конвертирует price_cost.xlsx в JSON-справочник {шифр таблицы: [работы]}."""
    import openpyxl

    logger.info(f"Сборка локального справочника расценок из {xlsx_path} ...")
    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)

    # --- стоимость по шифру расценки ---
    rows = wb[SHEET_PARAMS].iter_rows(values_only=True)
    header = [str(h or "").strip().lower() for h in next(rows)]
    col_map: Dict[int, str] = {}
    for i, h in enumerate(header):
        for prefixes, field in COST_COLUMNS:
            if field not in col_map.values() and any(h.startswith(p) for p in prefixes):
                col_map[i] = field
                break
    missing = {f for _, f in COST_COLUMNS} - set(col_map.values())
    if missing:
        logger.warning(f"В листе параметров не найдены колонки стоимости: {sorted(missing)}")

    costs: Dict[str, Dict[str, Optional[float]]] = {}
    for r in rows:
        pm = str(r[0] or "").strip()
        if pm and pm not in costs:
            costs[pm] = {field: _to_float(r[i]) for i, field in col_map.items()}

    # --- сами расценки ---
    rows = wb[SHEET_RATES].iter_rows(values_only=True)
    idx = {str(h or "").strip(): i for i, h in enumerate(next(rows))}

    def col(r, name):
        i = idx.get(name)
        return r[i] if i is not None and i < len(r) else None

    tables: Dict[str, List[Dict[str, Any]]] = {}
    seen = set()
    for r in rows:
        pm = str(col(r, "Шифр расценки") or "").strip()
        if not pm or pm in seen:
            continue
        seen.add(pm)
        work: Dict[str, Any] = {
            "id": _to_int(col(r, "id")),
            "pressmark": pm,
            "title": str(col(r, "Наименование") or "").strip(),
            "unitOfMeasure": str(col(r, "Единица измерения") or "").strip(),
            "status": str(col(r, "Статус") or "").strip(),
        }
        work.update(costs.get(pm, {}))
        tables.setdefault(table_code_of(pm), []).append(work)
    wb.close()

    for works in tables.values():
        works.sort(key=lambda w: _pm_sort_key(w["pressmark"]))

    payload = {
        "source": os.path.basename(xlsx_path),
        "generated_at": datetime.now(timezone.utc).isoformat().replace("+00:00", "Z"),
        "total_tables": len(tables),
        "total_works": len(seen),
        "with_costs": sum(1 for pm in seen if pm in costs),
        "tables": tables,
    }
    with open(out_path, "w", encoding="utf-8") as fh:
        json.dump(payload, fh, ensure_ascii=False)
    logger.info(
        f"Локальный справочник сохранён: {out_path} "
        f"({payload['total_tables']} таблиц, {payload['total_works']} расценок, "
        f"со стоимостью {payload['with_costs']})"
    )
    return payload


def load_local_works() -> Dict[str, Any]:
    """Справочник из кеша (при отсутствии кеша — собирается из Excel)."""
    global _cache
    if _cache is None:
        if not os.path.isfile(LOCAL_WORKS_JSON):
            build_local_works()
        with open(LOCAL_WORKS_JSON, "r", encoding="utf-8") as fh:
            data = json.load(fh)
        data["_by_id"] = {
            w["id"]: w
            for works in data["tables"].values()
            for w in works
            if w.get("id") is not None
        }
        _cache = data
    return _cache


def get_table_works(table_code: str) -> List[Dict[str, Any]]:
    """Все расценки таблицы (аналог catalog/work-process/list)."""
    works = load_local_works()["tables"].get(str(table_code).strip(), [])
    if not works:
        logger.warning(f"Таблица {table_code}: в локальном справочнике расценок нет")
    return [dict(w) for w in works]


def get_work_details(work_ids: List[int]) -> Dict[int, Dict[str, Any]]:
    """Стоимость расценок по id (аналог catalog/work-process/detail)."""
    by_id = load_local_works()["_by_id"]
    details = {int(wid): dict(by_id[int(wid)]) for wid in work_ids if int(wid) in by_id}
    logger.info(f"Локальная стоимость позиций: найдено {len(details)} из {len(work_ids)}")
    return details


if __name__ == "__main__":
    if len(sys.argv) > 1:
        for w in get_table_works(sys.argv[1]):
            print(w["pressmark"], "|", w["unitOfMeasure"], "|", w["title"][:80],
                  "| ЗП", w.get("curSalary"), "ЭМ", w.get("curOperationOfMachines"),
                  "МР", w.get("curCostOfMaterialResources"))
    else:
        info = build_local_works()
        print(f"Таблиц: {info['total_tables']}, расценок: {info['total_works']}, "
              f"со стоимостью: {info['with_costs']}")
