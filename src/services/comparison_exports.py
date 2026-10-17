"""Audit exports of comparison inputs, before any LLM matching or validation."""
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from .works_comparison.positions_extractor import PositionsExtractor
from .works_comparison.ifc_positions_extractor import IfcPositionsExtractor
from .works_comparison.units import normalize_positions
from .works_comparison.construction_matching import attach_constructions, group_constructions
from .works_comparison.semantic_matching import _customer_groups
from .works_comparison.utils import sum_positions_quantity

EXPORT_FILES = ('customer_vor_parsed.xlsx','customer_vor_grouped.xlsx','generated_vor.xlsx')


def _source(p):
    loc=p.get('sourceLocation') or {}
    place=', '.join(f'{label}: {loc[key]}' for key,label in [('page','стр.'),('sheet','лист'),('row','строка')] if loc.get(key) is not None)
    return p.get('sourceFile') or '; '.join(p.get('sourceFiles') or []),place


def _finish(workbook, target):
    for sheet in workbook:
        sheet.freeze_panes='C2'
        sheet.auto_filter.ref=sheet.dimensions
        sheet.sheet_view.showGridLines=False
        sheet.row_dimensions[1].height=42
        for cell in sheet[1]:
            cell.fill=PatternFill('solid',fgColor='244A73');cell.font=Font(bold=True,color='FFFFFF')
        for row in sheet:
            for cell in row:
                cell.alignment=Alignment(wrap_text=True,vertical='top')
                if isinstance(cell.value,str):cell.data_type='s'
                elif isinstance(cell.value,(int,float)):cell.number_format='#,##0.###'
        for index in range(1,sheet.max_column+1):
            title=str(sheet.cell(1,index).value or '')
            sheet.column_dimensions[get_column_letter(index)].width=52 if any(word in title for word in ('Наименование','Раздел','Конструкция','Секции')) else 22
    target.parent.mkdir(parents=True,exist_ok=True)
    tmp=target.with_suffix('.tmp.xlsx');workbook.save(tmp);tmp.replace(target)


def build_comparison_exports(folder, models=None, projects=None):
    folder=Path(folder)
    generated=normalize_positions(IfcPositionsExtractor(models).run()) if models is not None else []
    if models is not None:
        workbook=Workbook();sheet=workbook.active;sheet.title='Сформированный ВОР'
        sheet.append(['Модель','Группа','Тип элемента','Расположение / этажи','Материал','Код работы','Наименование работы','Ед. изм.','Объём','Статус объёма','Исходная единица','Исходный объём'])
        by_group={}
        for p in generated:by_group.setdefault((p.get('sourceFile'),p.get('groupId')),[]).append(p)
        for model in models:
            for group in model.get('elementGroups',[]):
                members=by_group.get((model.get('sourceFile'),group.get('groupId')),[])
                if not members:
                    el,props=group.get('element') or {},group.get('properties') or {}
                    sheet.append([model.get('sourceFile'),group.get('groupId'),el.get('typeName'),el.get('level'),props.get('material'),None,None,None,None,'Работы не подобраны',None,None])
                for p in members:
                    el,props=p.get('element') or {},p.get('properties') or {}
                    sheet.append([p.get('sourceFile'),p.get('groupId'),el.get('typeName'),el.get('level'),props.get('material'),p.get('code'),p.get('name'),p['unit'],p['quantity'],p.get('quantityStatus'),p.get('originalUnit'),p.get('originalQuantity')])
        _finish(workbook,folder/EXPORT_FILES[2])
    if projects is None:
        return
    raw=[]
    for source in projects:
        for document in source.get('result',{}).get('resultData',[]):
            for section in document.get('sections',[]):
                for position in section.get('positions',[]):
                    for row in position.get('rows',[]):
                        raw.append({**row,'sourceFiles':[source.get('sourceFile')],'sectionName':section.get('sectionName'),'positionNumber':position.get('number')})
    comparison_rows=PositionsExtractor(projects).run()
    workbook=Workbook();sheet=workbook.active;sheet.title='До объединения'
    sheet.append(['Источник','Раздел / секция','№ позиции','Наименование работы','Исходная единица','Распарсенный объём','Исходный текст объёма','Столбец объёма','Страница / лист / строка'])
    for p in raw:
        source,location=_source(p)
        sheet.append([source,p.get('sectionName'),p.get('positionNumber'),p.get('workAndCostName'),p.get('unit'),p.get('quantity'),p.get('sourceRawQuantity'),p.get('sourceQuantityColumn'),location])
    _finish(workbook,folder/EXPORT_FILES[0])
    grouped=_customer_groups(group_constructions(attach_constructions(normalize_positions(comparison_rows))),generated)
    workbook=Workbook();sheet=workbook.active;sheet.title='После объединения'
    sheet.append(['№ группы работ','Раздел','Конструкция','Наименование работы','Ед. изм.','Объём','Учтено исходных строк','Строк без объёма','Секции','Источник'])
    details=workbook.create_sheet('Исходные строки')
    details.append(['№ группы работ','Раздел / секция','Конструкция','№ позиции','Наименование работы','Ед. изм.','Объём','Роль строки','Источник','Страница / лист / строка'])
    for index,(p,members) in enumerate(grouped,1):
        construction=p.get('_construction_group') or {}
        quantities=sum_positions_quantity(members)
        sheet.append([index,construction.get('sectionName') or p.get('sectionName'),construction.get('label') or (p.get('_construction') or {}).get('name'),
                      '; '.join(dict.fromkeys(x.get('workAndCostName') or '' for x in members)),p['unit'],quantities['totalQuantity'],len(members),quantities['positionsWithoutQuantity'],
                      '; '.join(dict.fromkeys(x.get('sectionName') or '' for x in members)),_source(p)[0]])
        for member in members:
            for detail,role in [(member,'Учтена в итоге')]+[(x,'Детализация; повторно не учитывается') for x in member.get('sourceDetailRows',[])]+[(x,'Заголовок; в объём не входит') for x in member.get('sourceContextRows',[])]:
                source,location=_source(detail)
                details.append([index,detail.get('sectionName'),(detail.get('_construction') or {}).get('name'),detail.get('positionNumber'),detail.get('workAndCostName'),detail.get('unit'),detail.get('quantity'),role,source,location])
    _finish(workbook,folder/EXPORT_FILES[1])
