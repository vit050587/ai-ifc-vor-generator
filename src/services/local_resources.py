"""
Локальный справочник ресурсов (материалов) с ценами — из data/price_cost.xlsx,
лист «_Связаанные_ресурсы_select_wp_p». При первом обращении собирается
кеш data/local_resources.json.

    python -m src.services.local_resources            # пересобрать кеш
    python -m src.services.local_resources В7,5       # поиск ресурса по словам
"""
import json
import os
import re
import sys
from typing import Any, Dict, List, Optional

from src.core.logger import setup_logger

logger = setup_logger(__name__)

PRICE_COST_XLSX = os.getenv("PRICE_COST_PATH", "/app/data/price_cost.xlsx")
LOCAL_RES_JSON = os.getenv("LOCAL_RESOURCES_PATH", "/app/data/local_resources.json")
SHEET = "_Связаанные_ресурсы_select_wp_p"
_CODE = re.compile(r"^\d+\.\d+-\d+-\d+")
_cache: Optional[Dict[str, Dict[str, Any]]] = None


def _f(v: Any) -> Optional[float]:
    try:
        x = float(v)
        return x if x > 0 else None
    except (TypeError, ValueError):
        return None


def build(xlsx: str = PRICE_COST_XLSX, out: str = LOCAL_RES_JSON) -> Dict[str, Dict[str, Any]]:
    import openpyxl
    logger.info(f"Сборка справочника ресурсов из {xlsx} ...")
    wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
    rows = wb[SHEET].iter_rows(values_only=True)
    idx = {str(h or "").strip(): i for i, h in enumerate(next(rows))}

    def col(r, name):
        i = idx.get(name)
        return r[i] if i is not None and i < len(r) else None

    res: Dict[str, Dict[str, Any]] = {}
    for r in rows:
        code = str(col(r, "Шифр ресурса") or "").strip()
        if not _CODE.match(code) or code in res:
            continue
        res[code] = {"pressmark": code,
                     "title": str(col(r, "Наименование ресурса") or "").strip(),
                     "unit": str(col(r, "Единица измерения ресурса") or "").strip(),
                     "type": str(col(r, "Тип ресурса") or "").strip(),
                     "price": _f(col(r, "Сметная цена текущая")),
                     "price_base": _f(col(r, "Сметная цена базисная"))}
    wb.close()
    with open(out, "w", encoding="utf-8") as fh:
        json.dump(res, fh, ensure_ascii=False)
    logger.info(f"Справочник ресурсов: {len(res)} позиций, с ценой {sum(1 for d in res.values() if d['price'])}")
    return res


def load() -> Dict[str, Dict[str, Any]]:
    global _cache
    if _cache is None:
        if not os.path.isfile(LOCAL_RES_JSON):
            build()
        with open(LOCAL_RES_JSON, "r", encoding="utf-8") as fh:
            _cache = json.load(fh)
    return _cache


def resource_price(code: str) -> Optional[float]:
    d = load().get(str(code).strip())
    return d.get("price") if d else None


def concrete_mixes() -> List[Dict[str, Any]]:
    """Бетонные смеси справочника — в формате строки перечня (для подбора по классу)."""
    return [{"pressmark": d["pressmark"], "title": d["title"], "unit": d["unit"] or "м3",
             "formula": "", "is_resource": True}
            for d in load().values() if d["title"].lower().startswith("смесь бетонная тяжелого")]


if __name__ == "__main__":
    if len(sys.argv) > 1:
        words = [w.lower() for w in sys.argv[1:]]
        for d in load().values():
            if all(w in d["title"].lower() for w in words):
                print(d["pressmark"], "|", d["unit"], "|", d["price"], "руб. |", d["title"][:120])
    else:
        info = build()
        print(f"Ресурсов: {len(info)}")



# ===== Нормы расхода ресурсов в расценках (лист «Связанные ресурсы», колонка «Расход») =====
LOCAL_NORMS_JSON = os.getenv("LOCAL_NORMS_PATH", "/app/data/local_norms.json")
_norms = None


def build_norms(xlsx: str = PRICE_COST_XLSX, out: str = LOCAL_NORMS_JSON):
    import openpyxl
    logger.info("Сборка норм расхода ресурсов ...")
    wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
    rows = wb[SHEET].iter_rows(values_only=True)
    idx = {str(h or "").strip(): i for i, h in enumerate(next(rows))}

    def col(r, name):
        i = idx.get(name)
        return r[i] if i is not None and i < len(r) else None

    norms = {}
    for r in rows:
        rate = str(col(r, "Шифр расценки") or "").strip()
        res = str(col(r, "Шифр ресурса") or "").strip()
        v = _f(col(r, "Расход"))
        if rate and _CODE.match(res) and v:
            norms.setdefault(f"{rate}|{res}", v)
    wb.close()
    with open(out, "w", encoding="utf-8") as fh:
        json.dump(norms, fh, ensure_ascii=False)
    logger.info(f"Норм расхода: {len(norms)}")
    return norms


def norm(rate_code: str, res_code: str):
    """Расход ресурса на единицу измерения расценки."""
    global _norms
    if _norms is None:
        if not os.path.isfile(LOCAL_NORMS_JSON):
            build_norms()
        with open(LOCAL_NORMS_JSON, "r", encoding="utf-8") as fh:
            _norms = json.load(fh)
    return _norms.get(f"{rate_code}|{res_code}")
