"""
Эталон подбора работ и автоматическая проверка всех типов элементов модели.

    python -m src.services.kb_etalon init          # создать data/etalon.xlsx (если его нет)
    python -m src.services.kb_etalon init force    # пересоздать data/etalon.xlsx из встроенных правил
    python -m src.services.kb_etalon               # проверить все типы элементов по эталону
    python -m src.services.kb_etalon nocache       # то же, с пересчётом подбора

Эталон — правила по видам конструкций (не по отдельным элементам). Для каждого
правила: как найти элемент, какие расценки обязательны, какие недопустимы,
какой класс бетона ожидается, уверенность и источник. Коды: точный шифр
(«3.6-94-4») или таблица целиком (префикс с дефисом на конце: «3.6-97-»).
"""
import os
import re
import sys
import time
from typing import Any, Dict, List, Optional

from src.services import kb_model as km
from src.services import perechen_kb as kb
from src.services import perechen_pipeline as pp

ETALON_XLSX = os.getenv("ETALON_PATH", "/app/data/etalon.xlsx")
COLUMNS = ["Правило", "Имя (regex)", "Материал (regex)", "IFC класс (regex)", "Часть здания (regex)",
           "Должно быть (все)", "Должно быть (хотя бы одно)", "Не должно быть", "Класс бетона",
           "Уверенность", "Источник", "Комментарий"]

P, T, M = "перечень сметчиков", "правила ТСН", "предположение — уточнить у сметчика"
HEIGHTS_OTHER = "3.6-78-;3.6-82-;3.6-86-;3.6-90-"

DEFAULT_RULES = [
    ("Бетонная подготовка", r"подготов", "", "", "",
     "3.6-1-1", "", "3.6-70-;3.6-71-;3.6-72-;3.6-73-;3.6-77-;3.6-97-;3.6-98-", "7.5",
     "работы — высокая, бетон — средняя", f"{P}; {M}",
     "3.6-1-1 уже включает укладку бетона; В7,5 — типовой класс подготовки"),
    ("Приямок", r"приямок", "", "", "",
     "3.6-74-1;3.6-75-1;3.6-76-3;3.6-98-1", "3.6-77-", "3.6-94-;3.6-97-", "из имени",
     "средняя", P, "опалубка — две стороны стенок, толщина 200 мм (оценка)"),
    ("Фундаментная плита", r"фундаментная плита", "", "", "",
     "3.6-70-3;3.6-71-3;3.6-72-3;3.6-73-6;3.6-98-1", "", "3.6-1-1;3.8-;3.6-77-;3.6-97-", "из имени",
     "высокая", f"{P}; {T}", "плита бетонируется автобетононасосом (3.6-73-6)"),
    ("Стены ЖБ подземные и цокольные", r"стена_\d+мм_жб", "", "", r"подзем|цокол",
     "3.6-74-1;3.6-75-1;3.6-76-3;3.6-98-1", "3.6-77-", "3.6-94-;3.6-95-;3.6-97-;3.6-1-1", "из имени",
     "высокая", f"{P}; {T}", "таблицы 3.6-74…77 — подземная и цокольная части"),
    ("Стены ЖБ надземные", r"стена_\d+мм_жб", "", "", r"надзем",
     "3.6-94-1;3.6-95-1;3.6-96-3;3.6-98-1", "3.6-97-", f"3.6-74-;3.6-77-;3.6-1-1;{HEIGHTS_OTHER}", "из имени",
     "высокая", f"{P}; {T}", "высота 80 м → таблицы 3.6-94…97 (75–105 м)"),
    ("Перекрытия ЖБ надземные", r"перекрытие_\d+мм_жб", "", "", r"надзем",
     "3.6-94-4;3.6-95-4;3.6-96-5;3.6-98-1", "3.6-97-",
     f"3.6-104-;3.6-14-;3.9-72-;3.6-74-;3.6-77-;{HEIGHTS_OTHER}", "из имени",
     "высокая", f"{P}; {T}", "несъёмная опалубка, засыпка, анкеры — не для обычного перекрытия"),
    ("Перекрытия ЖБ цокольные и подземные", r"перекрытие_\d+мм_жб", "", "", r"цокол|подзем",
     "3.6-74-4;3.6-75-4;3.6-76-5;3.6-98-1", "3.6-77-", "3.6-94-;3.6-97-;3.6-14-", "из имени",
     "средняя", f"{P}; {T}", ""),
    ("Лестничные марши надземные", r"лестниц", "", r"stairflight|^ifcstair$", r"надзем",
     "3.6-94-5;3.6-95-5;3.6-97-12;3.6-98-1", "3.6-96-", "3.6-74-;3.6-77-", "из имени",
     "высокая", f"{P}; {T}", "арматура маршей — оценка по норме (нет данных в модели)"),
    ("Лестничные площадки надземные", r"лестниц", "", r"ifcslab", r"надзем",
     "3.6-94-4;3.6-95-4;3.6-96-5;3.6-97-13;3.6-98-1", "", "3.6-74-;3.6-77-;3.6-97-12;3.6-94-5;3.6-95-5;3.6-96-6", "из имени",
     "средняя", P, ""),
    ("Лестничные марши цокольные и подземные", r"лестниц", "", r"stairflight|^ifcstair$", r"цокол|подзем",
     "3.6-74-5;3.6-75-5;3.6-77-12;3.6-98-1", "3.6-76-", "3.6-94-;3.6-97-", "из имени",
     "средняя", f"{P}; {T}", ""),
    ("Лестничные площадки цокольные и подземные", r"лестниц", "", r"ifcslab", r"цокол|подзем",
     "3.6-74-4;3.6-75-4;3.6-76-5;3.6-77-13;3.6-98-1", "", "3.6-94-;3.6-97-", "из имени",
     "средняя", P, ""),
    ("Балка ЖБ надземная", r"балк", "", "", r"надзем",
     "3.6-94-;3.6-95-;3.6-96-5;3.6-98-1", "3.6-97-", "3.6-74-;3.6-77-", "из имени",
     "средняя", P, "в перечне для балок указана опалубка перекрытий"),
    ("Гидроизоляция Техноэласт", "", r"техноэласт", "", "",
     "3.8-2-11", "", "3.8-2-12;3.8-48-2;3.6-", "нет",
     "высокая", P, "рулонная наплавляемая; «в 2 слоя» — удвоенная площадь"),
    ("Мембрана ВИЛЛАДРЕЙН", "", r"вилладрейн", "", "",
     "3.8-2-12", "", "3.8-2-11;3.6-", "нет", "высокая", P, "профилированная мембрана"),
    ("Утеплитель стен подвала (ППС)", r"утеплитель", r"пенополист|ппс", "", r"подзем|цокол",
     "", "3.12-55-;3.26-", "3.6-;3.8-", "нет", "низкая", M,
     "в перечне КР нет — какую расценку ставить, уточнить у сметчика"),
    ("Термовкладыши / утеплитель ЭППС в перекрытиях", r"утеплитель", "", "", r"надзем",
     "", "3.26-27-;3.26-", "3.6-;3.8-", "нет", "средняя", P, "в перечне: «Устройство термовкладышей»"),
]


# ---------------------------------------------------------------- файл эталона

def write_default(path: str = ETALON_XLSX) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    wb = Workbook()
    ws = wb.active
    ws.title = "Эталон"
    ws.append(COLUMNS)
    for c in ws[1]:
        c.font = Font(bold=True)
    for rule in DEFAULT_RULES:
        ws.append(list(rule))
    for col, width in zip("ABCDEFGHIJKL", (34, 28, 18, 18, 16, 42, 22, 44, 12, 22, 30, 50)):
        ws.column_dimensions[col].width = width
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(path)


def load_rules(path: str = ETALON_XLSX) -> List[Dict[str, Any]]:
    if not os.path.exists(path):
        write_default(path)
    from openpyxl import load_workbook
    ws = load_workbook(path, read_only=True).active
    rows = list(ws.iter_rows(values_only=True))
    head = [str(h or "").strip() for h in rows[0]]
    rules = []
    for r in rows[1:]:
        d = {h: ("" if v is None else str(v).strip()) for h, v in zip(head, r)}
        if d.get("Правило"):
            rules.append(d)
    return rules


# ---------------------------------------------------------------- проверка

def _codes(text: str) -> List[str]:
    return [c.strip() for c in re.split(r"[;,\n]", text or "") if c.strip()]


def _has(pms, code: str) -> bool:
    return any(pm == code or (code.endswith("-") and pm.startswith(code)) for pm in pms)


def _re(pattern: str, value: str) -> bool:
    return not pattern or bool(re.search(pattern, value or "", re.I))


def find_rule(t: Dict[str, Any], rules) -> Optional[Dict[str, Any]]:
    el = t["el"]
    for rule in rules:
        name_mssk = f"{el.get('name') or ''} {kb.enrich(el).get('mssk_name') or ''}"
        if (_re(rule["Имя (regex)"], name_mssk) and _re(rule["Материал (regex)"], t["material"])
                and _re(rule["IFC класс (regex)"], el.get("ifc_class"))
                and _re(rule["Часть здания (regex)"], t["part"])):
            return rule
    return None


def check(t: Dict[str, Any], rule: Dict[str, Any], rows) -> List[str]:
    pms = {r["pressmark"] for r in rows}
    problems = []
    missing = [c for c in _codes(rule["Должно быть (все)"]) if not _has(pms, c)]
    if missing:
        problems.append("нет: " + ", ".join(missing))
    one_of = _codes(rule["Должно быть (хотя бы одно)"])
    if one_of and not any(_has(pms, c) for c in one_of):
        problems.append("нет ни одного из: " + ", ".join(one_of))
    extra = sorted({pm for pm in pms for c in _codes(rule["Не должно быть"])
                    if pm == c or (c.endswith("-") and pm.startswith(c))})
    if extra:
        problems.append("лишнее: " + ", ".join(extra))

    want = rule["Класс бетона"].lower()
    conc = [r for r in rows if r["pressmark"].startswith("1.3-1-")]
    got = sorted({pp._bwf(r["title"])[0] for r in conc if pp._bwf(r["title"])[0]})
    if want == "нет" and conc:
        problems.append(f"бетон не нужен, а стоит B{got}")
    elif want and want != "нет":
        exp = pp._bwf(f"{t['el'].get('name')} {t['material']}")[0] if want == "из имени" \
            else float(want.replace(",", "."))
        if exp is not None and exp not in got:
            problems.append(f"бетон: нужен B{exp:g}, стоит " + (", ".join(f"B{g:g}" for g in got) or "ничего"))
    return problems


def main() -> None:
    args = sys.argv[1:]
    if args and args[0] == "init":
        if os.path.exists(ETALON_XLSX) and "force" not in args:
            print(f"Эталон уже есть: {ETALON_XLSX} (для пересоздания: init force)")
        else:
            write_default()
            print(f"Эталон создан: {ETALON_XLSX} ({len(DEFAULT_RULES)} правил)")
        return

    rules = load_rules()
    if args and args[0] == "run":
        import glob, json
        path = sorted(glob.glob("/app/outputs/*/run_*/Финальный_перечень_работ.json"),
                      key=os.path.getmtime)[-1]
        data = json.load(open(path, encoding="utf-8"))
        print(f"Проверка результата интерфейса: {path}\n")
        ok = bad = norule = 0
        for n, g in enumerate(data.get("groups") or [], 1):
            el = g.get("el") or {"name": g.get("element")}
            t = {"el": el, "part": g["part"], "name": g["element"],
                 "material": g.get("material") or el.get("material") or ""}
            rule = find_rule(t, rules)
            label = f"{n:>3} {g['part']:<9} {g['element'][:48]:<48}"
            if not rule:
                norule += 1
                print(f"  — {label} | нет правила в эталоне")
                continue
            problems = check(t, rule, g.get("rows") or [])
            if problems:
                bad += 1
                print(f"  ❌ {label} | {rule['Правило']}\n       " + "\n       ".join(problems))
            else:
                ok += 1
                print(f"  ✅ {label} | {rule['Правило']}")
        print(f"\nИТОГО: верно {ok} из {ok + bad} (без правила: {norule})")
        return
    session = km._latest_session()
    types = km.collect_types(session)
    height = kb.building_height_for_run(os.path.join(session, "x", "y.json"), {})
    subs = kb.load_perechen()
    use_cache = "nocache" not in args
    print(f"Эталон: {ETALON_XLSX} ({len(rules)} правил) | модель: {len(types)} типов | "
          f"высота {height} м\n", flush=True)

    ok = bad = norule = 0
    by_rule: Dict[str, List[bool]] = {}
    t0 = time.time()
    only = {int(a) for a in args if a.isdigit()}
    for n, t in enumerate(types, 1):
        if only and n not in only:
            continue
        rule = find_rule(t, rules)
        label = f"{n:>3} {t['part']:<9} {t['name'][:48]:<48}"
        if not rule:
            norule += 1
            print(f"  — {label} | нет правила в эталоне", flush=True)
            continue
        q = {k: round(v, 3) if isinstance(v, float) else v for k, v in t["q"].items()}
        res = kb.pick_for_element(t["el"], t["part"], subs, height, use_cache=use_cache)
        rows = pp.finalize_rows(t["el"], q, res["works"], subs)
        problems = check(t, rule, rows)
        by_rule.setdefault(rule["Правило"], []).append(not problems)
        conf = "" if rule["Уверенность"].startswith("высок") else f" (уверенность: {rule['Уверенность']})"
        if problems:
            bad += 1
            print(f"  ❌ {label} | {rule['Правило']}{conf}\n       " + "\n       ".join(problems), flush=True)
        else:
            ok += 1
            print(f"  ✅ {label} | {rule['Правило']}", flush=True)

    print("\n" + "=" * 100)
    print(f"ИТОГО: верно {ok} из {ok + bad} (без правила: {norule}) | время {time.time() - t0:.0f} с")
    for name, results in by_rule.items():
        print(f"  {'✅' if all(results) else '❌'} {name}: {sum(results)}/{len(results)}")


if __name__ == "__main__":
    main()
