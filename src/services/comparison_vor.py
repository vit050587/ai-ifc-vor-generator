"""Comparison view preserving the generated VOR groups and work order."""
import copy
from collections import Counter
from functools import lru_cache
from pathlib import Path

LABELS = {'operation_linked':'Связь установлена через общую операцию','quantity_allocation_required':'Работы сопоставлены; общий объём по строкам не распределён','complex_matched':'Комплекс работ сопоставлен; объёмы операций приведены отдельно','matching_unresolved':'Требуется сопоставление по наименованию','matched':'Расхождение в пределах допуска','quantity_mismatch':'Объёмы различаются','only_project':'Только в ВОР заказчика','only_ifc':'Только в сформированном ВОР','ifc_quantity_incomplete':'Неполный объём сформированного ВОР','project_quantity_incomplete':'Неполный объём ВОР заказчика','unit_unresolved':'Требуется проверка единицы'}


TEMPLATE = Path(__file__).resolve().parents[2] / 'data' / 'comparison_vor_template.xlsx'


@lru_cache(maxsize=1)
def template_index():
    from openpyxl import load_workbook
    from .works_comparison.normalization_functions import code_normalize
    workbook = load_workbook(TEMPLATE, read_only=True, data_only=True)
    result = {}
    part = section = subgroup = ''
    work = ''
    for values in workbook.worksheets[0].iter_rows(min_row=5, values_only=True):
        heading = str(values[0] or '')
        if heading.startswith(('Подземная','Надземная','Цоколь')):
            part, section, subgroup = heading, '', ''
        elif heading.startswith('Раздел'):
            section, subgroup = heading, ''
        elif heading.startswith('Подраздел'):
            subgroup = heading
        if values[1]:
            work = str(values[1])
        code = str(values[15] or '').strip()
        if code:
            entry = (part, section, subgroup, work)
            key = code_normalize(code)
            if entry not in result.setdefault(key, []):
                result[key].append(entry)
    workbook.close()
    return result


def position_identity(p):
    return (p.get('sourceFile'),p.get('groupId'),p.get('code'),p.get('unit'),p.get('name') or p.get('workAndCostName'))


def input_statistics(models, projects):
    from .works_comparison.ifc_positions_extractor import IfcPositionsExtractor
    from .works_comparison.positions_extractor import PositionsExtractor
    generated=IfcPositionsExtractor(models).run()
    customer=PositionsExtractor(projects).run()
    from .works_comparison.customer_grouping import customer_group_manifest
    manifest=customer_group_manifest(projects)
    return {'customerGrouping':manifest,'generatedPositions':len(generated),'generatedWorks':sum(p.get('positionType')!='material' for p in generated),
            'generatedMaterials':sum(p.get('positionType')=='material' for p in generated),'customerParsedPositions':len(customer)}


def restore_unpaired_generated(comparison):
    """Recover works hidden by unconfirmed candidates in earlier saved reports."""
    visible=set()
    hidden={}
    for entry in comparison:
        semantic=entry.get('semanticMatch') or {}
        uncertain=bool(entry.get('projectGroup') and semantic and not semantic.get('confirmed') and entry.get('status')!='complex_matched')
        for p in (entry.get('ifcGroup') or {}).get('positions',[]):
            key=position_identity(p)
            if uncertain:
                hidden.setdefault(key,(p,entry))
            else:
                visible.add(key)
    from .works_comparison.utils import sum_positions_quantity
    for key,(p,entry) in hidden.items():
        if key in visible:continue
        positions=[copy.deepcopy(p)]
        group={'normalized_code':p.get('code'),'normalized_unit':p.get('unit'),'positions':positions,**sum_positions_quantity(positions)}
        comparison.append({'normalized_code':p.get('code'),'normalized_unit':p.get('unit'),'projectGroup':None,'ifcGroup':group,'status':'matching_unresolved'})
        visible.add(key)


def refresh_customer_parts(comparison):
    """Recheck saved links after improving location recognition."""
    from .works_comparison.matching_rules import object_scope
    from .works_comparison.utils import sum_positions_quantity
    result=[]
    for original in comparison:
        entry=copy.deepcopy(original)
        cp=(entry.get('projectGroup') or {}).get('positions',[])
        gp=(entry.get('ifcGroup') or {}).get('positions',[])
        accepted=[]; rejected=[]
        for p in cp:
            construction=p.get('_construction') or {}
            anchor={'workAndCostName':construction.get('name') or p.get('workAndCostName'),
                    'sectionName':construction.get('sectionName') or p.get('sectionName')}
            part=object_scope(anchor).get('part')
            if part:
                if p.get('_matching_scope') is not None:p['_matching_scope']['part']=part
                if construction.get('scope'):construction['scope']['part']=part
                group=p.get('_construction_group') or {}
                if group.get('scope'):
                    group['scope']['part']=part
                    if not any(word in group.get('label','') for word in ('подземная','надземная','цокольная')):
                        label={'under':'подземная часть','above':'надземная часть','plinth':'цокольная часть'}[part]
                        head,sep,tail=group.get('label','').partition(', бетон ')
                        group['label']=head+', '+label+sep+tail
            if gp and part and all(object_scope(g).get('part') and object_scope(g)['part']!=part for g in gp):
                rejected.append(p)
            else:accepted.append(p)
        if rejected:
            unpaired=copy.deepcopy(entry)
            unpaired['ifcGroup']=None
            unpaired['projectGroup']={**entry['projectGroup'],'positions':rejected,**sum_positions_quantity(rejected)}
            binding={'confirmed':False,'groups':[],'generated':[],'reason':'Расположение конструкции не совпадает с группой модели.'}
            for p in rejected:p['_construction_match']=binding
            unpaired['semanticMatch']={**entry.get('semanticMatch',{}),'confirmed':False,'constructionMatch':binding,'reason':binding['reason']}
            unpaired['status']='matching_unresolved'
            unpaired.pop('difference',None);unpaired.pop('percentageDifference',None)
            result.append(unpaired)
            entry['projectGroup']={**entry['projectGroup'],'positions':accepted,**sum_positions_quantity(accepted)} if accepted else None
        result.append(entry)
    return result


def consolidate_customer_comparisons(comparison):
    """Compare each generated work against its complete, traceable customer set."""
    import json
    from .works_comparison.matching_rules import operation, object_scope
    from .works_comparison.utils import sum_positions_quantity
    pooled={}
    result=[]
    for entry in comparison:
        gp=(entry.get('ifcGroup') or {}).get('positions',[])
        cp=(entry.get('projectGroup') or {}).get('positions',[])
        semantic=entry.get('semanticMatch') or {}
        binding=semantic.get('constructionMatch') or {}
        recovered=binding.get('confirmed') and 'Одна сформированная работа предложена разным комплексам' in semantic.get('reason','')
        valid=semantic.get('confirmed') or recovered or (semantic.get('relation')=='complex' and semantic.get('objectConfirmed') and binding.get('confirmed') is not False)
        if not gp or not cp or not valid:
            result.append(entry);continue
        key=(tuple(sorted((position_identity(p) for p in gp),key=str)),entry.get('normalized_unit'))
        if key not in pooled:
            clone=copy.deepcopy(entry)
            clone['semanticMatch']={**semantic,'confirmed':True}
            clone['_customer_components']=[]
            pooled[key]=clone;result.append(clone)
        clone=pooled[key]
        clone['_customer_components'].extend(cp)
    for entry in pooled.values():
        unique={}
        for p in entry.pop('_customer_components'):
            # Source coordinates identify a row even after report expansion.
            key=json.dumps({k:p.get(k) for k in ('sourceFiles','sourceLocation','rowNumber','positionNumber','sectionName','workAndCostName','quantity')},sort_keys=True,ensure_ascii=False)
            unique.setdefault(key,p)
        cp=list(unique.values())
        entry['projectGroup']={**entry['projectGroup'],'positions':cp,**sum_positions_quantity(cp)}
        entry['ifcGroup']={**entry['ifcGroup'],**sum_positions_quantity(entry['ifcGroup']['positions'])}
        scopes=[object_scope(p) for p in entry['ifcGroup']['positions']]
        missing=[]
        for dimension,label in [('waterResistance','W'),('frostResistance','F')]:
            if any(object_scope(p).get(dimension) for p in cp) and any(not scope.get(dimension) for scope in scopes):
                missing.append(label)
        entry['_scope_warning']=('В модели не указаны '+ '/'.join(missing)+ '; сопоставление выполнено по остальным характеристикам, эти показатели не проверены.') if missing else ''
        if len(cp)>1:
            entry['semanticMatch']['reason']='Сопоставлен общий объём работ заказчика для сформированной работы. Исходных строк заказчика: '+str(len(cp))+'.'
    # Identical customer sets shared by identical model work groups are compared once.
    shared={}; merged=[]
    for entry in result:
        gp=(entry.get('ifcGroup') or {}).get('positions',[])
        cp=(entry.get('projectGroup') or {}).get('positions',[])
        if gp and cp and (entry.get('semanticMatch') or {}).get('confirmed'):
            sources=tuple(sorted(json.dumps({k:p.get(k) for k in ('sourceFiles','sourceLocation','rowNumber','positionNumber','sectionName','workAndCostName','quantity')},sort_keys=True,ensure_ascii=False) for p in cp))
            works=tuple(sorted(set((p.get('sourceFile'),p.get('code'),p.get('name'),p.get('unit'),(p.get('element') or {}).get('typeName'),object_scope(p).get('part')) for p in gp),key=str))
            key=(sources,works)
            if key in shared:
                target=shared[key]
                positions={position_identity(p):p for p in target['ifcGroup']['positions']+gp}
                target['ifcGroup']={**target['ifcGroup'],'positions':list(positions.values()),**sum_positions_quantity(list(positions.values()))}
                target['_shared_customer_set']=True
                continue
            shared[key]=entry
        merged.append(entry)
    result=merged
    paired={position_identity(p) for entry in result if entry.get('projectGroup') and (entry.get('semanticMatch') or {}).get('confirmed') for p in (entry.get('ifcGroup') or {}).get('positions',[])}
    def source_key(p):
        return json.dumps({k:p.get(k) for k in ('sourceFiles','sourceLocation','rowNumber','positionNumber','sectionName','workAndCostName','quantity')},sort_keys=True,ensure_ascii=False)
    matched_customer={source_key(p) for entry in result if entry.get('ifcGroup') and (entry.get('semanticMatch') or {}).get('confirmed') for p in (entry.get('projectGroup') or {}).get('positions',[])}
    visible=[]
    for entry in result:
        if not entry.get('ifcGroup') and entry.get('projectGroup'):
            positions=[p for p in entry['projectGroup']['positions'] if source_key(p) not in matched_customer]
            if not positions:continue
            entry=copy.deepcopy(entry)
            entry['projectGroup']={**entry['projectGroup'],'positions':positions,**sum_positions_quantity(positions)}
        if not entry.get('projectGroup') and entry.get('ifcGroup'):
            positions=[p for p in entry['ifcGroup']['positions'] if position_identity(p) not in paired]
            if not positions:continue
            entry=copy.deepcopy(entry)
            entry['ifcGroup']={**entry['ifcGroup'],'positions':positions,**sum_positions_quantity(positions)}
        visible.append(entry)
    return visible


def expand_comparison_rows(comparison):
    """One work name per cell; repeated quantities carry explicit accounting flags."""
    from .works_comparison.utils import sum_positions_quantity
    from .works_comparison.matching_rules import operation
    from .works_comparison.compare_positions_groups import ComparePositionsGroups
    result=[]
    counted_customer=set()
    def buckets(positions, model=False, shared=False, operation_group=False):
        grouped={}
        for p in positions:
            name=p.get('name') or p.get('workAndCostName') or ''
            if model and operation_group:name=operation(p) or name
            key=((p.get('sourceFile'),name) if shared else (p.get('sourceFile'),p.get('groupId'),name)) if model else (operation(p) or name)
            grouped.setdefault(key,[]).append(p)
        return list(grouped.values()) or [[]]
    counted_areas=set()
    for index,entry in enumerate(comparison):
        gp=(entry.get('ifcGroup') or {}).get('positions',[])
        cp=(entry.get('projectGroup') or {}).get('positions',[])
        semantic=entry.get('semanticMatch') or {}
        formwork=bool(gp and cp and all(operation(p) in ('formwork_install','formwork_remove') for p in gp)
            and all(operation(p) in ('formwork','formwork_install','formwork_remove') for p in cp)
            and (semantic.get('objectConfirmed') or semantic.get('confirmed')))
        area_values={}
        if formwork:
            for position in gp:
                if position.get('quantity') is not None:
                    area_values.setdefault((position.get('sourceFile'),position.get('groupId')),set()).add(position['quantity'])
        area_conflict=any(len(values)>1 for values in area_values.values())
        generated_buckets=buckets(gp,model=True,shared=entry.get('_shared_customer_set',False),operation_group=entry.get('_cs_aggregate_operation',False))
        customer_buckets=buckets(cp)
        uncertain=bool(cp and semantic and not semantic.get('confirmed') and not formwork)
        if uncertain:generated_buckets=[gp]
        for gi,generated in enumerate(generated_buckets):
            for ci,customer in enumerate(customer_buckets):
                clone=copy.deepcopy(entry)
                clone['_expanded_work_row']=True
                clone['_comparison_group_id']=entry.get('_comparison_group_id',str(index))
                for field,positions in [('ifcGroup',generated),('projectGroup',customer)]:
                    if positions:
                        clone[field]={**(entry.get(field) or {}),'positions':copy.deepcopy(positions),**sum_positions_quantity(positions)}
                    else:clone[field]=None
                clone['_generated_count_once']=ci==0 and (not formwork or gi==0)
                customer_key=(clone['_comparison_group_id'],customer[0].get('workAndCostName') or customer[0].get('name')) if customer else None
                clone['_customer_count_once']=bool(customer_key and customer_key not in counted_customer)
                if customer_key:counted_customer.add(customer_key)
                if formwork:
                    fresh=[]
                    for position in generated:
                        key=(position.get('sourceFile'),position.get('groupId'),position.get('unit'))
                        if key not in counted_areas:
                            fresh.append(position);counted_areas.add(key)
                    totals=sum_positions_quantity(fresh)
                    clone['_generated_quantity_for_totals']=None if area_conflict else totals['totalQuantity'] if fresh else 0
                    clone['_area_conflict']=area_conflict
                    clone['_generated_count_once']=bool(fresh)
                if formwork:
                    clone['semanticMatch']={**semantic,'confirmed':True,'relation':'equivalent' if len(generated)==1 else 'aggregate',
                        'reason':'Опалубка заказчика сопоставлена с операцией сформированного ВОР. Монтаж и демонтаж относятся к одной площади; в общих объёмах площадь учитывается один раз.'}
                    binding=semantic.get('constructionMatch') or {}
                    if binding.get('confirmed') is False:
                        clone['semanticMatch']['confirmed']=False
                    compared=ComparePositionsGroups([clone['projectGroup']],[clone['ifcGroup']]).compare(['normalized_code','normalized_unit'])[0]
                    clone.update({k:v for k,v in compared.items() if k not in ('projectGroup','ifcGroup')})
                    if not clone['semanticMatch']['confirmed']:
                        clone['status']='matching_unresolved';clone.pop('difference',None);clone.pop('percentageDifference',None)
                elif gp and cp and not uncertain and len(generated_buckets)==1 and len(customer_buckets)==1:
                    compared=ComparePositionsGroups([clone['projectGroup']],[clone['ifcGroup']]).compare(['normalized_code','normalized_unit'])[0]
                    clone.update({k:v for k,v in compared.items() if k not in ('projectGroup','ifcGroup')})
                elif len(generated_buckets)>1 or len(customer_buckets)>1:
                    clone.pop('difference',None);clone.pop('percentageDifference',None)
                    if gp and cp and not uncertain:clone['status']='quantity_allocation_required'
                if clone.get('status')=='complex_matched':clone['status']='matching_unresolved'
                result.append(clone)
    return result


def prepare_report(report, models=None):
    from .works_comparison.positions_processor import PositionsProcessor
    from .works_comparison.ifc_positions_processor import IfcPositionsProcessor
    from .works_comparison.compare_positions_groups import ComparePositionsGroups
    from .works_comparison.validation_result_former import ValidationResultFormer
    result = copy.deepcopy(report)
    if result.get('csFinalModels'):
        from .works_comparison.ifc_positions_extractor import IfcPositionsExtractor
        from .works_comparison.units import normalize_positions
        from .cs_model_members import split_slab_roles
        role_models=split_slab_roles(result['csFinalModels'])
        final=normalize_positions(IfcPositionsExtractor(role_models).run())
        models=role_models
        result['inputStatistics']={**result.get('inputStatistics',{}),'generatedPositions':sum(len(g.get('positions',[])) for m in result['csFinalModels'] for g in m.get('elementGroups',[])),'generatedWorks':sum(p.get('positionType')!='material' for m in result['csFinalModels'] for g in m.get('elementGroups',[]) for p in g.get('positions',[])),'generatedMaterials':0}
        # Old CS JSON can use a different unit label; match original group and rate.
        rates={}
        for p in final:rates.setdefault((p.get('sourceFile'),p.get('_sourceGroupId') or p.get('groupId'),p.get('code')),[]).append(p)
        result['comparison']=[entry for entry in result.get('comparison',[]) if (entry.get('projectGroup') or {}).get('positions')]
        for entry in result.get('comparison',[]):
            group=entry.get('ifcGroup')
            if group:
                group['positions']=[copy.deepcopy(member) for p in group.get('positions',[]) if (key:=(p.get('sourceFile'),p.get('_sourceGroupId') or p.get('groupId'),p.get('code'))) in rates for member in rates[key]]
                if not group['positions']:entry['ifcGroup']=None
        represented={position_identity(p) for e in result.get('comparison',[]) for p in (e.get('ifcGroup') or {}).get('positions',[])}
        from .works_comparison.utils import sum_positions_quantity
        for p in final:
            if position_identity(p) not in represented:
                result.setdefault('comparison',[]).append({'normalized_code':p.get('code'),'normalized_unit':p.get('unit'),'status':'matching_unresolved','projectGroup':None,
                    'ifcGroup':{'positions':[copy.deepcopy(p)],**sum_positions_quantity([p])}})
    result['comparison']=[entry for entry in result.get('comparison',[]) if (entry.get('ifcGroup') or {}).get('positions') or (entry.get('projectGroup') or {}).get('positions')]
    result['comparison']=refresh_customer_parts(result.get('comparison',[]))
    if report.get('reconcileConstructionLinks'):
        from .comparison_reconciliation import reconcile_saved_links
        result['comparison']=reconcile_saved_links(result['comparison'],cs_mode=bool(result.get('csFinalModels')))
    customer = [p for g in result.get('comparison', []) for p in (g.get('projectGroup') or {}).get('positions', [])]
    generated = [p for g in report.get('comparison', []) for p in (g.get('ifcGroup') or {}).get('positions', [])]
    if report.get('comparison') and not any(g.get('semanticMatch') for g in report['comparison']):
        cg, keys = PositionsProcessor().run(customer)
        gg, _ = IfcPositionsProcessor().run(generated)
        comparison = ComparePositionsGroups(cg, gg).compare(keys)
        result = ValidationResultFormer(Path('/unused')).form(comparison, report.get('validationStatistics'))
    from .works_comparison.matching_rules import operation
    for entry in result.get('comparison',[]):
        cp=(entry.get('projectGroup') or {}).get('positions',[])
        gp=(entry.get('ifcGroup') or {}).get('positions',[])
        customer_ops={operation(p) for p in cp}-{None}
        generated_ops={operation(p) for p in gp}-{None}
        allowed=customer_ops|({'formwork_install','formwork_remove'} if 'formwork' in customer_ops else set())
        if cp and gp and customer_ops and generated_ops and not generated_ops.issubset(allowed):
            entry['semanticMatch']={**entry.get('semanticMatch',{}),'confirmed':False,'objectConfirmed':False,'reason':'Работы относятся к разным операциям; сопоставление отклонено.'}
            entry['status']='matching_unresolved';entry.pop('difference',None);entry.pop('percentageDifference',None)
    # Steel profiles and pipes are not reinforcement components.
    from .works_comparison.utils import sum_positions_quantity
    extra=[]
    for entry in result.get('comparison',[]):
        gp=(entry.get('ifcGroup') or {}).get('positions',[])
        cp=(entry.get('projectGroup') or {}).get('positions',[])
        if gp and cp and {operation(p) for p in gp}=={'reinforcement'}:
            other=[p for p in cp if any(word in (p.get('workAndCostName') or p.get('name') or '').lower() for word in ('труба','профиль','полоса сталь','уголок','пластина')) and operation(p)!='reinforcement']
            reinforcement=[p for p in cp if p not in other]
            if other:
                unpaired=copy.deepcopy(entry)
                unpaired['ifcGroup']=None
                unpaired['projectGroup']={**entry['projectGroup'],'positions':other,**sum_positions_quantity(other)}
                unpaired['semanticMatch']={**entry.get('semanticMatch',{}),'confirmed':False,'reason':'Профили и другие изделия не включены в объём арматуры; соответствующая работа не найдена.'}
                unpaired['status']='matching_unresolved'
                unpaired.pop('difference',None);unpaired.pop('percentageDifference',None)
                extra.append(unpaired)
                entry['projectGroup']={**entry['projectGroup'],'positions':reinforcement,**sum_positions_quantity(reinforcement)} if reinforcement else None
    result.setdefault('comparison',[]).extend(extra)
    restore_unpaired_generated(result.get('comparison',[]))
    result['comparison'] = expand_comparison_rows(consolidate_customer_comparisons(expand_comparison_rows(consolidate_customer_comparisons(result.get('comparison',[])))))
    result['reportVersion'] = 16
    if models is not None:
        from .works_comparison.ifc_positions_extractor import IfcPositionsExtractor
        from .works_comparison.units import normalize_positions
        positions=normalize_positions(IfcPositionsExtractor(models).run())
        by_group={}
        for p in positions:by_group.setdefault((p.get('sourceFile'),p.get('groupId')),[]).append(p)
        layout=[]
        for model in models:
            for group in model.get('elementGroups',[]):
                members=by_group.get((model.get('sourceFile'),group.get('groupId')),[])
                for p in members:
                    layout.append({'identity':list(position_identity(p)),'sourceFile':p.get('sourceFile'),'groupId':p.get('groupId'),'element':p.get('element') or {}})
                if not members:
                    layout.append({'identity':None,'sourceFile':model.get('sourceFile'),'groupId':group.get('groupId'),'element':group.get('element') or {},'empty':True})
        result['generatedVorLayout']=layout
    layout=result.get('generatedVorLayout') or [{'identity':list(position_identity(p)),'sourceFile':p.get('sourceFile'),'groupId':p.get('groupId'),'element':p.get('element') or {}} for entry in result.get('comparison',[]) for p in (entry.get('ifcGroup') or {}).get('positions',[])]
    work_order={}
    group_layout={}
    for index,item in enumerate(layout):
        if item.get('identity') is not None:work_order.setdefault(tuple(item['identity']),index)
        group_layout.setdefault((item['sourceFile'],item['groupId']),item)
    rows = []
    for group in result.get('comparison', []):
        gp = (group.get('ifcGroup') or {}).get('positions', [])
        cp = (group.get('projectGroup') or {}).get('positions', [])
        names = list(dict.fromkeys(str(p.get('name') or p.get('workAndCostName') or '') for p in gp + cp))
        groups = []
        for p in gp:
            element, props = p.get('element') or {}, p.get('properties') or {}
            label = ', '.join(str(v) for v in [element.get('category'), element.get('typeName'), element.get('level'), props.get('material')] if v)
            label = label or p.get('groupId') or 'Группа не указана'
            if label not in groups:
                groups.append(label)
        location=('','','')
        validations = [p.get('validation_result') for p in gp if isinstance(p.get('validation_result'), dict)]
        confidence = [v['confidence'] for v in validations if isinstance(v.get('confidence'), (float,int)) and not isinstance(v.get('confidence'),bool) and 0 <= v['confidence'] <= 1]
        score = min(confidence) if confidence else None
        verdicts = {v.get('verdict') for v in validations}
        verdict = next((v for v in ('incorrect','error','insufficient_data','uncertain','correct') if v in verdicts), 'not_checked')
        semantic = group.get('semanticMatch') or {}
        reasons = list(dict.fromkeys(v.get('reason','') for v in validations if v.get('reason')))
        if semantic:
            reasons.insert(0, 'Сопоставление: ' + str(semantic.get('reason', '')))
        construction_match = semantic.get('constructionMatch') or next((p.get('_construction_match') for p in cp if p.get('_construction_match')), {})
        element_confirmed = construction_match.get('confirmed') is True
        if not groups and element_confirmed:
            groups = list(dict.fromkeys(g.get('objectLabel') for g in construction_match.get('generated',[]) if g.get('objectLabel')))
        if not groups:
            groups = list(dict.fromkeys('; '.join(filter(None,[p.get('sectionName'),(p.get('_construction') or {}).get('name')])) for p in cp))
            groups = [label for label in groups if label]
        customer_groups = {p['_construction_group']['id']:p['_construction_group'] for p in cp if p.get('_construction_group')}
        customer_labels = [f"{g['label']} · секций: {len(g['sections'])} · объектов: {len(g['objects'])}" for g in customer_groups.values()]
        group_names = groups[0] if len(groups)==1 else f'Несколько групп ({len(groups)}), общий объём' if groups else 'Группа без наименования'
        paired = not (cp and semantic and not semantic.get('confirmed') and group['status']!='complex_matched')
        display_gp = gp if paired else []
        generated_work = '; '.join(dict.fromkeys(p.get('name') or p.get('workAndCostName') or '' for p in display_gp)) or None
        if group.get('_cs_aggregate_operation') and len({p.get('name') for p in display_gp})>1:
            op=operation(display_gp[0]) if display_gp else None
            generated_work={'concrete':'Бетонирование','reinforcement':'Армирование','formwork_install':'Монтаж опалубки','formwork_remove':'Демонтаж опалубки'}.get(op,generated_work)+' (общий объём по группам модели)'
        customer_work = '; '.join(dict.fromkeys(p.get('workAndCostName') or p.get('name') or '' for p in cp)) or None
        customer_names=list(dict.fromkeys(p.get('workAndCostName') or p.get('name') or '' for p in cp))
        customer_ops={operation(p) for p in cp}
        if len(customer_names)>1 or customer_ops=={'reinforcement'}:
            labels={'reinforcement':'Арматура','concrete':'Бетонирование','formwork':'Опалубка'}
            customer_work=labels.get(next(iter(customer_ops))) if len(customer_ops)==1 else None
            customer_work=customer_work or customer_names[0]
        object_names=list(dict.fromkeys((p.get('_construction_group') or {}).get('label') or (p.get('_construction') or {}).get('name') or p.get('sectionName') for p in cp))
        object_names=[name for name in object_names if name]
        if customer_work and object_names:
            customer_work += ' ('+'; '.join(object_names)+')'
        anchor=display_gp[0] if display_gp else None
        source_item=None
        if anchor:
            source_item={'sourceFile':anchor.get('sourceFile'),'groupId':anchor.get('groupId'),'element':anchor.get('element') or {}}
        elif element_confirmed and len(construction_match.get('groups',[]))==1:
            source_item=group_layout.get(tuple(construction_match['groups'][0]))
        if source_item:
            el=source_item.get('element') or {}
            location=(source_item.get('sourceFile') or 'Сформированный ВОР',
                      (el.get('typeName') or 'Группа элементов')+' ('+str(source_item.get('groupId') or '')+')','')
            if group.get('_shared_customer_set'):
                ids=list(dict.fromkeys(str(p.get('groupId')) for p in display_gp))
                location=(location[0],(el.get('typeName') or 'Группы элементов')+' ('+', '.join(ids)+')','')
            if el.get('ifcClass')=='IfcStair':
                location=(location[0],'Лестница — родительский объект IFC ('+str(source_item.get('groupId') or '')+')','')
            group_names=''
        elif element_confirmed:
            location=('Работы заказчика для нескольких групп модели',', '.join(str(g[1]) for g in construction_match.get('groups',[])),'')
            group_names=''
        else:
            location=('Работы ВОР заказчика без сопоставления',(cp[0].get('_construction_group') or {}).get('sectionName') or cp[0].get('sectionName') if cp else '','')
            group_names=''
        if anchor:
            display_order=(0,work_order.get(position_identity(anchor),len(work_order)))
        elif source_item:
            indices=[v for key,v in work_order.items() if key[:2]==(source_item['sourceFile'],source_item['groupId'])]
            display_order=(0,max(indices,default=len(work_order))+.5)
        else:
            display_order=(1 if element_confirmed else 2,str(cp[0].get('sourceFiles')) if cp else '',str(cp[0].get('sectionName')) if cp else '',str((cp[0].get('_construction_group') or {}).get('id')) if cp else '',str(operation(cp[0]) or cp[0].get('workAndCostName')) if cp else '')
        confidence_text = ('По правилам' if semantic.get('method')=='rules' else f"{semantic['confidence']:.0%}" if isinstance(semantic.get('confidence'),(int,float)) else 'Не оценивалась')
        if cp and not display_gp:
            confidence_text = 'Принадлежность к группе подтверждена; пара работ не найдена' if element_confirmed else 'Связь не установлена'
        comments = [LABELS.get(group['status'],group['status']), 'Ед. изм.: '+str(group.get('normalized_unit') or '')]
        if cp and not display_gp and element_confirmed:
            comments.insert(0,'Работа заказчика сохранена в подтверждённой группе элементов; соответствующая сформированная работа не установлена.')
        if group.get('difference') is not None and (group.get('ifcGroup') or {}).get('totalQuantity'):
            percentage=group['difference'] / abs(group['ifcGroup']['totalQuantity'])
            comments.append(('Заказчик завысил' if percentage>0 else 'Заказчик занизил' if percentage<0 else 'Объёмы совпадают')+(f' на {abs(percentage):.2%} относительно сформированного ВОР' if percentage else ''))
        if group.get('_area_conflict'):
            comments.append('Площади монтажа и демонтажа в исходном ВОР различаются; общий объём опалубки требует проверки.')
        if not group.get('_generated_count_once',True) or not group.get('_customer_count_once',True):
            comments.append('Повтор общей площади / объёма для отдельной работы; при суммировании повторно не учитывать.')
        if group.get('_scope_warning'):comments.append(group['_scope_warning'])
        allocations=[p.get('quantityCalculation',{}).get('allocation') for p in display_gp if isinstance(p.get('quantityCalculation'),dict) and p['quantityCalculation'].get('allocation')]
        if allocations:comments.append('Часть исходной группы выделена по GUID элементов и их конструктивному назначению. Объём распределён по исходным измерителям элементов; суммы частей равны объёму исходного ВОР.')
        if any(p.get('_cs_operation_family') for p in display_gp):
            comments.append('Альтернативные расценки ЦС на одну операцию. Общий объём учитывается один раз; применимость конкретной расценки этим сравнением не определяется.')
            if all(p.get('_cs_operation_primary') is False for p in display_gp):
                group['_generated_count_once']=False
                group['_generated_quantity_for_totals']=0
                comments.append('Объём операции сравнивается в основной строке, повторно не суммировать.')
        if any((p.get('element') or {}).get('ifcClass')=='IfcStair' for p in display_gp):
            comments.append('Исходное имя в IFC: '+str((display_gp[0].get('element') or {}).get('typeName') or '')+'. Родительский объект лестницы IfcStair. Объёмы дочерних маршей и площадок здесь не сравниваются; название родителя не подтверждает способ изготовления дочерних элементов.')
        if group.get('_shared_customer_set'):
            comments.append('Общий объём по группам модели: '+ '; '.join(str(p.get('groupId'))+': '+str(p.get('quantity'))+' '+str(p.get('unit')) for p in display_gp)+'. Объём заказчика сравнивается с суммой этих групп один раз.')
        if customer_ops=={'reinforcement'}:
            details=list(dict.fromkeys((child.get('workAndCostName') or child.get('name') or '') for p in cp for child in p.get('sourceDetailRows',[])))
            details+= [name for name in customer_names if name and name[0].lower() in ('d','д','ø','ф','⌀')]
            if details:comments.append('Состав арматуры: '+', '.join(dict.fromkeys(details)))
        if validations:
            comments.extend('Проверка подбора: '+v['reason'] for v in validations if v.get('reason'))
        rows.append({'displayOrder':display_order,'quantityAccountingKey':group.get('_comparison_group_id'),'generatedQuantityForTotals':(group.get('_generated_quantity_for_totals',(group.get('ifcGroup') or {}).get('totalQuantity') if group.get('_generated_count_once',True) else 0)),
                     'customerQuantityForTotals':((group.get('projectGroup') or {}).get('totalQuantity') if group.get('_customer_count_once',True) else 0),'generatedWork':generated_work,'customerWork':customer_work,'elementConfirmed':element_confirmed,
                     'modelReason':str(semantic.get('reason') or '; '.join(reasons)),
                     'comparisonConfidence':confidence_text,'comments':'; '.join(comments),'part':location[0],'section':location[1],'subsection':location[2], 'elementGroup':group_names,'elementGroups':groups,
                     'work':'; '.join(names),'code':None if str(group.get('normalized_code','')).startswith(('__uncoded__:', '__semantic__')) else group.get('normalized_code'),'unit':group.get('normalized_unit'),
                     'customerQuantity':(group.get('projectGroup') or {}).get('totalQuantity'),
                     'generatedQuantity':(group.get('ifcGroup') or {}).get('totalQuantity') if display_gp else None,
                     'presence':'Требуется сопоставление' if group['status']=='matching_unresolved' else 'В обоих ВОР' if gp and cp else 'Только в сформированном ВОР' if gp else 'Только в ВОР заказчика',
                     'status':group['status'],'result':('Объём не определён или неполон' if group['status'] in ('only_ifc','only_project') and any((side or {}).get('positionsWithoutQuantity',0)>0 for side in (group.get('ifcGroup'),group.get('projectGroup'))) else LABELS.get(group['status'],group['status'])),
                     'difference':group.get('difference'),'differencePercent':None if group.get('difference') is None or not (group.get('ifcGroup') or {}).get('totalQuantity') else group['difference'] / abs(group['ifcGroup']['totalQuantity']) * 100,
                     'validationVerdict':verdict,'confidence':score,'confidenceLabel':'Не оценивалась' if score is None else 'Высокая' if score >= .85 else 'Средняя' if score >= .6 else 'Низкая',
                     'customerConstructionGroups':customer_labels,'matchingMethod':semantic.get('method'),'matchingConfidence':semantic.get('confidence'),'matchingRelation':semantic.get('relation'),'reason':reasons[0] if reasons else '','allocationNote':'',
                     'groupCount':len({(p.get('sourceFile'),p.get('groupId')) for p in gp}),
                     'detail':group})
    from .works_comparison.cs_operations import link_operation_rows
    link_operation_rows(rows)
    parents={}
    def root(key):
        parents.setdefault(key,key)
        if parents[key]!=key:parents[key]=root(parents[key])
        return parents[key]
    for row in rows:
        positions=(row['detail'].get('ifcGroup') or {}).get('positions',[]) if row['generatedWork'] else []
        keys=list(dict.fromkeys((p.get('sourceFile'),p.get('groupId')) for p in positions))
        for key in keys[1:]:parents[root(key)]=root(keys[0])
    components={}
    for key in parents:components.setdefault(root(key),[]).append(key)
    for row in rows:
        positions=(row['detail'].get('ifcGroup') or {}).get('positions',[]) if row['generatedWork'] else []
        if not positions:continue
        key=(positions[0].get('sourceFile'),positions[0].get('groupId'))
        component=components.get(root(key),[key])
        if len(component)>1:
            component.sort(key=lambda key:min((v for identity,v in work_order.items() if identity[:2]==key),default=len(layout)))
            first=group_layout.get(component[0],{})
            row['part']=first.get('sourceFile') or row['part']
            row['section']=((first.get('element') or {}).get('typeName') or 'Группы элементов')+' ('+', '.join(str(key[1]) for key in component)+')'
            first_index=min((v for identity,v in work_order.items() if identity[:2] in component),default=len(layout))
            row['displayOrder']=(0,first_index,row['displayOrder'][1])
    for row in rows:
        positions=(row['detail'].get('ifcGroup') or {}).get('positions',[]) if row['generatedWork'] else []
        if not positions:continue
        key=(positions[0].get('sourceFile'),positions[0].get('groupId'))
        component=components.get(root(key),[key])
        identities=list(dict.fromkeys(((group_layout.get(member,{}) or {}).get('element') or {}).get('sourceGroupId') or member[1] for member in component))
        role=next(((p.get('element') or {}).get('structuralRole') for p in positions if (p.get('element') or {}).get('structuralRole')),None)
        from .works_comparison.matching_rules import object_scope
        scope=object_scope(positions[0])
        if role or len(identities)>4:
            title={'basement_cover':'Перекрытия подвала','upper_slabs':'Перекрытия надземной части'}.get(role)
            if not title:title={'flight':'Лестничные марши','landing':'Лестничные площадки'}.get(scope.get('subtype'),'Группы элементов')
            row['section']=title+' ('+str(len(identities))+' исходных групп)'
        if any('::' in str(p.get('groupId')) for p in positions):
            row['comments']+=' Исходные группы: '+', '.join(str(i) for i in identities)+'.'
    for index,item in enumerate(layout):
        if not item.get('empty'):continue
        element=item.get('element') or {}
        label=(element.get('typeName') or element.get('materialGroup') or 'Группа элементов')+' ('+str(item.get('groupId') or '')+')'
        comment='Работы для группы не подобраны.'
        if element.get('count') is not None:comment+=' Количество элементов: '+str(element['count'])+'.'
        rows.append({'isEmptyGroup':True,'displayOrder':(0,index),'part':item.get('sourceFile') or 'Сформированный ВОР','section':label,'subsection':'','elementGroup':'',
                     'generatedWork':None,'customerWork':None,'generatedQuantity':None,'customerQuantity':None,'difference':None,'differencePercent':None,
                     'modelReason':'','comparisonConfidence':'','comments':comment,'unit':None,'work':'','code':None,'status':'empty_group','presence':'Группа без работ','result':'Работы не подобраны',
                     'quantityAccountingKey':None,'generatedQuantityForTotals':0,'customerQuantityForTotals':0,'detail':{}})
    rows.sort(key=lambda row:tuple(row['displayOrder']))
    result['matchingStatistics'] = next((g['matchingStatistics'] for g in result.get('comparison', []) if g.get('matchingStatistics')), {})
    from .comparison_similarity import composition_similarity
    result['compositionSimilarity']=composition_similarity(rows)
    result['vorRows'] = rows
    result['vorSummary'] = dict(Counter(row['presence'] for row in rows if not row.get('isEmptyGroup')))
    visible_generated = {(p.get('sourceFile'),p.get('_sourceGroupId') or p.get('groupId'),p.get('code'),p.get('unit'),p.get('name') or p.get('workAndCostName')) for row in rows if row['generatedWork'] for p in (row['detail'].get('ifcGroup') or {}).get('positions',[])}
    paired_generated = {(p.get('sourceFile'),p.get('_sourceGroupId') or p.get('groupId'),p.get('code'),p.get('unit'),p.get('name') or p.get('workAndCostName')) for row in rows if row['generatedWork'] and row['customerWork'] for p in (row['detail'].get('ifcGroup') or {}).get('positions',[])}
    matched_families={p.get('_cs_operation_family') for row in rows if row['generatedWork'] and row['customerWork'] for p in (row['detail'].get('ifcGroup') or {}).get('positions',[]) if p.get('_cs_operation_family')}
    paired_generated.update((p.get('sourceFile'),p.get('_sourceGroupId') or p.get('groupId'),p.get('code'),p.get('unit'),p.get('name') or p.get('workAndCostName')) for row in rows if row['generatedWork'] for p in (row['detail'].get('ifcGroup') or {}).get('positions',[]) if p.get('_cs_operation_family') in matched_families)
    source_rows=set()
    def trace_customer(position):
        import json
        source_rows.add((tuple(position.get('sourceFiles') or []),json.dumps(position.get('sourceLocation'),sort_keys=True),position.get('rowNumber'),position.get('sectionName'),position.get('positionNumber'),position.get('workAndCostName')))
        for field in ('sourceDetailRows','sourceContextRows'):
            for child in position.get(field,[]):trace_customer(child)
    for row in rows:
        for position in (row['detail'].get('projectGroup') or {}).get('positions',[]):trace_customer(position)
    result['comparisonStatistics'] = {**result.get('comparisonStatistics',{}),'totalGroups':len(result.get('comparison',[])),
        'byStatus':dict(Counter(entry.get('status') for entry in result.get('comparison',[])))}
    result['displayStatistics'] = {'customerRepresentedSourceRows':len(source_rows),'generatedVisiblePositions':len(visible_generated),'generatedUnpairedPositions':len(visible_generated-paired_generated),
        'customerWorkGroups':len({row['quantityAccountingKey'] for row in rows if row['customerWork']}),'customerUnpairedGroups':len({row['quantityAccountingKey'] for row in rows if row['customerWork'] and not row['generatedWork']}),
        'matchedWorkGroups':len({row['quantityAccountingKey'] for row in rows if row['customerWork'] and row['generatedWork']}),'quantityMismatches':sum(row['status']=='quantity_mismatch' for row in rows)}
    compact=[]
    for row in rows:
        positions=(row['detail'].get('ifcGroup') or {}).get('positions',[]) if row['generatedWork'] else []
        if positions and all(p.get('_cs_operation_primary') is False for p in positions):continue
        codes=sorted({str(code) for p in positions for code in p.get('_cs_operation_codes',[])})
        if len(codes)>1:row['comments']+=' Варианты расценок одной операции: '+', '.join(codes)+'. Полный состав сохранён на листе «Подробности».'
        compact.append(row)
    result['vorRows']=compact
    manifest=(result.get('inputStatistics') or {}).get('customerGrouping')
    if manifest:
        from .works_comparison.customer_grouping import customer_group_coverage
        result['displayStatistics'].update(customer_group_coverage(manifest,rows))
        result['matchingStatistics']['reportReconciled']=True
        similarity=result['compositionSimilarity']
        coverage=result['displayStatistics']['matchedWorkGroups']/manifest['totalGroups'] if manifest['totalGroups'] else None
        similarity.update(customerTotal=manifest['totalGroups'],customerMatched=result['displayStatistics']['matchedWorkGroups'],customerCoverage=None if coverage is None else round(coverage*100,1))
        generated_coverage=similarity['generatedMatched']/similarity['generatedTotal'] if similarity['generatedTotal'] else 0
        similarity['score']=None if not manifest['totalGroups'] and not similarity['generatedTotal'] else round(200*generated_coverage*(coverage or 0)/(generated_coverage+(coverage or 0)),1) if generated_coverage+(coverage or 0) else 0
    return result
