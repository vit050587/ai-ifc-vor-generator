"""
Итоговый перечень работ АР по перечню сметчиков (perechen_kb) — замена
works_final_selector при WORKS_SELECTOR=perechen.

  * листовая группа из filtered_elements_grouped_AR.json делится по ТИПУ элемента
    (имя без ID + материал): каждый тип подбирается отдельно (с кэшем);
  * объёмы — поэлементно из filtered_elements.xlsx, суммой по типу;
  * стоимость — ЗП/ЭМ/МР из локального справочника × объём.
Выход: Финальный_перечень_работ.xlsx / .json в папке запуска.
"""
import json
import os
import re
from typing import Any, Dict, List, Optional, Tuple

from src.core.logger import setup_logger
from src.services import perechen_kb as kb

logger = setup_logger(__name__)

FINAL_XLSX = "Финальный_перечень_работ.xlsx"
FINAL_JSON = "Финальный_перечень_работ.json"

PART_ORDER = ("Подземная", "Цоколь", "Надземная")
PART_TITLES = {"Подземная": "ПОДЗЕМНАЯ ЧАСТЬ", "Цоколь": "ЦОКОЛЬНАЯ ЧАСТЬ",
               "Надземная": "НАДЗЕМНАЯ ЧАСТЬ"}
PART_TOTALS = {"Подземная": "ИТОГО по подземной части:", "Цоколь": "ИТОГО по цокольной части:",
               "Надземная": "ИТОГО по надземной части:"}
HEADERS = ["Шифр ТСН", "Наименование расценки/ресурса", "Ед. изм.", "Объём работ",
           "ЗП", "ЭМ", "МР", "Стоимость"]

_local: Optional[Dict[str, Dict[str, Any]]] = None


def _local_rate(pressmark: str) -> Optional[Dict[str, Any]]:
    global _local
    if _local is None:
        _local = {}
        try:
            from src.services.local_works import load_local_works
            for works in load_local_works()["tables"].values():
                for w in works:
                    _local[w["pressmark"]] = w
        except Exception as exc:
            logger.warning(f"Локальный справочник недоступен: {exc}")
    return _local.get(pressmark)


# ---------------------------------------------------------------- дерево групп

def _children(node: Dict[str, Any]) -> List[Dict[str, Any]]:
    return node.get("children") or node.get("subgroups") or node.get("groups") or []


def _leaves(nodes: List[Dict[str, Any]], path: Tuple[str, ...] = ()):
    for n in nodes or []:
        p = path + (str(n.get("name", "")),)
        ch = _children(n)
        if ch:
            yield from _leaves(ch, p)
        else:
            yield n, p


def _part_from_path(path: Tuple[str, ...]) -> str:
    for name in path:
        p = kb._part_in(name)
        if p:
            return p
    return "Надземная"


def _leaf_area(node: Dict[str, Any]) -> float:
    areas = node.get("total_areas") or {}
    v = areas.get("Площадь, м2")
    if isinstance(v, (int, float)) and v > 0:
        return float(v)
    nums = [x for x in areas.values() if isinstance(x, (int, float))]
    return float(max(nums)) if nums else 0.0


# ---------------------------------------------------------------- объёмы по элементам

def _element_quantities(run_dir: str, filename: str = "filtered_elements.xlsx") -> List[Dict[str, float]]:
    """Объём, площадь, площадь опалубки и расход арматуры каждой строки Excel.

    Площадь опалубки считается ровно так же, как в режиме ЦС (КР) — той же
    функцией `group_excel._element_formwork_area_m2`: периметр × толщина
    (боковые грани; для перекрытий дополнительно площадь плиты). У стен,
    балок и прочих элементов периметра/толщины нет — там значение равно
    нулю, а объём опалубки берётся из «Площадь, м2» (см. `_formwork_area`).
    """
    import pandas as pd
    path = os.path.join(run_dir, filename)
    try:
        try:
            df = pd.read_excel(path, sheet_name="Данные")
        except Exception:
            df = pd.read_excel(path)
    except Exception as exc:
        logger.warning(f"Нет поэлементных объёмов ({path}): {exc}")
        return []

    # Функция и колонки — те же, что использует КР (ifc_reference_builder /
    # group_excel) при расчёте площади опалубки групп.
    from src.services.group_excel import (
        _element_formwork_area_m2,
        _find_formwork_columns,
    )

    cols = [str(c) for c in df.columns]
    vol_cols = [c for c in cols if ("объём" in c.lower() or "volume" in c.lower())
                and "литр" not in c.lower()]
    vol_cols.sort(key=lambda c: 0 if "netvolume" in c.lower() else 1 if "объём, м3" in c.lower() else 2)
    area_main = "Площадь, м2" if "Площадь, м2" in cols else None
    area_cols = [c for c in cols if "площад" in c.lower() and "м2" in c]
    rebar_col = next((c for c in cols if "reinforcementvolumeratio" in c.lower()), None)
    perim_cols, depth_cols = _find_formwork_columns(cols)

    def num(row, c) -> float:
        try:
            v = float(str(row[c]).replace(",", "."))
            return v if v == v and v > 0 else 0.0
        except (TypeError, ValueError, KeyError):
            return 0.0

    out = []
    df.columns = cols
    for _, r in df.iterrows():
        vol = next((num(r, c) for c in vol_cols if num(r, c) > 0), 0.0)
        area = num(r, area_main) if area_main else 0.0
        if not area and area_cols:
            area = max(num(r, c) for c in area_cols)
        fw = _element_formwork_area_m2(r.to_dict(), perim_cols, depth_cols)
        out.append({"volume": vol, "area": area, "fw_hint": fw,
                    "rebar_kg": num(r, rebar_col) * vol if rebar_col else 0.0})
    return out


def _divisor(unit: str) -> float:
    m = re.match(r"\s*(\d+(?:[.,]\d+)?)", unit or "")
    return float(m.group(1).replace(",", ".")) if m else 1.0


def _element_kind(el: Dict[str, Any]) -> str:
    """Вид конструкции — по имени (IFC-класс в моделях часто неточен), затем по классу."""
    n = str(el.get("name") or "").lower()
    for kind, words in (("wall", ("стена", "стен", "приямок", "приямк", "парапет")), ("column", ("колонн", "пилон")),
                        ("beam", ("балк", "ригел")), ("stair", ("лестн", "марш")),
                        ("slab", ("перекрыт", "плита", "покрыт", "площадк", "фундамент"))):
        if any(w in n for w in words):
            if kind == "stair" and el.get("ifc_class") == "IfcSlab":
                return "slab"   # лестничная площадка — плита
            return kind
    return {"IfcWall": "wall", "IfcWallStandardCase": "wall", "IfcColumn": "column",
            "IfcBeam": "beam", "IfcSlab": "slab", "IfcFooting": "slab",
            "IfcStair": "stair", "IfcStairFlight": "stair"}.get(el.get("ifc_class"), "other")


def _formwork_area(el: Dict[str, Any], q: Dict[str, float]) -> float:
    """Площадь опалубки группы — по правилам ЦС (КР).

    Раньше площадь считалась «на глаз» по геометрии (для стен — 2 × объём /
    толщина, для плит — объём / толщина), из-за чего объёмы опалубки в ИИ
    расходились с ЦС: у стен ровно в 2 раза, у плит — на десятки процентов.
    Теперь правила те же, что в КР (`api_works_lookup._calculate_work_volume`):

      * плиты с данными периметра/толщины (фундаментные плиты, перекрытия,
        приямки, лестничные площадки) — периметр × толщина, для перекрытий
        дополнительно площадь плиты (`q["fw_hint"]`, сумма по элементам);
      * стены, балки и прочие элементы (периметра/толщины нет) —
        «Площадь, м2» группы (`q["area"]`);
      * плитные группы (плита/перекрытие) без рассчитанной площади опалубки —
        объём не считается: площадь плиты не является площадью опалубки.
    """
    fw = float(q.get("fw_hint") or 0.0)
    if fw > 0:
        return fw
    name = str(el.get("name") or "").lower()
    if "плит" in name or "перекрыт" in name:
        return 0.0
    return float(q.get("area") or 0.0)


def _curing_rate(el: Dict[str, Any], existing) -> Optional[Dict[str, Any]]:
    """Уход за бетоном обязателен для монолита — добавляем, если не выбран."""
    if not kb._is_concrete(el) or kb.is_precast(el):
        return None
    if "подготов" in f"{el.get('name') or ''} {el.get('mssk_name') or ''}".lower():
        return None   # бетонная подготовка — без ухода за бетоном
    if any(str(pm).startswith(("3.6-98", "3.6-61")) for pm in existing):
        return None
    local = _local_rate("3.6-98-1") or {}
    return {"pressmark": "3.6-98-1", "unit": local.get("unitOfMeasure") or "1 м3",
            "title": local.get("title") or "Уход за бетоном", "is_resource": False}


def _layers(el: Dict[str, Any]) -> int:
    """«Техноэласт в 2 слоя» -> 2 (для площадных работ слоёв)."""
    m = re.search(r"(\d)\s*сло", f"{el.get('material') or ''} {el.get('name') or ''}".lower())
    return max(1, min(int(m.group(1)), 4)) if m else 1


# Типовые нормы армирования, кг/м3 — если в модели нет ReinforcementVolumeRatio
REBAR_NORM = {"stair": 100, "slab": 110, "wall": 90, "beam": 150, "column": 180}


def _volume(rate: Dict[str, Any], el: Dict[str, Any], q: Dict[str, float]) -> Optional[float]:
    unit = (rate.get("unit") or "").lower().replace("²", "2").replace("³", "3")
    title = kb._full_title(rate).lower()
    tokens = unit.replace(".", " ").split()
    div = _divisor(unit)
    if "м3" in unit:
        base = q["volume"]
        if not base and q["area"]:
            t = (kb.element_thickness(el) or 0) / 1000
            base = q["area"] * t if t else 0.0
    elif "м2" in unit:
        base = _formwork_area(el, q) if "опалуб" in title else q["area"] * _layers(el)
    elif tokens and tokens[-1] == "т":
        if rate["is_resource"] or "отдельных стержней" not in title:
            return None
        kg = q["rebar_kg"] or REBAR_NORM.get(_element_kind(el), 100) * q["volume"]
        base = kg / 1000.0
    elif "шт" in unit:
        if rate["is_resource"]:
            return None   # сколько штук материала — модель не знает
        base = q["count"]
    else:
        return None
    if not base:
        return None
    return round(base / div, 4)


def _costs(rate: Dict[str, Any], vol: Optional[float]):
    if vol is None:
        return None, None, None, None
    if rate["is_resource"]:
        try:
            from src.services.local_resources import resource_price
            price = resource_price(rate["pressmark"])
        except Exception:
            price = None
        if not price:
            return None, None, None, None
        mr = round(price * vol, 2)
        return None, None, mr, mr
    w = _local_rate(rate["pressmark"])
    if not w:
        return None, None, None, None
    zp = (w.get("curSalary") or 0) * vol
    em = (w.get("curOperationOfMachines") or 0) * vol
    mr = (w.get("curCostOfMaterialResources") or 0) * vol
    return round(zp, 2), round(em, 2), round(mr, 2), round(zp + em + mr, 2)


_conc_res: Optional[Dict[str, Dict[str, Any]]] = None


def _bwf(text: str):
    """Класс бетона из текста: (B, W, F), например «B30_W4_F75» -> (30.0, '4', '75')."""
    t = str(text or "")
    b = re.search(r"[ВB]\s?(\d+(?:[.,]\d+)?)", t)
    w = re.search(r"W\s?(\d+)", t)
    f = re.search(r"F\s?(\d+)", t)
    return (float(b.group(1).replace(",", ".")) if b else None,
            w.group(1) if w else None, f.group(1) if f else None)


def _concrete_resource(el: Dict[str, Any], subs) -> Optional[Dict[str, Any]]:
    """Марка бетона из перечня, ближайшая к классу элемента (B обязателен, затем W, F)."""
    global _conc_res
    if not kb._is_concrete(el) or kb.is_precast(el):
        return None
    if _conc_res is None:
        _conc_res = {}
        for s in subs:
            for w in s["works"]:
                for r in w["rates"]:
                    if r["pressmark"].startswith("1.3-1-"):
                        _conc_res.setdefault(r["pressmark"], r)
        try:
            from src.services.local_resources import concrete_mixes
            for r in concrete_mixes():
                _conc_res.setdefault(r["pressmark"], r)
        except Exception as exc:
            logger.warning(f"Справочник ресурсов недоступен: {exc}")
    b, w, f = _bwf(f"{el.get('name')} {el.get('material')}")
    if b is None and "подготов" in f"{el.get('name') or ''} {el.get('mssk_name') or ''}".lower():
        b = 7.5   # бетонная подготовка — типовой класс В7,5
    if b is None:
        return None
    best, best_score = None, None
    for r in _conc_res.values():
        rb, rw, rf = _bwf(r["title"])
        if rb is None:
            continue
        score = ((100 if rb == b else -abs(rb - b)) + (10 if w and rw == w else 0)
                 + (1 if f and rf == f else 0) + (0.5 if "гранит" in r["title"].lower() else 0))
        if best_score is None or score > best_score:
            best, best_score = r, score
    # только точное совпадение класса прочности: неверный материал хуже отсутствующего
    return best if best_score is not None and best_score >= 100 else None


def _row(rate: Dict[str, Any], el: Dict[str, Any], q: Dict[str, float], work: str) -> Dict[str, Any]:
    vol = _volume(rate, el, q)
    zp, em, mr, total = _costs(rate, vol)
    return {"pressmark": rate["pressmark"], "title": kb._full_title(rate), "unit": rate["unit"],
            "volume": vol, "zp": zp, "em": em, "mr": mr, "total": total,
            "is_resource": rate["is_resource"], "work": work}


def finalize_rows(el: Dict[str, Any], q: Dict[str, float], works, subs) -> List[Dict[str, Any]]:
    """Строки перечня для элемента: без дублей, бетон по классу элемента, уход за бетоном."""
    el = kb.enrich(el)
    rows, seen, seen_keys = [], set(), set()
    conc = _concrete_resource(el, subs)
    b_exp = _bwf(f"{el.get('name')} {el.get('material')}")[0]
    if b_exp is None and "подготов" in f"{el.get('name') or ''} {el.get('mssk_name') or ''}".lower():
        b_exp = 7.5
    for w in works:
        for r in w["rates"]:
            if (r["pressmark"].startswith("1.3-1-") and b_exp is not None
                    and _bwf(r["title"])[0] not in (None, b_exp)):
                continue   # бетонная смесь не того класса — неверный материал хуже отсутствующего
            key = (r["pressmark"], w["name"])
            if key in seen_keys or (conc and r["pressmark"].startswith("1.3-1-")):
                continue
            seen_keys.add(key)
            seen.add(r["pressmark"])
            label = "Уход за бетоном" if r["pressmark"].startswith(("3.6-98", "3.6-61")) else w["name"]
            row = _row(r, el, q, label)
            if (not r["is_resource"] and re.search(r"не модел", (w.get("ifc_class") or "").lower())
                    and not re.search(r"опалуб|арматур|стержн", row["title"].lower())):
                # работа не из модели: объём не берётся от бетона — по проекту
                row.update(volume=None, zp=None, em=None, mr=None, total=None)
            t_low = row["title"].lower()
            if row["volume"] is None and (r["is_resource"] and "арматур" in t_low or "каркас" in t_low):
                continue   # диаметры и каркасы: весь расход арматуры — в «отдельные стержни»
            if row["volume"] is None and not r["is_resource"]:
                row["work"] += " — объём по проекту (нет данных в модели)"
            if "отдельных стержней" in t_low and not q.get("rebar_kg") and row["volume"]:
                row["work"] += " — ОЦЕНКА по норме армирования"
            rows.append(row)
    if conc:
        seen.add(conc["pressmark"])
        rows.append(_row(conc, el, q, "Бетон (по классу элемента)"))
    cure = _curing_rate(el, seen)
    if cure:
        rows.append(_row(cure, el, q, "Уход за бетоном (добавлено правилом)"))
    return rows


# ---------------------------------------------------------------- основной шаг

def build_final_from_perechen(tables_json_path: str, run_dir: str) -> str:
    data = json.load(open(tables_json_path, encoding="utf-8"))
    elements = data.get("elements") or []
    groups = json.load(open(os.path.join(run_dir, "filtered_elements_grouped_AR.json"),
                            encoding="utf-8"))
    height = kb.building_height_for_run(tables_json_path, data)
    subs = kb.load_perechen()
    per_el = _element_quantities(run_dir)
    logger.info(f"Подбор по перечню: высота {height} м, подразделов {len(subs)}, "
                f"элементов {len(elements)}, объёмов по элементам {len(per_el)}")

    results: List[Dict[str, Any]] = []
    for node, path in _leaves(groups if isinstance(groups, list) else [groups]):
        idx = sorted(i for i in (node.get("indices") or [])
                     if isinstance(i, int) and 0 <= i < len(elements))
        if not idx:
            continue
        part = _part_from_path(path)

        # делим листовую группу по типу элемента (имя без ID + материал)
        buckets: Dict[Tuple[str, str], List[int]] = {}
        for i in idx:
            el = elements[i].get("element", {})
            buckets.setdefault((kb.type_key(el.get("name")), str(el.get("material") or "")), []).append(i)
        single = len(buckets) == 1

        for (tkey, material), ids in buckets.items():
            el = elements[ids[0]].get("element", {})
            q = {"volume": 0.0, "area": 0.0, "rebar_kg": 0.0, "fw_hint": 0.0, "count": len(ids)}
            if per_el:
                for i in ids:
                    if i < len(per_el):
                        for k in ("volume", "area", "rebar_kg", "fw_hint"):
                            q[k] += per_el[i].get(k, 0.0)
            if single:  # итоги группировки надёжнее, если группа однородна
                leaf = {"volume": float(node.get("total_volume") or 0), "area": _leaf_area(node),
                        "rebar_kg": float(node.get("total_reinforcement") or 0)}
                for k, v in leaf.items():
                    if v > 0:
                        q[k] = v
            q = {k: round(v, 4) if isinstance(v, float) else v for k, v in q.items()}

            try:
                res = kb.pick_for_element(el, part, subs, height)
            except Exception as exc:
                logger.error(f"Подбор по перечню не удался ({tkey}): {exc}", exc_info=True)
                res = {"status": "error", "subsection": None, "works": [], "reason": str(exc)}

            rows = finalize_rows(el, q, res["works"], subs)
            results.append({"part": part, "element": tkey, "material": material,
                            "el": {k: el.get(k) for k in ("name", "ifc_class", "mssk_code", "material", "predefined_type")},
                            "count": len(ids), "path": list(path), "quantity": q,
                            "status": res["status"],
                            "subsection": (res["subsection"] or {}).get("title"),
                            "reason": res["reason"], "rows": rows})
            logger.info(f"[{part}] {tkey} ×{len(ids)}: {res['status']}, строк {len(rows)}")

    results.sort(key=lambda g: PART_ORDER.index(g["part"]) if g["part"] in PART_ORDER else 9)
    with open(os.path.join(run_dir, FINAL_JSON), "w", encoding="utf-8") as fh:
        json.dump({"source": "perechen_kr", "building_height_m": height, "groups": results},
                  fh, ensure_ascii=False, indent=2)
    xlsx_path = os.path.join(run_dir, FINAL_XLSX)
    _write_xlsx(results, xlsx_path)
    logger.info(f"Итоговый перечень по перечню сметчиков: {xlsx_path} ({len(results)} групп)")
    return xlsx_path


def _write_xlsx(results: List[Dict[str, Any]], path: str) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "Данные"
    ws.append(HEADERS)
    for c in ws[1]:
        c.font = Font(bold=True)
    part_fill = PatternFill("solid", fgColor="A6A6A6")
    group_fill = PatternFill("solid", fgColor="D9D9D9")

    def styled(values, fill=None, bold=False):
        ws.append(values)
        for c in ws[ws.max_row]:
            if fill:
                c.fill = fill
            if bold:
                c.font = Font(bold=True)

    grand = [0.0, 0.0, 0.0, 0.0]
    for part in PART_ORDER:
        items = [g for g in results if g["part"] == part]
        if not items:
            continue
        styled(["", PART_TITLES[part], "", "", "", "", "", ""], part_fill, True)
        ptot = [0.0, 0.0, 0.0, 0.0]
        for g in items:
            head = g["element"] + (f" ({g['material']})" if g["material"] else "") + f"  Кол-во: {g['count']}"
            if g["subsection"]:
                head += f"  [{g['subsection']}]"
            styled(["", head, "", "", "", "", "", ""], group_fill, True)
            if not g["rows"]:
                ws.append(["", f"Работы не подобраны: {str(g['reason'])[:250]}"])
            for r in g["rows"]:
                money = [r["zp"], r["em"], r["mr"], r["total"]]
                vol = r["volume"] if r["volume"] is not None else ("" if r["is_resource"] else "по проекту")
                title = r["title"]
                if [x["pressmark"] for x in g["rows"]].count(r["pressmark"]) > 1:
                    title += f"  [{r['work'].split(' — ')[0][:60]}]"
                ws.append([r["pressmark"], title, r["unit"], vol, *money])
                for i, v in enumerate(money):
                    if v:
                        ptot[i] += v
            ws.append([])
        styled(["", PART_TOTALS[part], "", "", *[round(v, 2) for v in ptot]], None, True)
        ws.append([])
        grand = [a + b for a, b in zip(grand, ptot)]
    styled(["", "ИТОГО:", "", "", *[round(v, 2) for v in grand]], None, True)

    for col, width in zip("ABCDEFGH", (14, 90, 12, 14, 14, 14, 14, 16)):
        ws.column_dimensions[col].width = width
    ws.auto_filter.ref = f"A1:H{ws.max_row}"
    wb.save(path)



# =====================================================================
#  ДОРАБОТКИ ПО ЗАМЕЧАНИЯМ (обёртки поверх базовой логики)
# =====================================================================

# (6) Бетон: точный класс прочности, по W и F — ближайшая марка сверху
def _concrete_resource(el, subs):
    global _conc_res
    if not kb._is_concrete(el) or kb.is_precast(el):
        return None
    if _conc_res is None:
        _conc_res = {}
        for sb in subs:
            for w in sb["works"]:
                for r in w["rates"]:
                    if r["pressmark"].startswith("1.3-1-"):
                        _conc_res.setdefault(r["pressmark"], r)
        try:
            from src.services.local_resources import concrete_mixes
            for r in concrete_mixes():
                _conc_res.setdefault(r["pressmark"], r)
        except Exception as exc:
            logger.warning(f"Справочник ресурсов недоступен: {exc}")
    b, w, f = _bwf(f"{el.get('name')} {el.get('material')}")
    if b is None and "подготов" in f"{el.get('name') or ''} {el.get('mssk_name') or ''}".lower():
        b = 7.5
    if b is None:
        return None

    def grade(req, got, exact, above):
        if not req or not got:
            return 0.0
        rq, gt = int(req), int(got)
        if gt == rq:
            return exact
        if gt > rq:
            return above - (gt - rq) * above / 400
        return -exact

    best, best_score = None, None
    for r in _conc_res.values():
        rb, rw, rf = _bwf(r["title"])
        if rb != b:
            continue
        gw, gf = grade(w, rw, 10, 6), grade(f, rf, 1, 0.6)
        score = gw + gf + (0.5 if "гранит" in r["title"].lower() else 0) + (3 if gw >= 0 and gf >= 0 and _has_price(r["pressmark"]) else 0)
        if best_score is None or score > best_score:
            best, best_score = r, score
    return best


# (2) Огрунтовка: площадь поверхности один раз, без умножения на слои
_volume_base = _volume


def _volume(rate, el, q):
    v = _volume_base(rate, el, q)
    title = kb._full_title(rate).lower()
    if v and re.search(r"огрунт|праймер", title) and _layers(el) > 1:
        return round(v / _layers(el), 4)
    return v


# (5) Лестницы: площадке — не работы маршей, маршу — не работы площадок
def _dedupe_stair_works(el, works):
    e = kb.enrich(el)
    text = f"{e.get('name') or ''} {e.get('mssk_name') or ''}".lower()
    ifc = str(e.get("ifc_class") or "").lower()
    landing = "площадк" in text or ("лестн" in text and ifc == "ifcslab")
    flight = "марш" in text or ifc in ("ifcstairflight", "ifcstair")
    if landing == flight:
        return works
    drop = "марш" if landing else "площад"

    def txt(w):
        return (w["name"] + " " + " ".join(r["title"] for r in w["rates"] if not r["is_resource"])).lower()

    kept = [w for w in works if drop not in txt(w)]
    return kept or works


# (3) Цена ресурса: текущая, иначе базисная × индекс группы (с пометкой)
_idx_cache = {}


def _estimated_price(code):
    from src.services.local_resources import load
    res = load()
    d = res.get(code)
    if not d:
        return None, "цены нет в справочнике"
    if d.get("price"):
        return d["price"], ""
    if not d.get("price_base"):
        return None, "цены нет в справочнике"
    grp = "-".join(code.split("-")[:2])
    if grp not in _idx_cache:
        ratios = sorted(x["price"] / x["price_base"] for k, x in res.items()
                        if k.startswith(grp + "-") and x.get("price") and x.get("price_base"))
        _idx_cache[grp] = ratios[len(ratios) // 2] if ratios else None
    k = _idx_cache[grp]
    if not k:
        return None, "текущей цены нет в справочнике"
    return round(d["price_base"] * k, 2), f"цена оценена: базисная × индекс группы {k:.2f}"


# (1) Материал утеплителя, если его нет в строках (qwen выбирает из справочника)
def _insulation_material_row(el, q):
    from src.services.local_resources import load
    key = "mat|" + kb.type_key(el.get("name")) + "|" + str(el.get("material") or "")
    cache = kb._load_cache()
    if key in cache:
        code = cache[key]
    else:
        mat = f"{el.get('material') or ''} {el.get('name') or ''}".lower()
        skip = {"базовая", "стена", "утеплитель", "перекрытие", "изоляция"}
        stems = list(dict.fromkeys(w[:7] for w in re.findall(r"[а-яё]{5,}", mat) if w not in skip))
        cands = []
        for d in load().values():
            t = d["title"].lower()
            sc = sum(1 for st in stems if st in t)
            unit = d["unit"].lower().replace("³", "3").replace("²", "2")
            if sc and unit in ("м3", "м2"):
                cands.append((sc, d))
        cands.sort(key=lambda x: (-x[0], len(x[1]["title"])))
        cands = [d for _, d in cands[:25]]
        code = None
        if cands:
            lines = "\n".join(f"{i}: {d['pressmark']} | {d['unit']} | {d['title'][:150]}" for i, d in enumerate(cands))
            ans = kb._ask_llm(
                "Ты опытный сметчик. Выбери из списка материал (ресурс ТСН), соответствующий "
                "утеплителю элемента BIM-модели. Выбери наиболее близкий по виду материала (плиты того же типа), даже если марка или плотность отличаются; null — только если материалов такого вида нет вообще.\n\n"
                f"Элемент: {el.get('name')}; материал: {el.get('material')}\n\n"
                f"Материалы (номер: шифр | ед. | наименование):\n{lines}\n\n"
                'Ответь строго JSON: {"id": <номер или null>}')
            i = ans.get("id")
            if isinstance(i, int) and 0 <= i < len(cands):
                code = cands[i]["pressmark"]
        cache[key] = code
        kb._save_cache(cache)
    if not code or code not in load():
        return None
    d = load()[code]
    rate = {"pressmark": code, "title": d["title"], "unit": d["unit"], "is_resource": True, "formula": ""}
    return _row(rate, el, q, "Материал утеплителя (подобран по справочнику)")


_finalize_rows_base = finalize_rows


def finalize_rows(el, q, works, subs):
    el = kb.enrich(el)
    rows = _finalize_rows_base(el, q, _dedupe_stair_works(el, works), subs)
    codes = [r["pressmark"] for r in rows]

    # (2) огрунтовка + мастика под наплавляемую рулонную гидроизоляцию
    if "3.8-2-11" in codes and not any(c.startswith("3.8-48-") for c in codes):
        from src.services.local_resources import load
        loc = _local_rate("3.8-48-1") or {}
        primer = {"pressmark": "3.8-48-1", "unit": loc.get("unitOfMeasure") or "100 м2",
                  "title": loc.get("title") or "Огрунтовка поверхности стен, фундаментов",
                  "is_resource": False, "formula": ""}
        new = [_row(primer, el, q, "Огрунтовка под наплавляемую гидроизоляцию (добавлено правилом)")]
        mast = load().get("1.1-1-1693")
        if mast:
            new.append(_row({"pressmark": "1.1-1-1693", "unit": mast["unit"], "title": mast["title"],
                             "is_resource": True, "formula": ""}, el, q, new[0]["work"]))
        i = codes.index("3.8-2-11")
        rows[i:i] = new

    # (1) материал утеплителя
    if kb._layer_kind(el) == "insulation" and rows and not any(r["is_resource"] for r in rows):
        mrow = _insulation_material_row(el, q)
        if mrow:
            rows.append(mrow)

    # (4) количество ресурса без объёма — по норме расхода в своей расценке
    try:
        from src.services.local_resources import norm
    except Exception:
        norm = None
    last = None
    for r in rows:
        if not r["is_resource"]:
            last = r if r["volume"] is not None else None
            continue
        if r["volume"] is None and last is not None and norm:
            n = norm(last["pressmark"], r["pressmark"]) or _norm_similar(last["pressmark"], r["pressmark"])
            if n:
                r["volume"] = round(last["volume"] * n, 4)
                r["work"] += " — по норме расхода"

    # (3) цена ресурса, если не найдена
    for r in rows:
        if r["is_resource"] and r["volume"] is not None and r["total"] is None:
            price, note = _estimated_price(r["pressmark"])
            if price:
                r["mr"] = r["total"] = round(price * r["volume"], 2)
            if note:
                r["title"] += f"  ({note})"
        elif r["is_resource"] and r["volume"] is None:
            r["title"] += "  (количество по проекту)"
    return rows



def _has_price(code):
    try:
        from src.services.local_resources import load
        return bool((load().get(code) or {}).get("price"))
    except Exception:
        return False


def _norm_similar(rate, res):
    """Норма похожего ресурса (по первому слову названия) в этой расценке или её таблице."""
    from src.services import local_resources as lr
    lr.norm(rate, res)
    resd = lr.load()
    words = re.findall(r"[а-яё]{4,}", (resd.get(res) or {}).get("title", "").lower())
    if not words:
        return None
    stem = words[0][:6]
    table = rate.rsplit("-", 1)[0] + "-"
    for exact in (True, False):
        for k, v in lr._norms.items():
            r_code, res_code = k.split("|", 1)
            if (r_code == rate if exact else r_code.startswith(table)) and \
                    stem in (resd.get(res_code) or {}).get("title", "").lower():
                return v
    return None



# ===== Константы проекта (из интерфейса / ПОС): сезон ухода за бетоном =====
PROJECT_CONSTANTS = {}


def _cold_season():
    vals = [str(v).lower() for k, v in PROJECT_CONSTANTS.items()
            if any(x in str(k).lower() for x in ("curing", "season", "сезон"))]
    vals.append(os.getenv("CURING_SEASON", "").lower())
    return any(("холод" in v or "зим" in v or "cold" in v) for v in vals)


_finalize_rows_v2 = finalize_rows


def finalize_rows(el, q, works, subs):
    rows = _finalize_rows_v2(el, q, works, subs)
    if not _cold_season():
        return rows
    for i, r in enumerate(rows):
        if r["pressmark"].startswith("3.6-98-"):
            loc = _local_rate("3.6-61-1") or {}
            rate = {"pressmark": "3.6-61-1", "unit": loc.get("unitOfMeasure") or "100 м2 поверхности",
                    "title": loc.get("title") or "Уход за бетоном с применением тепловлагозащитного покрытия",
                    "is_resource": False, "formula": ""}
            rows[i] = _row(rate, el, q, "Уход за бетоном в холодный период (тепловлагозащита, по ПОС)")
    return rows


_build_base = build_final_from_perechen


def build_final_from_perechen(tables_json_path, run_dir):
    global PROJECT_CONSTANTS
    try:
        d = json.load(open(tables_json_path, encoding="utf-8"))
        PROJECT_CONSTANTS = dict(((d.get("elements") or [{}])[0].get("applied_constants")) or {})
        PROJECT_CONSTANTS.update(d.get("global_constants") or {})
    except Exception:
        PROJECT_CONSTANTS = {}
    logger.info(f"Константы проекта для подбора: {PROJECT_CONSTANTS}")
    return _build_base(tables_json_path, run_dir)


# ==== PD_FULL: подача бетона по ПОС (бетононасос вместо кран-бадьи) ====
import os, re, json
PD_FULL = {}
_PUMP_CANDIDATES = ["3.6-73-4", "3.6-73-5", "3.6-73-6", "3.6-110-1", "3.6-110-2", "3.6-110-3",
                    "3.6-128-1", "3.6-128-2", "3.6-128-3", "3.6-128-4", "3.6-128-5"]
_KINDS = ["фундаментн", "ростверк", "ленточн", "стен", "колонн", "плит", "перекрыт", "марш", "площад", "балочн"]


def _pd_load(run_dir):
    p = os.path.abspath(run_dir)
    for _ in range(4):
        fp = os.path.join(p, "pd_full.json")
        if os.path.isfile(fp):
            try:
                return {r["id"]: r for r in json.load(open(fp, encoding="utf-8"))}
            except Exception:
                return {}
        p = os.path.dirname(p)
    return {}


def _supply_is_pump(title_low):
    """True — бетононасос, False — кран-бадья, None — нет данных / расчёт без ПОС/ПЗ."""
    if not PROJECT_CONSTANTS:
        return None                       # выбран режим «без учёта ПОС и ПЗ»
    under = "подземн" in title_low and "надземн" not in title_low
    rec = PD_FULL.get("concrete_supply_under" if under else "concrete_supply_above") or {}
    v = f"{rec.get('value') or ''} {rec.get('quote') or ''}".lower()
    if "насос" in v:
        return True
    if "бадь" in v:
        return False
    return any("бетононасос" in str(x).lower() for x in PROJECT_CONSTANTS.values()) or None


def _height_marks(t):
    return set(re.findall(r"высот\w* здани\w* (?:до|более|от)[^,;]*?\d+", t))


_TREE_TITLES = None


def _table_title(code):
    """'3.6-97-3' -> заголовок таблицы 3.6-97 из tree_work_compact.json"""
    global _TREE_TITLES
    if _TREE_TITLES is None:
        _TREE_TITLES = {}
        try:
            d = json.load(open("/app/data/tree_work_compact.json", encoding="utf-8"))
            def walk(o):
                if isinstance(o, dict):
                    for k, v in o.items():
                        for x in (k, v if isinstance(v, str) else ""):
                            m = re.match(r"Таблица (\d+\.\d+-\d+)\.", x)
                            if m:
                                _TREE_TITLES.setdefault(m.group(1), x.lower())
                        walk(v)
                elif isinstance(o, list):
                    for v in o:
                        walk(v)
            walk(d)
        except Exception:
            pass
    return _TREE_TITLES.get(code.rsplit("-", 1)[0], "")


def _hrange(t):
    m = re.search(r"высоте здания (?:более |от )?(\d+) до (\d+)", t)
    return (m.group(1), m.group(2)) if m else None


def _pump_rate_for(old_title, old_code=""):
    t = re.sub(r"между осями колонн или стен", "", old_title.lower())
    kinds = [k for k in _KINDS if k in t]
    tt_old = _table_title(old_code) if old_code else ""
    under = "подземн" in (tt_old or t) and "надземн" not in (tt_old or t)
    h_old = _hrange(tt_old) or _hrange(t)
    best, best_s = None, 0
    for code in _PUMP_CANDIDATES:
        loc = _local_rate(code) or {}
        ct = (loc.get("title") or "").lower()
        if not ct:
            continue
        tt = _table_title(code)
        if under:
            if "надземн" in tt or "надземн" in ct or not code.startswith("3.6-73"):
                continue                      # подземная часть — только фундаменты автобетононасосом
        else:
            if code.startswith("3.6-73"):
                continue
            if _hrange(tt) != h_old:
                continue                      # насосная таблица для другой высоты здания
        ck = [k for k in _KINDS if k in ct]
        if not kinds or not set(kinds) & set(ck):
            continue
        if "более 300" in t and "до 300" in ct:
            continue
        s = len(set(kinds) & set(ck)) * 10
        if s > best_s:
            best, best_s = (code, loc), s
    return best


_fin_before_pump = finalize_rows


def finalize_rows(el, q, works, subs):
    rows = _fin_before_pump(el, q, works, subs)
    for i, r in enumerate(rows):
        title = (_local_rate(str(r.get("pressmark") or "")) or {}).get("title") or " ".join(str(v) for v in r.values() if isinstance(v, str))
        if "кран-бадья" not in title.lower():
            continue
        if not _supply_is_pump(title.lower()):
            continue
        pr = _pump_rate_for(title, str(r.get('pressmark') or ''))
        if not pr:
            continue
        code, loc = pr
        rate = {"pressmark": code, "unit": loc.get("unitOfMeasure") or "100 м3",
                "title": loc.get("title") or "Бетонирование бетононасосом", "is_resource": False, "formula": ""}
        try:
            rows[i] = _row(rate, el, q, "Бетонирование бетононасосом (по ПОС)")
            logger.info(f"ПОС: {r.get('pressmark')} (кран-бадья) -> {code} (бетононасос)")
        except Exception as e:
            logger.warning(f"замена на бетононасос не удалась: {e}")
    return rows


_build_before_pd = build_final_from_perechen


def build_final_from_perechen(tables_json_path, run_dir):
    global PD_FULL
    PD_FULL = _pd_load(run_dir) or _pd_load(os.path.dirname(os.path.abspath(tables_json_path)))
    logger.info(f"ПОС/ПЗ (полный разбор): {sum(1 for r in PD_FULL.values() if r.get('value'))} значений")
    return _build_before_pd(tables_json_path, run_dir)


# ==== ANALYTIC: краткая аналитическая справка по расчёту ====
_AN_STATS = []


def _an_text(r):
    return " ".join(str(v) for v in r.values() if isinstance(v, str))


_fin_before_an = finalize_rows


def finalize_rows(el, q, works, subs):
    rows = _fin_before_an(el, q, works, subs)
    try:
        _AN_STATS.append({
            "el": str(el.get("name") or el.get("type_name") or "элемент")[:80],
            "rows": [(str(r.get("pressmark") or ""), _an_text(r).lower()) for r in rows if isinstance(r, dict)],
        })
    except Exception:
        pass
    return rows


def _an_build(run_dir):
    st = _AN_STATS
    nrows = sum(len(s["rows"]) for s in st)
    ok, doubt = [], []

    def els(pred, limit=4):
        names = []
        for s in st:
            if any(pred(pm, t) for pm, t in s["rows"]) and s["el"] not in names:
                names.append(s["el"])
        return names

    use_pd = bool(PROJECT_CONSTANTS)
    if not use_pd:
        ok.append("Расчёт выполнен БЕЗ учёта ПОС и ПЗ — только по данным модели и значениям по умолчанию.")
    else:
        pc = {str(k).lower(): str(v) for k, v in PROJECT_CONSTANTS.items()}
        h = next((v for k, v in pc.items() if "building" in k and "height" in k), "") or \
            next((v for k, v in pc.items() if "height" in k and "floor" not in k and "этаж" not in k), "")
        if h:
            hv = (PD_FULL.get("height") or {}).get("value", "")
            ok.append(f"Высота здания {h} м — расценки взяты из таблиц ТСН для этой высоты" + (f" (ПОС: {hv[:60]})." if hv else "."))
    n_w = sum(1 for s in st for pm, t in s["rows"] if pm.startswith("3.6-61"))
    if n_w:
        ok.append(f"Холодный период по ПОС — уход за бетоном заменён на тепловлагозащиту (3.6-61-1): {n_w} поз.")
    n_pump = sum(1 for s in st for pm, t in s["rows"] if "бетононасосом (по пос)" in t)
    n_bucket = sum(1 for s in st for pm, t in s["rows"] if "кран-бадья" in t)
    sup = " ".join(str((PD_FULL.get(k) or {}).get("value", "")) for k in ("concrete_supply_under", "concrete_supply_above")).lower()
    if n_pump:
        ok.append(f"Подача бетона насосом по ПОС: {n_pump} поз. переведены с «кран-бадьи» на бетононасос.")
    if use_pd and "насос" in sup and n_bucket:
        doubt.append({"text": f"В ПОС подача бетона насосом, но в ТСН для зданий этой высоты расценок на насос нет — "
                              f"оставлена «кран-бадья» ({n_bucket} поз.). Решение за сметчиком.", "items": []})
    wm = (PD_FULL.get("winter_method") or {}).get("value", "")
    if use_pd and wm:
        doubt.append({"text": f"ПОС: зимний прогрев бетона ({wm[:70]}). Отдельной расценки в ТСН нет — "
                              "учесть зимним удорожанием.", "items": []})
    ok.append(f"Подобрано {nrows} позиций для {len(st)} типов элементов. Бетон — по классу из модели (B, W, F).")

    np_ = els(lambda pm, t: "по проекту" in t)
    if np_:
        doubt.append({"text": f"Объём не определяется по модели — взять из проекта ({len(np_)} эл.): закладные, "
                              "материалы без количества и т.п.", "items": np_[:4]})
    nc = els(lambda pm, t: "цены нет" in t)
    if nc:
        doubt.append({"text": f"Нет цены в справочнике — стоимость не посчитана ({len(nc)} эл.).", "items": nc[:4]})
    empty = [s["el"] for s in st if not s["rows"]]
    if empty:
        doubt.append({"text": f"Работы не подобраны ({len(empty)} эл.) — проверить вручную.", "items": empty[:4]})
    if use_pd and PD_FULL:
        miss = [r.get("title", "").split(":")[0].split("(")[0].strip() for k, r in PD_FULL.items()
                if not r.get("value") and k in ("concrete_class", "waterproof", "insulation", "prep", "rebar", "cover")]
        if miss:
            doubt.append({"text": "В ПОС/ПЗ нет данных: " + ", ".join(miss) + " — взято из модели (нужен раздел КР для проверки).",
                          "items": []})
    res = {"ok": ok, "doubt": doubt, "rows": nrows, "elements": len(st), "with_pd": use_pd}
    try:
        json.dump(res, open(os.path.join(run_dir, "analytic.json"), "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    except Exception as e:
        logger.warning(f"аналитическая справка не записана: {e}")
    return res


_build_before_an = build_final_from_perechen


def build_final_from_perechen(tables_json_path, run_dir):
    _AN_STATS.clear()
    out = _build_before_an(tables_json_path, run_dir)
    try:
        a = _an_build(run_dir)
        logger.info(f"Аналитическая справка: {len(a['ok'])} учтено, {len(a['doubt'])} сомнений")
    except Exception as e:
        logger.warning(f"аналитическая справка: {e}")
    return out


# ==== ANALYTIC_V2: понятная справка для пользователя ====
def _an_short(name):
    parts = [p for p in str(name).split(":") if p and not p.strip().isdigit()]
    n = parts[1] if len(parts) >= 2 else (parts[0] if parts else str(name))
    n = re.sub(r"[_]+", " ", n)
    n = re.sub(r"\s+", " ", n).strip()
    return n[:60]


def _an_ref(pid):
    r = PD_FULL.get(pid) or {}
    if not r.get("value"):
        return ""
    return f"{r.get('src') or 'док.'}" + (f", стр. {r.get('page')}" if r.get("page") else "")


def _an_build(run_dir):
    st = _AN_STATS
    nrows = sum(len(s["rows"]) for s in st)
    use_pd = bool(PROJECT_CONSTANTS)
    ok, doubt, info = [], [], []

    def els(pred):
        names = []
        for s in st:
            if any(pred(pm, t) for pm, t in s["rows"]):
                n = _an_short(s["el"])
                if n not in names:
                    names.append(n)
        return names

    def cnt(pred):
        return sum(1 for s in st for pm, t in s["rows"] if pred(pm, t))

    def ex(names, k=3):
        return "; ".join(names[:k]) + (" и др" if len(names) > k else "")

    # --- учтено
    if not use_pd:
        ok.append("ПОС и ПЗ НЕ учитывались — расчёт только по модели и значениям по умолчанию.")
    else:
        pc = {str(k).lower(): str(v) for k, v in PROJECT_CONSTANTS.items()}
        h = next((v for k, v in pc.items() if "building" in k and "height" in k), "") or \
            next((v for k, v in pc.items() if "height" in k and "floor" not in k), "")
        hv = re.search(r"\d+[.,]?\d*\s*м", (PD_FULL.get("height") or {}).get("value", ""))
        if h:
            ok.append(f"Высота здания{' ' + hv.group(0) if hv else ''} ({_an_ref('height') or 'ПОС'}) → расценки для зданий {h} м.")
        n_w = cnt(lambda pm, t: pm.startswith("3.6-61"))
        if n_w:
            ok.append(f"Холодный период (ПОС) → уход за бетоном с тепловлагозащитой: {n_w} поз.")
        n_pump = cnt(lambda pm, t: "бетононасосом (по пос)" in t)
        if n_pump:
            ok.append(f"Подача бетона насосом ({_an_ref('concrete_supply_under') or 'ПОС'}) → расценки на бетононасос: {n_pump} поз.")
    ok.append(f"Класс бетона (B, W, F), толщины и объёмы — из модели.")

    # --- проверить
    if use_pd:
        sup = " ".join(str((PD_FULL.get(k) or {}).get("value", "")) for k in ("concrete_supply_under", "concrete_supply_above")).lower()
        n_b = cnt(lambda pm, t: "кран-бадья" in t)
        if "насос" in sup and n_b:
            doubt.append(f"Подача бетона: в ПОС насос ({_an_ref('concrete_supply_above') or 'ПОС'}), а в ТСН для зданий этой высоты "
                         f"есть только «кран-бадья» — она и стоит ({n_b} поз.). Если нужен насос — решает сметчик.")
        if (PD_FULL.get("winter_method") or {}).get("value"):
            doubt.append(f"Зимний прогрев бетона ({_an_ref('winter_method')}): отдельной расценки нет — учесть зимним удорожанием.")
    np_ = els(lambda pm, t: "по проекту" in t)
    if np_:
        doubt.append(f"Объём не определить по модели — взять из проекта ({cnt(lambda pm, t: 'по проекту' in t)} поз., "
                     f"{len(np_)} эл.): {ex(np_)}.")
    nc = els(lambda pm, t: "цены нет" in t)
    if nc:
        doubt.append(f"Нет цены в справочнике — стоимость не посчитана ({cnt(lambda pm, t: 'цены нет' in t)} поз.): {ex(nc)}.")
    empty = [_an_short(s["el"]) for s in st if not s["rows"]]
    if empty:
        doubt.append(f"Работы не подобраны ({len(empty)} эл.) — добавить вручную: {ex(empty)}.")
    if use_pd and PD_FULL:
        miss = [(PD_FULL[k].get("title") or "").split(":")[0].split("(")[0].strip()
                for k in ("concrete_class", "waterproof", "insulation", "prep", "cover")
                if k in PD_FULL and not PD_FULL[k].get("value")]
        if miss:
            doubt.append("В ПОС/ПЗ нет: " + ", ".join(m[:1].lower() + m[1:] for m in miss) + " — взято из модели. Для проверки нужен раздел КР.")

    # --- не входит в перечень
    if use_pd and PD_FULL:
        other = [(r.get("title") or "").split("(")[0].split(":")[0].strip() + f" ({_an_ref(k)})"
                 for k, r in PD_FULL.items() if r.get("use") == "other" and r.get("value")]
        if other:
            info.append("По ПОС нужны работы других разделов сметы (в этот перечень не входят): " + "; ".join(other) + ".")

    res = {"ok": ok, "doubt": [{"text": d, "items": []} for d in doubt], "info": info,
           "rows": nrows, "elements": len(st), "with_pd": use_pd}
    try:
        json.dump(res, open(os.path.join(run_dir, "analytic.json"), "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    except Exception as e:
        logger.warning(f"аналитическая справка не записана: {e}")
    return res


# ==== TSN_COEFFS: поправки из сборников ТСН (да -> в перечень, неясно -> в справку) ====
_TSN_HITS = []


def _tsn_extra():
    ex = {}
    for k, v in (PROJECT_CONSTANTS or {}).items():
        ex[str(k)] = str(v)[:60]
    for k, r in (PD_FULL or {}).items():
        if r.get("value"):
            ex["ПОС/ПЗ: " + (r.get("title") or k).split("(")[0].strip()] = str(r["value"])[:90]
    return ex


_fin_before_tsn = finalize_rows


def finalize_rows(el, q, works, subs):
    rows = _fin_before_tsn(el, q, works, subs)
    try:
        from src.services import tsn_coeffs
    except Exception:
        return rows
    extra = _tsn_extra()
    for r in rows:
        if not isinstance(r, dict):
            continue
        pm = str(r.get("pressmark") or "")
        title = (_local_rate(pm) or {}).get("title") or ""
        try:
            res = tsn_coeffs.decide(pm, title, el, q, extra)
        except Exception as e:
            logger.warning(f"ТСН-поправки {pm}: {e}")
            continue
        for x in res:
            if x["decision"] == "нет":
                continue
            _TSN_HITS.append(dict(x, pressmark=pm, el=str(el.get("name") or "элемент")))
            if x["decision"] == "да":
                note = f" [поправка {x['doc']} п.{x['clause']}: {x['coef']}]"
                r.setdefault("tsn_coeffs", []).append(x)
                head = title[:25].lower()
                for k, v in r.items():
                    if isinstance(v, str) and head and v.lower().startswith(head) and note not in v:
                        r[k] = v + note
                        break
    return rows


_an_build_before_tsn = _an_build


def _an_build(run_dir):
    res = _an_build_before_tsn(run_dir)
    try:
        from src.services import tsn_coeffs
        docs = tsn_coeffs.docs()
    except Exception:
        docs = []
    if not docs:
        res.setdefault("info", []).append("Сборники ТСН с поправками не загружены — поправочные коэффициенты не проверялись.")
    else:
        groups = {}
        for h in _TSN_HITS:
            k = (h["decision"], h["doc"], h["clause"], h["coef"], h["condition"])
            groups.setdefault(k, []).append(h)
        n_yes = n_q = 0
        for (dec, doc, cl, coef, cond), hs in groups.items():
            names = []
            for h in hs:
                n = _an_short(h["el"])
                if n not in names:
                    names.append(n)
            why = hs[0].get("why") or ""
            cond_s = cond if len(cond) <= 80 else cond[:78] + "…"
            if dec == "да":
                n_yes += 1
                res["ok"].append(f"Поправка {doc} п.{cl} «{cond_s}» → {coef}. Применена к: {'; '.join(names[:3])}"
                                 + (" и др" if len(names) > 3 else "") + (f" — потому что {why}" if why else "") + ".")
            else:
                n_q += 1
                res["doubt"].append({"text": f"Возможна поправка {doc} п.{cl} «{cond_s}» → {coef} для: "
                                             f"{'; '.join(names[:3])}" + (" и др" if len(names) > 3 else "")
                                             + (f". {why[:1].upper() + why[1:]}" if why else "")
                                             + ". Проверьте и примените, если условие выполняется.", "items": []})
        lst = ", ".join(f"{d['code']} {d['title']}".strip() for d in docs)
        if not n_yes and not n_q:
            res.setdefault("info", []).append(f"Поправки проверены по сборникам: {lst}. К подобранным работам они не относятся.")
        else:
            res.setdefault("info", []).append(f"Поправки проверены по сборникам: {lst}.")
    try:
        json.dump(res, open(os.path.join(run_dir, "analytic.json"), "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    except Exception:
        pass
    return res


_build_before_tsn = build_final_from_perechen


def build_final_from_perechen(tables_json_path, run_dir):
    _TSN_HITS.clear()
    return _build_before_tsn(tables_json_path, run_dir)


# ==== TSN_MODE: учитывать сборники ТСН или нет (переключатель на странице) ====
TSN_ON = True


def _tsn_mode_load(run_dir):
    p = os.path.abspath(run_dir)
    for _ in range(4):
        fp = os.path.join(p, "tsn_mode.json")
        if os.path.isfile(fp):
            try:
                return json.load(open(fp, encoding="utf-8")).get("mode", "on") != "off"
            except Exception:
                return True
        p = os.path.dirname(p)
    return True


_fin_with_tsn = finalize_rows


def finalize_rows(el, q, works, subs):
    if TSN_ON:
        return _fin_with_tsn(el, q, works, subs)
    return _fin_before_tsn(el, q, works, subs)


_an_with_tsn = _an_build


def _an_build(run_dir):
    if TSN_ON:
        res = _an_with_tsn(run_dir)
    else:
        res = _an_build_before_tsn(run_dir)
        res.setdefault("info", []).append(
            "Сборники ТСН НЕ учитывались (выбрано «без учёта сборников») — поправочные коэффициенты не проверялись.")
    res["with_tsn"] = TSN_ON
    try:
        json.dump(res, open(os.path.join(run_dir, "analytic.json"), "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    except Exception:
        pass
    return res


_build_before_tsnmode = build_final_from_perechen


def build_final_from_perechen(tables_json_path, run_dir):
    global TSN_ON
    TSN_ON = _tsn_mode_load(run_dir) and _tsn_mode_load(os.path.dirname(os.path.abspath(tables_json_path)))
    logger.info(f"Сборники ТСН: {'учитываются' if TSN_ON else 'НЕ учитываются'}")
    return _build_before_tsnmode(tables_json_path, run_dir)


# ==== AN_GROUPS: справка по источникам (ПОС / ПЗ / сборники ТСН / модель) ====
_AN_SRC = ["Из ПОС", "Из ПЗ", "Из ПОС и ПЗ", "Из сборников ТСН", "Из модели и справочника цен"]


def _an_src(t):
    if "ТСН-" in t or "сборник" in t.lower():
        return "Из сборников ТСН"
    if "ПОС/ПЗ" in t or ("ПОС" in t and "ПЗ" in t):
        return "Из ПОС и ПЗ"
    if "ПОС" in t:
        return "Из ПОС"
    if "ПЗ" in t:
        return "Из ПЗ"
    return "Из модели и справочника цен"


_an_before_groups = _an_build


def _an_build(run_dir):
    res = _an_before_groups(run_dir)
    ok = {k: [] for k in _AN_SRC}
    db = {k: [] for k in _AN_SRC}
    for t in res.get("ok", []):
        ok[_an_src(t)].append(t)
    for d in res.get("doubt", []):
        t = d["text"] if isinstance(d, dict) else str(d)
        db[_an_src(t)].append(t)
    info = []
    for t in res.get("info", []):
        if t.startswith("Поправки проверены") or t.startswith("Сборники ТСН"):
            ok["Из сборников ТСН"].append(t)
        else:
            info.append(t)
    # по источникам, из которых ничего не взято, — явная строка, чтобы было понятно
    use_pd = res.get("with_pd", True)
    if use_pd and not ok["Из ПЗ"] and not db["Из ПЗ"]:
        ok["Из ПЗ"].append("Из ПЗ в этот перечень ничего не применено: данных, меняющих подбор монолитных работ, в ПЗ нет.")
    if not use_pd:
        for k in ("Из ПОС", "Из ПЗ"):
            ok[k] = [f"{k[3:]} не учитывался (выбрано «без учёта ПОС и ПЗ»)."]
        ok["Из ПОС и ПЗ"] = []
    res["groups"] = {
        "ok": [{"src": k, "items": v} for k, v in ok.items() if v],
        "doubt": [{"src": k, "items": v} for k, v in db.items() if v],
    }
    res["info"] = info
    try:
        json.dump(res, open(os.path.join(run_dir, "analytic.json"), "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    except Exception:
        pass
    return res


# ==== IFC_FEATURES: признаки по геометрии модели -> в проверку поправок и в справку ====
IFC_FEAT = {}
try:
    from src.services import ifc_features as _ifcf
    from src.services import tsn_coeffs as _tsnc
    _el_facts_orig = _tsnc._el_facts

    def _el_facts_geo(el, q):
        f = _el_facts_orig(el, q)
        try:
            for k, v in _ifcf.facts_for(str(el.get("name") or ""), IFC_FEAT).items():
                f["по модели: " + k] = v
        except Exception:
            pass
        return f
    _tsnc._el_facts = _el_facts_geo
except Exception as _e:
    logger.warning(f"признаки модели не подключены: {_e}")
    _ifcf = None


def _session_dir(run_dir):
    p = os.path.abspath(run_dir)
    return os.path.dirname(p) if os.path.basename(p).startswith("run_") else p


_build_before_feat = build_final_from_perechen


def build_final_from_perechen(tables_json_path, run_dir):
    global IFC_FEAT
    IFC_FEAT = {}
    if _ifcf is not None:
        try:
            IFC_FEAT = _ifcf.load(_session_dir(run_dir))
            logger.info(f"Признаки модели: типов {len(IFC_FEAT)}")
        except Exception as e:
            logger.warning(f"признаки модели: {e}")
    return _build_before_feat(tables_json_path, run_dir)


_an_before_feat = _an_build


def _an_build(run_dir):
    res = _an_before_feat(run_dir)
    if not IFC_FEAT:
        return res
    t = list(IFC_FEAT.values())
    cur = sum(1 for x in t if x.get("curved") or x.get("curved_contour"))
    tilt = sum(1 for x in t if x.get("tilt_max"))
    tap = sum(1 for x in t if x.get("tapered"))
    zs = [x["z_min"] for x in t if x.get("z_min") is not None]
    ze = [x["z_max"] for x in t if x.get("z_max") is not None]
    hs = sorted({x["floor_h_max"] for x in t if x.get("floor_h_max") and x["floor_h_max"] >= 2.5})
    parts = []
    parts.append(f"криволинейных элементов: {cur} тип." if cur else "криволинейных элементов нет")
    parts.append(f"наклонных: {tilt} тип." if tilt else "наклонных нет")
    if tap:
        parts.append(f"с переменным сечением: {tap} тип.")
    if zs and ze:
        parts.append(f"отметки от {min(zs):+.1f} до {max(ze):+.1f} м от ±0.000".replace(".", ","))
    if hs:
        parts.append("высоты этажей " + (f"{hs[0]}" if len(hs) == 1 else f"от {hs[0]} до {hs[-1]}").replace(".", ",") + " м")
    line = "Геометрия по модели: " + "; ".join(parts) + " — учтено при проверке поправок из сборников."
    g = res.setdefault("groups", {}).setdefault("ok", [])
    tgt = next((x for x in g if x.get("src") == "Из модели и справочника цен"), None)
    if tgt is None:
        tgt = {"src": "Из модели и справочника цен", "items": []}
        g.append(tgt)
    tgt["items"].append(line)
    try:
        json.dump(res, open(os.path.join(run_dir, "analytic.json"), "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    except Exception:
        pass
    return res


# ==== CONFLICTS: противоречия между ПОС, ПЗ, сборниками ТСН, моделью и расценками перечня ====
import hashlib, re, os, json
CONFLICT_SRC = "Противоречия между документами"
_CONF_CACHE = "/app/data/tsn_conflicts_cache.json"


def _model_facts():
    out = []
    try:
        zs = [x["z_min"] for x in (IFC_FEAT or {}).values() if x.get("z_min") is not None]
        ze = [x["z_max"] for x in (IFC_FEAT or {}).values() if x.get("z_max") is not None]
        if zs and ze:
            out.append(f"отметки элементов от {min(zs)} до {max(ze)} м от ±0.000")
    except Exception:
        pass
    for k, v in (PROJECT_CONSTANTS or {}).items():
        if str(k).startswith("_"):
            continue
        out.append(f"параметр расчёта {k} = {v}")
    classes = set()
    for s in _AN_STATS:
        n = s["el"]
        b = re.search(r"[BВ](\d{2}(?:[.,]\d)?)", n)
        if b:
            w = re.search(r"W\d+", n); f = re.search(r"F\d+", n)
            classes.add(" ".join(x for x in ("B" + b.group(1), w and w.group(0), f and f.group(0)) if x))
    if classes:
        out.append("классы бетона в модели: " + ", ".join(sorted(classes)[:12]))
    return out


def _conflicts(pms_titles):
    try:
        from src.services import tsn_coeffs
        from src.services.pd_full import ask
    except Exception:
        return []
    tech = tsn_coeffs.tech_for(pms_titles.keys())
    docs = [r for r in (PD_FULL or {}).values() if r.get("value")]
    if not docs and not tech:
        return []
    doc_txt = "\n".join(f"- [{r.get('src') or 'ПОС'}, стр. {r.get('page')}] {r['title'].split('(')[0].strip()}: {r['value']}"
                        f" (цитата: «{(r.get('quote') or '')[:140]}»)" for r in docs)
    tech_txt = ""
    for code, t in tech.items():
        tech_txt += f"\n### {code}, «1. Общие указания»:\n"
        tech_txt += "\n".join(f"п.{c['num']}: {c['text'][:380]}" for c in t["general"])
    model_txt = "\n".join("- " + x for x in _model_facts())
    rates_txt = "\n".join(f"- {pm}: {tt[:110]} ({n} поз.)" for pm, (tt, n) in sorted(pms_titles.items()))
    key = hashlib.md5((doc_txt + tech_txt + model_txt + rates_txt).encode()).hexdigest()
    try:
        cache = json.load(open(_CONF_CACHE, encoding="utf-8"))
    except Exception:
        cache = {}
    if key in cache:
        return cache[key]
    prompt = (
        "Ты опытный инженер-сметчик. Найди ПРОТИВОРЕЧИЯ между источниками:\n"
        "ПОС — проект организации строительства (технология работ); ПЗ — пояснительная записка (что проектируется);\n"
        "ТСН — общие указания сборника (когда какие расценки применять); МОДЕЛЬ — данные цифровой модели;\n"
        "ПЕРЕЧЕНЬ — расценки, которые стоят в перечне работ.\n"
        "Противоречие — когда один источник ЯВНО говорит одно, а другой — другое. Примеры: ПОС — бетон бадьёй, "
        "а в перечне насос; ПОС и ПЗ дают разную высоту/этажность; ПЗ — класс бетона B25, а в модели B30; "
        "сборник требует для таких условий другую расценку, чем стоит в перечне.\n"
        "Строго: только реальные явные расхождения, не выдумывай, не считай противоречием отсутствие данных. "
        "Мелкие расхождения в пределах округления (80,2 м и 80,6 м разных корпусов) — не противоречие. "
        "Если противоречий нет — пустой список.\n"
        "Для каждого: topic (2-4 слова); a — что в первом источнике; a_ref — ссылка: «ПОС, стр. N», «ПЗ, стр. N», "
        "«<код сборника> п.X.Y», «модель» или «перечень»; b, b_ref — то же для второго источника; "
        "why — простыми словами до 150 символов, в чём расхождение и что проверить.\n"
        'Ответ строго JSON: {"conflicts": [{"topic":"","a":"","a_ref":"","b":"","b_ref":"","why":""}]}\n'
        f"\n=== ПОС и ПЗ ===\n{doc_txt or '(нет)'}\n\n=== ТСН ==={tech_txt or ' (нет)'}"
        f"\n\n=== МОДЕЛЬ ===\n{model_txt or '(нет)'}\n\n=== ПЕРЕЧЕНЬ ===\n{rates_txt}"
    )
    try:
        ans = ask(prompt)
    except Exception as e:
        logger.warning(f"противоречия: LLM недоступна: {e}")
        return []
    clauses = {f"{code} п.{c['num']}" for code, t in tech.items() for c in t["general"]}
    pages = {(r.get("src") or "ПОС", str(r.get("page"))) for r in docs}

    def ref_ok(ref):
        ref = str(ref or "")
        m = re.search(r"(ТСН-[\d.]+-\d+)\D+(\d\.\d+(?:\.\d+)?)", ref)
        if m:
            return f"{m.group(1)} п.{m.group(2)}" in clauses
        m = re.search(r"(ПОС|ПЗ)\D*(\d+)", ref)
        if m:
            return any(src == m.group(1) for src, _ in pages)   # документ есть (страница могла быть соседней)
        return ref.strip().lower() in ("модель", "перечень") or "модел" in ref.lower() or "перечень" in ref.lower()

    def kind(ref):
        ref = str(ref)
        for k in ("ТСН", "ПОС", "ПЗ"):
            if k in ref:
                return k
        return "модель" if "модел" in ref.lower() else "перечень"

    out = []
    for c in (ans.get("conflicts") or []):
        if not isinstance(c, dict):
            continue
        if not (ref_ok(c.get("a_ref")) and ref_ok(c.get("b_ref"))):
            continue                               # ссылка на несуществующий пункт — отбрасываем
        if kind(c.get("a_ref")) == kind(c.get("b_ref")):
            continue
        out.append({k: str(c.get(k) or "")[:220] for k in ("topic", "a", "a_ref", "b", "b_ref", "why")})
    cache[key] = out
    try:
        json.dump(cache, open(_CONF_CACHE, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    except Exception:
        pass
    return out


_an_before_conf = _an_build


def _an_build(run_dir):
    res = _an_before_conf(run_dir)
    if not res.get("with_pd") and not res.get("with_tsn", True):
        return res
    pms = {}
    for s in _AN_STATS:
        for pm, t in s["rows"]:
            if re.match(r"^3\.\d+-\d+-\d+$", pm):
                tt = (_local_rate(pm) or {}).get("title") or ""
                pms[pm] = (tt, pms.get(pm, (tt, 0))[1] + 1)
    try:
        conf = _conflicts(pms)
    except Exception as e:
        logger.warning(f"противоречия: {e}")
        conf = []
    res["conflicts"] = conf
    g = res.setdefault("groups", {})
    g.setdefault("doubt", [])[:] = [x for x in g.get("doubt", []) if x.get("src") != CONFLICT_SRC]
    g.setdefault("ok", [])[:] = [x for x in g.get("ok", []) if x.get("src") != CONFLICT_SRC]
    if conf:
        g["doubt"].insert(0, {"src": CONFLICT_SRC, "items": [
            f"{c['topic']}: {c['a']} ({c['a_ref']}) ↔ {c['b']} ({c['b_ref']}). {c['why']} Решение за сметчиком."
            for c in conf]})
    else:
        g["ok"].append({"src": CONFLICT_SRC, "items": [
            "Проверено: противоречий между ПОС, ПЗ, общими указаниями сборников ТСН, моделью и расценками перечня не найдено."]})
    try:
        json.dump(res, open(os.path.join(run_dir, "analytic.json"), "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    except Exception:
        pass
    return res


# ==== CONFLICTS_FIX1: параметры расчёта не источник (они сами из ПОС) ====
def _model_facts():
    out = []
    try:
        zs = [x["z_min"] for x in (IFC_FEAT or {}).values() if x.get("z_min") is not None]
        ze = [x["z_max"] for x in (IFC_FEAT or {}).values() if x.get("z_max") is not None]
        if zs and ze:
            out.append(f"отметки элементов от {min(zs)} до {max(ze)} м от ±0.000 (высота здания по модели ≈ {max(ze)} м)")
    except Exception:
        pass
    classes = set()
    for s in _AN_STATS:
        n = s["el"]
        b = re.search(r"[BВ](\d{2}(?:[.,]\d)?)", n)
        if b:
            w = re.search(r"W\d+", n); f = re.search(r"F\d+", n)
            classes.add(" ".join(x for x in ("B" + b.group(1), w and w.group(0), f and f.group(0)) if x))
    if classes:
        out.append("классы бетона в модели: " + ", ".join(sorted(classes)[:12]))
    return out


# ==== CONFLICTS_FIX2: пункты сборника — только про таблицы перечня ====
def _conflicts(pms_titles):
    try:
        from src.services import tsn_coeffs
        from src.services.pd_full import ask
    except Exception:
        return []
    tech = tsn_coeffs.tech_for(pms_titles.keys())
    docs = [r for r in (PD_FULL or {}).values() if r.get("value")]
    if not docs and not tech:
        return []
    doc_txt = "\n".join(f"- [{r.get('src') or 'ПОС'}, стр. {r.get('page')}] {r['title'].split('(')[0].strip()}: {r['value']}"
                        f" (цитата: «{(r.get('quote') or '')[:140]}»)" for r in docs)
    tech_txt = ""
    for code, t in tech.items():
        tech_txt += f"\n### {code}, «1. Общие указания»:\n"
        tech_txt += "\n".join(f"п.{c['num']}: {c['text'][:380]}" for c in t["general"])
    model_txt = "\n".join("- " + x for x in _model_facts())
    rates_txt = "\n".join(f"- {pm}: {tt[:110]} ({n} поз.)" for pm, (tt, n) in sorted(pms_titles.items()))
    tabs = sorted({pm.rsplit("-", 1)[0][2:] for pm in pms_titles}, key=lambda x: [int(y) for y in re.findall(r"\d+", x)])
    rates_txt += "\nТаблицы сборников, из которых взяты расценки перечня: " + ", ".join(tabs)
    key = hashlib.md5((doc_txt + tech_txt + model_txt + rates_txt).encode()).hexdigest()
    try:
        cache = json.load(open(_CONF_CACHE, encoding="utf-8"))
    except Exception:
        cache = {}
    if key in cache:
        return cache[key]
    prompt = (
        "Ты опытный инженер-сметчик. Найди ПРОТИВОРЕЧИЯ между источниками:\n"
        "ПОС — проект организации строительства (технология работ); ПЗ — пояснительная записка (что проектируется);\n"
        "ТСН — общие указания сборника (когда какие расценки применять); МОДЕЛЬ — данные цифровой модели;\n"
        "ПЕРЕЧЕНЬ — расценки, которые стоят в перечне работ.\n"
        "Противоречие — когда один источник ЯВНО говорит одно, а другой — другое. Примеры: ПОС — бетон бадьёй, "
        "а в перечне насос; ПОС и ПЗ дают разную высоту/этажность; ПЗ — класс бетона B25, а в модели B30; "
        "сборник требует для таких условий другую расценку, чем стоит в перечне.\n"
        "Строго: только реальные явные расхождения, не выдумывай, не считай противоречием отсутствие данных. "
        "Мелкие расхождения в пределах округления (80,2 м и 80,6 м разных корпусов) — не противоречие. "
        "ВАЖНО про сборник: пункт общих указаний учитывай, только если он относится к таблицам, из которых взяты "
        "расценки перечня (список таблиц дан в конце). Если пункт описывает другой отдел/другие таблицы "
        "(например, «в отделе 1 … деревянная щитовая опалубка», а расценки перечня из отдела 2) — это НЕ противоречие.\n"
        "Если противоречий нет — пустой список.\n"
        "Для каждого: topic (2-4 слова); a — что в первом источнике; a_ref — ссылка: «ПОС, стр. N», «ПЗ, стр. N», "
        "«<код сборника> п.X.Y», «модель» или «перечень»; b, b_ref — то же для второго источника; "
        "why — простыми словами до 150 символов, в чём расхождение и что проверить.\n"
        'Ответ строго JSON: {"conflicts": [{"topic":"","a":"","a_ref":"","b":"","b_ref":"","why":""}]}\n'
        f"\n=== ПОС и ПЗ ===\n{doc_txt or '(нет)'}\n\n=== ТСН ==={tech_txt or ' (нет)'}"
        f"\n\n=== МОДЕЛЬ ===\n{model_txt or '(нет)'}\n\n=== ПЕРЕЧЕНЬ ===\n{rates_txt}"
    )
    try:
        ans = ask(prompt)
    except Exception as e:
        logger.warning(f"противоречия: LLM недоступна: {e}")
        return []
    clauses = {f"{code} п.{c['num']}" for code, t in tech.items() for c in t["general"]}
    pages = {(r.get("src") or "ПОС", str(r.get("page"))) for r in docs}

    def ref_ok(ref):
        ref = str(ref or "")
        m = re.search(r"(ТСН-[\d.]+-\d+)\D+(\d\.\d+(?:\.\d+)?)", ref)
        if m:
            return f"{m.group(1)} п.{m.group(2)}" in clauses
        m = re.search(r"(ПОС|ПЗ)\D*(\d+)", ref)
        if m:
            return any(src == m.group(1) for src, _ in pages)   # документ есть (страница могла быть соседней)
        return ref.strip().lower() in ("модель", "перечень") or "модел" in ref.lower() or "перечень" in ref.lower()

    def kind(ref):
        ref = str(ref)
        for k in ("ТСН", "ПОС", "ПЗ"):
            if k in ref:
                return k
        return "модель" if "модел" in ref.lower() else "перечень"

    out = []
    for c in (ans.get("conflicts") or []):
        if not isinstance(c, dict):
            continue
        if not (ref_ok(c.get("a_ref")) and ref_ok(c.get("b_ref"))):
            continue                               # ссылка на несуществующий пункт — отбрасываем
        if kind(c.get("a_ref")) == kind(c.get("b_ref")):
            continue
        out.append({k: str(c.get(k) or "")[:220] for k in ("topic", "a", "a_ref", "b", "b_ref", "why")})
    cache[key] = out
    try:
        json.dump(cache, open(_CONF_CACHE, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    except Exception:
        pass
    return out




# ==== CONFLICTS_RULES: противоречия по жёстким правилам (без qwen — находятся всегда) ====
def _rule_conflicts():
    """Подача бетона: ПОС по части здания vs расценки кран-бадья / насос в перечне."""
    out, ok = [], []
    if not PD_FULL:
        return out, ok
    parts = {"under": ("concrete_supply_under", "подземной части"), "above": ("concrete_supply_above", "надземной части")}
    cnt = {"under": {"b": set(), "p": set()}, "above": {"b": set(), "p": set()}}
    for s in _AN_STATS:
        for pm, t in s["rows"]:
            title = ((_local_rate(pm) or {}).get("title") or t).lower()
            tt = _table_title(pm) if re.match(r"^3\.\d+-\d+-\d+$", pm) else ""
            ctx = (tt or title)
            part = "under" if ("подземн" in ctx and "надземн" not in ctx) or "фундамент" in title else "above"
            if "кран-бадья" in title or "кран-бадья" in tt:
                cnt[part]["b"].add(pm)
            elif "бетононасос" in title:
                cnt[part]["p"].add(pm)
    for part, (pid, label) in parts.items():
        r = PD_FULL.get(pid) or {}
        v = f"{r.get('value') or ''} {r.get('quote') or ''}".lower()
        if not r.get("value"):
            continue
        ref = f"{r.get('src') or 'ПОС'}, стр. {r.get('page')}"
        pump, bucket = "насос" in v, "бадь" in v
        b, p = sorted(cnt[part]["b"]), sorted(cnt[part]["p"])
        if pump and not bucket and b:
            out.append(f"Подача бетона в {label}: ПОС — только насосом ({ref}) ↔ в перечне расценки «кран-бадья» "
                       f"({', '.join(b[:4])}{' и др.' if len(b) > 4 else ''}). Насосных расценок для этих конструкций "
                       f"в ТСН нет — выбрать: оставить кран-бадью или учесть подачу насосом отдельно. Решение за сметчиком.")
        elif bucket and not pump and p:
            out.append(f"Подача бетона в {label}: ПОС — кран-бадьёй ({ref}) ↔ в перечне расценки на бетононасос "
                       f"({', '.join(p[:4])}). Решение за сметчиком.")
        elif pump and bucket and b:
            ok.append(f"Подача бетона в {label}: ПОС допускает кран-бадью или насос ({ref}) — в перечне кран-бадья "
                      f"({len(b)} расц.), это соответствует ПОС.")
        elif pump and p:
            ok.append(f"Подача бетона в {label}: по ПОС насосом ({ref}) — в перечне расценки на бетононасос ({len(p)} расц.).")
    return out, ok


_an_before_rules = _an_build


def _an_build(run_dir):
    res = _an_before_rules(run_dir)
    if not res.get("with_pd"):
        return res
    try:
        confl, oks = _rule_conflicts()
    except Exception as e:
        logger.warning(f"правила противоречий: {e}")
        return res
    g = res.setdefault("groups", {})
    # старую общую строку про подачу бетона убираем — теперь точнее, по частям здания
    for sect in ("doubt", "ok"):
        for x in g.get(sect, []):
            x["items"] = [i for i in x["items"] if not i.startswith("Подача бетона: в ПОС насос")]
        g[sect] = [x for x in g.get(sect, []) if x["items"]]
    if confl:
        g["ok"] = [x for x in g.get("ok", []) if x.get("src") != CONFLICT_SRC]
        tgt = next((x for x in g.setdefault("doubt", []) if x.get("src") == CONFLICT_SRC), None)
        if tgt is None:
            tgt = {"src": CONFLICT_SRC, "items": []}
            g["doubt"].insert(0, tgt)
        llm = [i for i in tgt["items"] if "одач" not in i.lower() or "бетон" not in i.lower()]   # дубли qwen про подачу
        tgt["items"] = confl + llm
    if oks:
        tgt = next((x for x in g.setdefault("ok", []) if x.get("src") == "Из ПОС"), None)
        if tgt is None:
            tgt = {"src": "Из ПОС", "items": []}
            g["ok"].insert(0, tgt)
        tgt["items"] += oks
    try:
        json.dump(res, open(os.path.join(run_dir, "analytic.json"), "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    except Exception:
        pass
    return res


# ==== BOOK_MATERIALS: материалы, не учтённые в расценках (по сборнику ТСН) ====
import re, os, json
_BOOK_HITS = []


def _unit_mult(u):
    m = re.match(r"\s*(\d+(?:[.,]\d+)?)", str(u or ""))
    return float(m.group(1).replace(",", ".")) if m else 1.0


def _stems(name):
    words = re.findall(r"[а-яёa-z]+", str(name).lower())
    return [w[:6] for w in words if len(w) >= 5 and w not in ("смеси", "тяжелого", "легкого", "отдельные", "изделия")]


_fin_before_mat = finalize_rows


def finalize_rows(el, q, works, subs):
    rows = _fin_before_mat(el, q, works, subs)
    if not TSN_ON:
        return rows
    try:
        from src.services import tsn_coeffs
    except Exception:
        return rows
    # материал ищем только среди строк-ресурсов (не среди работ: «железобетонных» есть в названии любой работы)
    texts = [(i, str(r.get("title") or "").lower()) for i, r in enumerate(rows)
             if isinstance(r, dict) and (r.get("is_resource") or not str(r.get("pressmark") or "").startswith("3."))]
    for i, r in enumerate(rows):
        if not isinstance(r, dict) or r.get("is_resource") or r.get("_bad"):
            continue
        b = tsn_coeffs.book_rate(r.get("pressmark"))
        if not b or not b["not_included"]:
            continue
        vol = r.get("volume")
        k = _unit_mult(r.get("unit")) / _unit_mult(b["unit"])       # «100 м3» в перечне vs «100 м3» в сборнике
        for m in b["not_included"]:
            val = m.get("value")
            need = round(float(val) * float(vol) * k, 3) if isinstance(val, (int, float)) and isinstance(vol, (int, float)) else None
            st = _stems(m.get("name"))
            have = [(j, t) for j, t in texts if j != i and st and any(s_ in t for s_ in st)]
            have_vol = None
            for j, _ in have:
                rv = rows[j].get("volume")
                if isinstance(rv, (int, float)) and str(rows[j].get("unit") or "").strip().lower().endswith(str(m.get("unit") or "").lower()):
                    have_vol = (have_vol or 0) + rv * _unit_mult(rows[j].get("unit"))
            hit = {"el": str(el.get("name") or "элемент"), "pm": r.get("pressmark"), "doc": b["doc"], "table": b["table"],
                   "material": m.get("name"), "unit": m.get("unit"), "per": val, "per_unit": b["unit"],
                   "need": need, "present": bool(have), "have_vol": have_vol}
            _BOOK_HITS.append(hit)
            if not have and need:
                note = f" [по сборнику {b['doc']} табл. {b['table']} отдельно нужен материал: {m.get('name')} ≈ {need} {m.get('unit')}]"
                if note not in str(r.get("title")):
                    r["title"] = str(r.get("title") or "") + note
    return rows


_an_before_mat = _an_build


def _an_build(run_dir):
    res = _an_before_mat(run_dir)
    if not res.get("with_tsn", True) or not _BOOK_HITS:
        return res
    g = res.setdefault("groups", {})
    agg = {}
    for h in _BOOK_HITS:
        key = (h["material"], h["unit"])
        a = agg.setdefault(key, {"need": 0.0, "have": 0.0, "have_known": False, "missing": [], "present": 0, "doc": h["doc"],
                                 "tables": set(), "per": h["per"], "per_unit": h["per_unit"]})
        a["tables"].add(h["table"])
        if h["need"]:
            a["need"] += h["need"]
        if h["present"]:
            a["present"] += 1
            if h["have_vol"]:
                a["have"] += h["have_vol"]; a["have_known"] = True
        else:
            n = _an_short(h["el"])
            if n not in a["missing"]:
                a["missing"].append(n)
    ok_items, dbt_items = [], []
    fmt = lambda x: f"{x:,.1f}".replace(",", " ").replace(".", ",")
    for (mat, unit), a in agg.items():
        tabs = ", ".join(sorted(a["tables"]))
        mat_s = re.sub(r"\s*\([\d,\s]+\)", "", mat).strip()
        mat_s = mat_s if len(mat_s) <= 60 else mat_s[:58] + "…"
        if a["present"]:
            s = (f"«{mat_s}» — не входит в расценки (сборник {a['doc']}, табл. {tabs}: {str(a['per']).replace('.', ',')} {unit} "
                 f"на {a['per_unit']}) — в перечне есть отдельной строкой.")
            if a["have_known"] and a["need"]:
                d = (a["have"] - a["need"]) / a["need"] * 100
                s += f" По норме сборника нужно ≈ {fmt(a['need'])} {unit}, в перечне {fmt(a['have'])} {unit}"
                s += " — совпадает." if abs(d) <= 3 else f" — расхождение {d:+.0f}%, проверьте (в норме сборника учтены потери)."
            ok_items.append(s)
        if a["missing"]:
            dbt_items.append(f"«{mat_s}» — по сборнику {a['doc']} (табл. {tabs}) не входит в расценку и должен быть отдельной "
                             f"строкой, а в перечне его нет: {'; '.join(a['missing'][:3])}{' и др' if len(a['missing']) > 3 else ''}. "
                             + (f"По норме нужно ≈ {fmt(a['need'])} {unit}." if a["need"] >= 0.05 else "Количество — по проекту (объём работы по модели не определён).") + " Добавьте в перечень.")
    for sect, items in (("ok", ok_items), ("doubt", dbt_items)):
        if not items:
            continue
        tgt = next((x for x in g.setdefault(sect, []) if x.get("src") == "Из сборников ТСН"), None)
        if tgt is None:
            tgt = {"src": "Из сборников ТСН", "items": []}
            g[sect].append(tgt)
        tgt["items"] += items
    try:
        json.dump(res, open(os.path.join(run_dir, "analytic.json"), "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    except Exception:
        pass
    return res


_build_before_mat = build_final_from_perechen


def build_final_from_perechen(tables_json_path, run_dir):
    _BOOK_HITS.clear()
    return _build_before_mat(tables_json_path, run_dir)


# ==== REBAR_STEEL: сама арматурная сталь — отдельной строкой (по сборнику не входит в расценку установки) ====
_REBAR_ADDED = []
_REBAR_DEFAULT = "1.3-4-22"      # А-III, Ø12–14 мм — принят, если диаметр в модели не указан
_REBAR_OK_CLASS = re.compile(r"(А|A)\s*500|А-?III|A-?III", re.I)


_RES_CACHE = None


def _rebar_price_info():
    global _RES_CACHE
    if _RES_CACHE is None:
        try:
            # Путь/кеш ресурсов — через модуль local_resources (env-переменные,
            # а не жёстко /app/data — чтобы работало и в рабочей копии).
            from src.services import local_resources as _lr
            _RES_CACHE = _lr.load()
        except Exception:
            _RES_CACHE = {}
    return _RES_CACHE.get(_REBAR_DEFAULT) or {}


def _rebar_class(el):
    for k, v in (el or {}).items():
        if isinstance(v, str) and ("ReinforceStrengthClass" in str(k) or "класс арматуры" in str(k).lower()):
            return v.strip()
    txt = " ".join(str(v) for v in (el or {}).values() if isinstance(v, str))
    m = re.search(r"[АA]\s?500\s?[СC]?|[АA]-?III", txt)
    return m.group(0) if m else ""


# встраиваемся ПОД проверку материалов, чтобы она уже видела добавленную сталь
_fin_before_rebar = _fin_before_mat


def _finalize_with_rebar(el, q, works, subs):
    rows = _fin_before_rebar(el, q, works, subs)
    try:
        if any(isinstance(r, dict) and (r.get("is_resource") or not str(r.get("pressmark") or "").startswith("3."))
               and "арматур" in str(r.get("title") or "").lower() for r in rows):
            return rows                                            # сталь уже есть
        idx = next((i for i, r in enumerate(rows) if isinstance(r, dict) and not r.get("is_resource")
                    and "арматур" in str(r.get("title") or "").lower() and "закладн" not in str(r.get("title") or "").lower()
                    and str(r.get("unit") or "").strip().lower().endswith("т")), None)
        if idx is None:
            return rows
        work = rows[idx]
        vol = work.get("volume")
        if not isinstance(vol, (int, float)) or vol <= 0:
            return rows
        cls = _rebar_class(el) or "А500С"
        if not _REBAR_OK_CLASS.search(cls):
            _REBAR_ADDED.append({"el": str(el.get("name") or ""), "t": 0, "cls": cls, "skipped": True})
            return rows
        info = _rebar_price_info()
        price = info.get("price")
        t = round(float(vol) * _unit_mult(work.get("unit")), 4)
        title = (f"Арматура {cls} (А-III), сталь периодического профиля — не учтена в расценке установки "
                 f"(диаметр по проекту; цена принята для Ø12–14 мм, уточнить по спецификации КР)")
        row = {"pressmark": _REBAR_DEFAULT, "title": title, "unit": "т", "volume": t,
               "zp": 0, "em": 0, "mr": round(price * t, 2) if isinstance(price, (int, float)) else None,
               "total": round(price * t, 2) if isinstance(price, (int, float)) else None,
               "is_resource": True, "work": work.get("work") or "Арматура"}
        rows.insert(idx + 1, row)
        _REBAR_ADDED.append({"el": str(el.get("name") or ""), "t": t, "cls": cls, "priced": price is not None})
    except Exception as e:
        logger.warning(f"арматура-сталь: {e}")
    return rows


_fin_before_mat = _finalize_with_rebar


_an_before_rebar = _an_build


def _an_build(run_dir):
    res = _an_before_rebar(run_dir)
    added = [x for x in _REBAR_ADDED if not x.get("skipped")]
    skipped = [x for x in _REBAR_ADDED if x.get("skipped")]
    if not added and not skipped:
        return res
    g = res.setdefault("groups", {})
    fmt = lambda x: f"{x:,.1f}".replace(",", " ").replace(".", ",")
    tot = sum(x["t"] for x in added)
    classes = sorted({x["cls"] for x in added})

    def put(sect, src, text):
        tgt = next((x for x in g.setdefault(sect, []) if x.get("src") == src), None)
        if tgt is None:
            tgt = {"src": src, "items": []}
            g[sect].append(tgt)
        tgt["items"].append(text)
    if added:
        put("ok", "Из сборников ТСН",
            f"Арматурная сталь добавлена в перечень отдельными строками: {fmt(tot)} т ({len(added)} эл.), класс по модели "
            f"{', '.join(classes)} — по сборнику сталь не входит в расценки установки арматуры.")
        put("doubt", "Из модели и справочника цен",
            f"Диаметры арматуры в модели не указаны — цена стали ({fmt(tot)} т) принята для Ø12–14 мм (1.3-4-22). "
            f"Уточните по спецификации КР: цена зависит от диаметра.")
    if skipped:
        put("doubt", "Из модели и справочника цен",
            f"Арматура класса {', '.join(sorted({x['cls'] for x in skipped}))} — сталь не добавлена автоматически "
            f"({len(skipped)} эл.), добавьте по проекту.")
    # строки «арматура — в перечне нет» от проверки материалов больше не нужны там, где сталь добавлена
    for x in g.get("doubt", []):
        x["items"] = [i for i in x["items"] if not (i.startswith("«Арматурные заготовки") and added)]
    g["doubt"] = [x for x in g.get("doubt", []) if x["items"]]
    try:
        json.dump(res, open(os.path.join(run_dir, "analytic.json"), "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    except Exception:
        pass
    return res


_build_before_rebar = build_final_from_perechen


def build_final_from_perechen(tables_json_path, run_dir):
    _REBAR_ADDED.clear()
    return _build_before_rebar(tables_json_path, run_dir)


# ==== PD_MODE: значения из ПОС подставляются на сервере (не зависят от того, успела ли страница) ====
def _pd_mode_on(run_dir):
    p = os.path.abspath(run_dir)
    for _ in range(4):
        fp = os.path.join(p, "pd_mode.json")
        if os.path.isfile(fp):
            try:
                return json.load(open(fp, encoding="utf-8")).get("mode", "on") != "off"
            except Exception:
                return True
        p = os.path.dirname(p)
    return True


def _pos_constants(session_dir):
    import glob
    out = {}
    for fp in glob.glob(os.path.join(session_dir, "*ПОС*конст*.json")) + glob.glob(os.path.join(session_dir, "pos_constants*.json")):
        try:
            d = json.load(open(fp, encoding="utf-8"))
        except Exception:
            continue
        c = d.get("constants", d) if isinstance(d, dict) else {}
        for k, v in (c.items() if isinstance(c, dict) else []):
            val = v.get("value") if isinstance(v, dict) else v
            found = v.get("found", True) if isinstance(v, dict) else True
            if found and val not in (None, "", [], {}):
                out[k] = val
    return out


_PD_RUN_DIR = None
_build_base_orig = _build_base


def _build_base(tables_json_path, run_dir):
    """Вызывается изнутри обёртки зимнего ухода — PROJECT_CONSTANTS уже прочитаны из запроса страницы."""
    global PROJECT_CONSTANTS
    try:
        sd = os.path.dirname(os.path.abspath(run_dir)) if os.path.basename(os.path.abspath(run_dir)).startswith("run_") else run_dir
        if _pd_mode_on(run_dir):
            added = []
            for k, v in _pos_constants(sd).items():
                if not PROJECT_CONSTANTS.get(k):
                    PROJECT_CONSTANTS[k] = v
                    added.append(k)
            if added:
                logger.info(f"ПОС: подставлены на сервере {added}")
        else:
            PROJECT_CONSTANTS = {}          # режим «без учёта ПОС и ПЗ»
            logger.info("ПОС и ПЗ не учитываются (режим на странице)")
    except Exception as e:
        logger.warning(f"ПОС на сервере: {e}")
    return _build_base_orig(tables_json_path, run_dir)


# ==== RATE_FIX: расценка из «чужого» сборника — проверка qwen и замена на подходящую ====
import re, os, json, hashlib, collections
_RATE_FIX = []
_RATE_FIX_CACHE = "/app/data/rate_fix_cache.json"


def _coll(pm):
    m = re.match(r"^3\.(\d+)-", str(pm or ""))
    return m.group(1) if m else None


def _catalog_rates(table):
    out = []
    for k in range(1, 60):
        loc = _local_rate(f"{table}-{k}")
        if loc and loc.get("title"):
            out.append((f"{table}-{k}", loc["title"], loc.get("unitOfMeasure") or ""))
        elif k > 3 and not out:
            break
    return out


def _fix_ask(work, el_name, pm, title, main, tables):
    from src.services.pd_full import ask
    key = hashlib.md5(json.dumps([work, pm, main], ensure_ascii=False).encode()).hexdigest()
    try:
        cache = json.load(open(_RATE_FIX_CACHE, encoding="utf-8"))
    except Exception:
        cache = {}
    if key in cache:
        return cache[key]
    tlist = "\n".join(f"{c}: {t[:140]}" for c, t in tables)
    a1 = ask(
        "Ты инженер-сметчик. Для работы подобрана расценка из другого сборника ТСН. Проверь, подходит ли она.\n"
        f"Элемент: {el_name}\nРабота: {work}\nПодобрана: {pm} — {title}\n"
        f"Если расценка НЕ соответствует работе, выбери подходящую таблицу из сборника {main} (список ниже) или null, если подходящей нет.\n"
        'Ответ JSON: {"fits": true|false, "table": "3.N-NN" или null, "why": "до 100 символов простыми словами"}\n\n'
        f"Таблицы сборника {main}:\n{tlist}")
    res = {"fits": bool(a1.get("fits")), "why": str(a1.get("why") or "")[:140], "rate": None}
    tb = a1.get("table")
    mt = re.search(r"(\d+)\s*-\s*(\d+)", str(tb or ""))
    tb = f"3.{mt.group(1)}-{mt.group(2)}" if mt else None
    if not res["fits"] and tb and any(tb == c for c, _ in tables):
        cands = _catalog_rates(tb)
        if cands:
            a2 = ask(
                "Выбери одну расценку для работы из списка (или null, если ни одна не подходит).\n"
                f"Элемент: {el_name}\nРабота: {work}\n\n" + "\n".join(f"{c}: {t[:160]} [{u}]" for c, t, u in cands)
                + '\n\nОтвет JSON: {"rate": "3.N-NN-N" или null}')
            r = a2.get("rate")
            mr = re.search(r"(\d+)\s*-\s*(\d+)\s*-\s*(\d+)", str(r or ""))
            r = f"3.{mr.group(1)}-{mr.group(2)}-{mr.group(3)}" if mr else None
            if r and any(r == c for c, _, _ in cands):
                res["rate"] = r
    cache[key] = res
    try:
        json.dump(cache, open(_RATE_FIX_CACHE, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    except Exception:
        pass
    return res


_fin_before_fix = _fin_before_mat


def _finalize_fix(el, q, works, subs):
    rows = _fin_before_fix(el, q, works, subs)
    try:
        from src.services import tsn_coeffs
        colls = collections.Counter(_coll(r.get("pressmark")) for r in rows
                                    if isinstance(r, dict) and not r.get("is_resource") and _coll(r.get("pressmark")))
        if not colls:
            return rows
        main = colls.most_common(1)[0][0]
        tables = tsn_coeffs.tables_of(main)
        if not tables:
            return rows
        for i, r in enumerate(rows):
            if not isinstance(r, dict) or r.get("is_resource"):
                continue
            n = _coll(r.get("pressmark"))
            if not n or n == main:
                continue
            title = (_local_rate(r["pressmark"]) or {}).get("title") or str(r.get("title") or "")
            work = str(r.get("work") or "")
            res = _fix_ask(work, str(el.get("name") or ""), r["pressmark"], title, main, tables)
            if res.get("fits"):
                continue
            rec = {"el": str(el.get("name") or ""), "work": work, "old": r["pressmark"], "old_title": title[:80],
                   "new": res.get("rate"), "why": res.get("why")}
            if res.get("rate"):
                loc = _local_rate(res["rate"]) or {}
                rate = {"pressmark": res["rate"], "unit": loc.get("unitOfMeasure") or r.get("unit"),
                        "title": loc.get("title") or res["rate"], "is_resource": False, "formula": ""}
                try:
                    rows[i] = _row(rate, el, q, work)
                    rec["new_title"] = (loc.get("title") or "")[:80]
                except Exception as e:
                    logger.warning(f"замена расценки {r['pressmark']}: {e}")
                    rec["new"] = None
            if not rec["new"]:
                r["title"] = str(r.get("title") or "") + " [расценка не соответствует работе — подобрать вручную]"
                r["_bad"] = True
                r["total"] = None
            _RATE_FIX.append(rec)
            logger.info(f"Проверка расценки: {rec}")
    except Exception as e:
        logger.warning(f"проверка расценок: {e}")
    return rows


_fin_before_mat = _finalize_fix


_an_before_fix = _an_build


def _an_build(run_dir):
    res = _an_before_fix(run_dir)
    if not _RATE_FIX:
        return res
    g = res.setdefault("groups", {})

    def put(sect, text):
        tgt = next((x for x in g.setdefault(sect, []) if x.get("src") == "Из сборников ТСН"), None)
        if tgt is None:
            tgt = {"src": "Из сборников ТСН", "items": []}
            g[sect].append(tgt)
        tgt["items"].append(text)
    seen = set()
    for x in _RATE_FIX:
        k = (x["work"], x["old"], x.get("new"))
        if k in seen:
            continue
        seen.add(k)
        els = [_an_short(y["el"]) for y in _RATE_FIX if (y["work"], y["old"], y.get("new")) == k]
        els = list(dict.fromkeys(els))
        who = "; ".join(els[:3]) + (" и др" if len(els) > 3 else "")
        if x.get("new"):
            put("ok", f"Работа «{x['work'][:60]}» ({who}): расценка {x['old']} ({x['old_title'][:50]}…) не подходила — "
                      f"заменена на {x['new']} ({x.get('new_title', '')[:60]}…). {x.get('why') or ''}")
        else:
            put("doubt", f"Работа «{x['work'][:60]}» ({who}): расценка {x['old']} ({x['old_title'][:50]}…) не соответствует "
                         f"работе, подходящей в сборниках не найдено — подберите вручную. {x.get('why') or ''}")
    try:
        json.dump(res, open(os.path.join(run_dir, "analytic.json"), "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    except Exception:
        pass
    return res


_build_before_fix = build_final_from_perechen


def build_final_from_perechen(tables_json_path, run_dir):
    _RATE_FIX.clear()
    return _build_before_fix(tables_json_path, run_dir)
