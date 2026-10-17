"""Read CS final works exactly as exported, retaining group boundaries and empty groups."""
import hashlib
from pathlib import Path
from openpyxl import load_workbook

FINAL_WORKBOOK='ОБЩИЙ_Финальный_перечень_работ.xlsx'


def read_final_works(folder, selected_groups):
    from .api_works_lookup import _resolve_building_part, BUILDING_PART_ORDER, _BUILDING_PART_TITLES
    path=Path(folder)/FINAL_WORKBOOK
    if not path.is_file():return None
    order=[i for part in BUILDING_PART_ORDER for i,g in enumerate(selected_groups) if _resolve_building_part(g)==part]
    if len(order)!=len(selected_groups):raise ValueError('Не определена часть здания группы ЦС')
    result={};headers=[];active=None
    workbook=load_workbook(path,read_only=True,data_only=True)
    try:
        sheet=workbook['Данные']
        for row_number,values in enumerate(sheet.iter_rows(min_row=2,values_only=True),2):
            code,name,unit,quantity=(*values[:4],)
            if not code:
                if name in _BUILDING_PART_TITLES.values():active=None
                elif name and 'Кол-во:' in str(name):
                    if len(headers)>=len(order):raise ValueError('В финальном перечне ЦС лишняя группа')
                    active=order[len(headers)];headers.append(str(name));result[active]=[]
                continue
            if active is None:raise ValueError(f'Работа ЦС вне группы: строка {row_number}')
            if quantity is None or str(quantity).strip() in ('','-'):number=None
            else:
                try:number=float(str(quantity).replace('\u00a0','').replace(' ','').replace(',','.'))
                except ValueError as exc:raise ValueError(f'Некорректный объём ЦС в строке {row_number}') from exc
            result[active].append({'positionType':'work','code':str(code),'name':str(name or ''),'unit':str(unit or ''),
                                  'quantity':number,'quantityStatus':'missing_data' if number is None else 'calculated',
                                  'selectionParameters':{},'quantityCalculation':{'source':'final_cs_workbook','file':FINAL_WORKBOOK,'sheet':'Данные','row':row_number,'quantityUnit':str(unit or '')}})
        if len(headers)!=len(order):raise ValueError(f'Группы финального перечня ЦС не совпали: {len(headers)} вместо {len(order)}')
    finally:workbook.close()
    return result


def synchronize_cs_model(model, folder):
    import json,copy
    selected_path=Path(folder)/'selected_elements_grouped.json'
    if not selected_path.is_file() or not (Path(folder)/FINAL_WORKBOOK).is_file():return model
    revision=hashlib.sha256((Path(folder)/FINAL_WORKBOOK).read_bytes()).hexdigest()
    if model.get('workSource')==FINAL_WORKBOOK and model.get('workSourceRevision')==revision and model.get('workSourceReaderVersion')==2:return model
    selected=json.loads(selected_path.read_text(encoding='utf-8'))
    works=read_final_works(folder,selected)
    groups=model.get('elementGroups',[])
    if len(groups)!=len(selected):raise ValueError('Не совпали группы модели и финального перечня ЦС')
    output=copy.deepcopy(model)
    for index,group in enumerate(output['elementGroups']):
        if group.get('groupId')!=f'group-{index+1:04d}':raise ValueError('Нарушен порядок групп ЦС')
        group['positions']=works[index]
        group['workSelectionNote']='Работы не подобраны в исходном финальном перечне ЦС' if not works[index] else ''
    from .cs_model_members import exact_members,apply_members
    apply_members(output,exact_members(folder,selected))
    output['workSource']=FINAL_WORKBOOK
    output['workSourceRevision']=revision
    output['workSourceReaderVersion']=2
    return output
